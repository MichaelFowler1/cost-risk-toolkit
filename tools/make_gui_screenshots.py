#!/usr/bin/env python3
# Copyright 2026 Michael Fowler
# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0

"""
Generate docs/gui-*.png - the pictures of the window in the README and the
getting-started guide.

Opens the real window (``ce-core gui``), puts real results in it from the
bundled examples (every number invented), and captures the screen where the
window is. Nothing is mocked up: the tiles, text and chart are what a run
shows. The one change is the folder names shown, which are swapped for a
neutral ``C:\\Users\\you\\Documents\\cost-core`` so no machine's paths end up
in the docs.

Windows only (it captures the screen). Needs Pillow, which comes with
matplotlib.

Run:  python tools/make_gui_screenshots.py
"""

from __future__ import annotations

import contextlib
import ctypes
import io
import logging
import os
import sys
import tempfile
import time
from ctypes import wintypes
from pathlib import Path

from PIL import ImageGrab

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT))

from cost_core import cli, gui  # noqa: E402

OUT = ROOT / "docs"
SHOWN = r"C:\Users\you\Documents\cost-core"


def run_demo(topic: str, out: Path) -> str:
    """Run a bundled example and return what it printed."""
    buf = io.StringIO()
    logging.disable(logging.WARNING)
    with contextlib.redirect_stdout(buf):
        cli.main(["demo", topic, "--out", str(out)])
    logging.disable(logging.NOTSET)
    return buf.getvalue()


def window_box(root) -> tuple:
    """The window's rectangle on screen, without the drop shadow."""
    hwnd = int(root.wm_frame(), 16)
    rect = wintypes.RECT()
    # DWMWA_EXTENDED_FRAME_BOUNDS = 9: the visible frame, not the shadow
    ctypes.windll.dwmapi.DwmGetWindowAttribute(hwnd, 9, ctypes.byref(rect),
                                               ctypes.sizeof(rect))
    return rect.left, rect.top, rect.right, rect.bottom


def shoot(root, name: str) -> Path:
    root.attributes("-topmost", True)
    root.lift()
    for _ in range(10):
        root.update()
        time.sleep(0.05)
    time.sleep(0.4)
    image = ImageGrab.grab(bbox=window_box(root), all_screens=True)
    path = OUT / name
    image.save(path, optimize=True)
    print(f"wrote {path.relative_to(ROOT)} ({image.width}x{image.height})")
    return path


def main() -> None:
    if sys.platform != "win32":
        sys.exit("This captures the screen, so it runs on Windows only.")
    import tkinter as tk

    work = Path(tempfile.mkdtemp(prefix="gui-shots-"))
    os.environ["APPDATA"] = str(work / "appdata")      # leave real preferences alone
    risk_out = work / "my_estimate results"
    aoa_out = work / "my_aoa results"
    risk_text = run_demo("cost-risk", risk_out)
    aoa_text = run_demo("aoa", aoa_out)

    gui.sharp()
    root = tk.Tk()
    app = gui.App(root)
    root.geometry("1120x760+40+40")
    shown = str(work)
    plain_say = app.say
    app.say = lambda parts: plain_say([(t.replace(shown, SHOWN), tag) for t, tag in parts])

    # 1. The answer to "how sure is my estimate?"
    app.choose("cost-risk")
    app.file_var.set(SHOWN + r"\my_estimate.xlsx")
    app.finish(0, risk_text, risk_out, "", None, "cost-risk")
    app.status.configure(text="Done")
    shoot(root, "gui-answer.png")

    # 2. The chart tab, on the AoA
    app.choose("aoa")
    app.file_var.set(SHOWN + r"\my_aoa.xlsx")
    app.finish(0, aoa_text, aoa_out, "", None, "aoa")
    root.update()
    app.book.select(app.chart_page)
    shoot(root, "gui-chart.png")

    # 3. A workbook with a mistake in it, checked before running
    from openpyxl import load_workbook

    book = gui.make_template(gui.task("cost-risk"), work / "my_estimate.xlsx")
    wb = load_workbook(book)
    ws = wb["Elements"]
    heads = [c.value for c in ws[1]]
    ws.cell(row=4, column=heads.index("High") + 1).value = "TBD"
    wb.save(book)
    app.choose("cost-risk")
    app.file_var.set(SHOWN + r"\my_estimate.xlsx")
    found = gui.check_workbook(gui.task("cost-risk"), book)
    app.checked(book, found)
    shoot(root, "gui-check.png")

    root.destroy()


if __name__ == "__main__":
    main()
