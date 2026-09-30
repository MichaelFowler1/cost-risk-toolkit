# Copyright 2026 Michael Fowler
# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0

"""Wrong answers found by the 2026-09-30 correctness sweep, one test each.

Every number here is invented; the lot data come from a Crawford 85% curve
with a first-unit cost of 1,000 (lot cost = sum over its units of
1000 * u ** log2(0.85)).
"""

import numpy as np
import pandas as pd
import pytest

from cost_core import cli


def lots_csv(tmp_path, units, costs, name="lots.csv", **extra):
    path = tmp_path / name
    pd.DataFrame({"units": units, "cost": costs, **extra}).to_csv(path, index=False)
    return path


def refused(argv, caplog, capsys):
    with pytest.raises(SystemExit):
        cli.main([str(a) for a in argv])
    return caplog.text + capsys.readouterr().err


# ------------------------------------------------------------ lot engine
RAMP = [5, 10, 15, 20, 25, 30]
NOISY_RAMP = [4285.61, 5399.69, 7320.85, 8280.21, 9353.93, 10378.31]


def test_an_lc_rate_fit_that_never_settles_is_not_selected(tmp_path):
    # Lot sizes that rise in step with their midpoints: the plain fixed-point
    # iteration never settled, and the unsettled LC+Rate fit won anyway with
    # a 0.00% learning slope and a first-unit cost of zero.
    from cost_core.lots import LotSeries, analyse_lots

    report = analyse_lots(LotSeries.read(lots_csv(tmp_path, RAMP, NOISY_RAMP),
                                         dollar_year=2025))
    assert report.selected_model == "LC"
    slope = 2 ** report.fit.ctx["b_lc"]
    assert 0.80 < slope < 0.90


def test_a_file_that_crashed_the_solver_now_fits(tmp_path, capsys):
    path = lots_csv(tmp_path, [10, 20, 30, 40, 50],
                    [6766.64, 10806.61, 13566.49, 14327.98, 15396.22])
    cli.main(["fit-lots", "--csv", str(path), "--dollar-year", "2025"])
    out = capsys.readouterr().out
    assert "LC selected" in out and "85." in out


def test_the_comparison_shows_the_rate_t_statistic_and_the_reason(tmp_path):
    from cost_core.lots import LotSeries, analyse_lots

    report = analyse_lots(LotSeries.read(lots_csv(tmp_path, RAMP, NOISY_RAMP),
                                         dollar_year=2025))
    items = list(report.model_comparison()["Item"])
    assert "t (rate coefficient)" in items and "Selection basis" in items


def test_a_lot_with_no_cost_in_the_middle_is_refused_not_renumbered(tmp_path):
    from cost_core.lots import LotInputError, LotSeries

    costs = [4031.09, 5829.97, None, 8422.45, 9481.82, 10445.65]
    with pytest.raises(LotInputError, match="renumber the units of every later lot"):
        LotSeries.read(lots_csv(tmp_path, RAMP, costs), dollar_year=2025)


def test_lots_with_no_cost_after_the_last_costed_lot_are_just_left_out(tmp_path):
    from cost_core.lots import LotSeries

    costs = [4031.09, 5829.97, 7229.60, 8422.45, None, None]
    assert len(LotSeries.read(lots_csv(tmp_path, RAMP, costs), dollar_year=2025).quantities) == 4


@pytest.mark.parametrize("basis,column", [("recurring", "recurring_cost"),
                                          ("total", "total_cost")])
def test_with_two_cost_columns_the_declared_basis_decides(tmp_path, basis, column):
    from cost_core.lots import LotSeries

    recurring = [7116.06, 6260.43, 6348.95, 6581.41, 6859.52, 7153.38]
    total = [r + n for r, n in zip(recurring, [20000, 8000, 3000, 0, 0, 0])]
    path = tmp_path / "both.csv"
    pd.DataFrame({"units": [10, 12, 14, 16, 18, 20], "recurring_cost": recurring,
                  "total_cost": total}).to_csv(path, index=False)
    series = LotSeries.read(path, dollar_year=2025, cost_basis=basis)
    expected = recurring if column == "recurring_cost" else total
    assert np.allclose(series.costs, expected)


def test_forecast_reprices_under_the_reports_own_selection_rule(tmp_path):
    from cost_core.lots import LotSeries, analyse_lots

    path = lots_csv(tmp_path, [10, 25, 15, 30, 20, 25],
                    [7565.38, 11278.99, 6297.79, 11091.64, 6864.72, 8203.75])
    report = analyse_lots(LotSeries.read(path, dollar_year=2025), t_gate=10.0,
                          forecast=[20, 20])
    assert report.selected_model == "LC"
    assert np.allclose(report.forecast([20, 20])["Lot Cost ($)"].to_numpy(dtype=float),
                       report.intervals()["Lot Cost ($)"].to_numpy(dtype=float))


def test_simulate_needs_a_forecast(tmp_path, caplog, capsys):
    path = lots_csv(tmp_path, [20] * 6, [12010.76, 8557.34, 7910.12, 7506.22, 7282.08, 6668.22])
    msg = refused(["fit-lots", "--csv", path, "--dollar-year", 2025, "--simulate", 2000],
                  caplog, capsys)
    assert "needs --forecast" in msg


def test_a_unit_fit_has_no_confidence_interval_for_a_lot():
    from cost_core.learning_curve import FitError, fit_curve

    units = np.arange(1, 21)
    fit = fit_curve(theory="crawford", units=units, costs=1000 * units ** np.log2(0.85))
    with pytest.raises(FitError, match="no confidence interval"):
        fit.forecast_lots([[21, 30]], kind="confidence")
    assert len(fit.forecast_lots([[21, 30]], kind="prediction")) == 1


def test_program_lots_must_be_in_fiscal_year_order():
    from cost_core.program.rollup import Program, ProgramError, fitted

    analogy = pd.DataFrame({"Lot": [1, 2, 3, 4], "Fiscal Year": [2020, 2021, 2022, 2023],
                            "Lot Quantity": [10, 12, 14, 16],
                            "AUC ($K)": [700.0, 620.0, 600.0, 580.0]})
    with pytest.raises(ProgramError, match="fiscal-year order"):
        Program("P", [2028, 2027, 2029], [fitted("Airframe", analogy, [10, 20, 30])]).validate()


# --------------------------------------------------------------- ingest
def _ingest(frames, **kw):
    import logging
    from cost_core.ingest import Crosswalk, InflationTable, normalize

    logging.disable(logging.CRITICAL)
    try:
        return normalize(frames, crosswalk=Crosswalk.default(),
                         inflation=InflationTable.from_rate(0.0235, base_year=2020),
                         base_year=2020, **kw)
    finally:
        logging.disable(logging.NOTSET)


ROW = dict(program="DEMO", period_fy=2023, lot=1, wbs_element_name="Airframe",
           recurring_flag=True, basis="BY", dollar_year=2020)


def test_rows_with_no_report_date_are_kept():
    df = pd.DataFrame([{**ROW, "cost_incurred_period": 1e6},
                       {**ROW, "wbs_element_name": "Avionics", "cost_incurred_period": 5e5}])
    ds = _ingest({"dd1921": df}, strict=True)
    assert len(ds.rows) == 2 and ds.rows["dollars"].sum() == pytest.approx(1.5e6)


def test_the_later_submission_wins_whatever_the_date_format():
    df = pd.DataFrame([{**ROW, "report_date": "3/31/2024", "cost_incurred_period": 1.0e6},
                       {**ROW, "report_date": "10/15/2024", "cost_incurred_period": 1.2e6}])
    ds = _ingest({"dd1921": df}, strict=True)
    assert list(ds.rows["dollars"]) == [1.2e6]


def test_a_then_year_row_with_no_dollar_year_is_deflated_from_its_period():
    row = {k: v for k, v in ROW.items() if k != "dollar_year"}
    df = pd.DataFrame([{**row, "basis": "TY", "period_fy": 2025, "report_date": "2025-01-01",
                        "cost_incurred_period": 1e6}])
    ds = _ingest({"dd1921": df}, strict=True)
    assert ds.rows["dollars"].iloc[0] == pytest.approx(1e6 / 1.0235 ** 5)


def test_the_reconciliation_runs_without_a_report_type_column():
    from cost_core.ingest import IngestError
    from cost_core.synth import PathologyConfig, generate_program

    reports = generate_program(3, pathologies=PathologyConfig.clean()).reports
    frames = {k: reports[k].drop(columns="report_type") for k in ("dd1921", "dd1921_1")}
    frames["dd1921_1"].loc[0, "dollars"] *= 10
    frames["flexfile"] = reports["flexfile"]
    with pytest.raises(IngestError, match="cross_report_reconciliation"):
        _ingest(frames, strict=True)


def test_flexfile_units_are_read_whatever_their_case():
    from cost_core.ingest import IngestError

    base = dict(program="DEMO", period_fy=2020, lot=1, wbs_element_name="Airframe",
                functional_category="manufacturing", recurring_flag=True, basis="BY",
                dollar_year=2020, report_date="2020-06-30")
    df = pd.DataFrame([{**base, "unit": "Hours", "value": 1000.0},
                       {**base, "unit": "Dollars", "value": 96250.0}])
    ds = _ingest({"flexfile": df})
    assert ds.rows["hours"].sum() == 1000.0 and ds.rows["dollars"].sum() == 96250.0
    with pytest.raises(IngestError, match="each must be hours or dollars"):
        _ingest({"flexfile": df.assign(unit="Each")})


def test_software_size_is_summed_across_builds():
    import logging
    from cost_core.ingest import normalize_program
    from cost_core.synth import PathologyConfig, generate_program

    logging.disable(logging.CRITICAL)
    program = generate_program(3, pathologies=PathologyConfig.clean())
    sw = normalize_program(program).software_input()
    logging.disable(logging.NOTSET)
    assert sw["equivalent_sloc"].iloc[0] == pytest.approx(494300.0)


def test_learning_curve_units_keep_the_gaps_where_lots_are_missing():
    import logging
    from cost_core.ingest import normalize_program
    from cost_core.synth import generate_program

    logging.disable(logging.CRITICAL)
    program = generate_program(4)
    lc = normalize_program(program, strict=False).learning_curve_input()
    logging.disable(logging.NOTSET)
    source = program.reports["dd1921_2"].drop_duplicates("lot").set_index("lot")
    for r in lc.itertuples():
        assert (r.first_unit, r.last_unit) == (source.loc[r.lot, "first_unit"],
                                               source.loc[r.lot, "last_unit"])


def test_then_year_sar_rows_carry_no_growth():
    from cost_core.public.panel import _growth

    uc = pd.DataFrame({"dollars": ["BY", "TY"], "base_year": [2010, None],
                       "baseline_cost": [1, 1], "baseline_quantity": [1, 1],
                       "baseline_unit_cost": [100.0, 110.0], "current_cost": [1, 1],
                       "current_quantity": [1, 1], "current_unit_cost": [120.0, 150.0],
                       "reported_pct_change": [None, None]})
    out = _growth(uc)
    assert out["unit_cost_growth_pct"].iloc[0] == pytest.approx(20.0)
    assert np.isnan(out["unit_cost_growth_pct"].iloc[1])


def test_fractional_quantities_are_refused_not_truncated(tmp_path):
    from cost_core.data_io import ValidationError, load_cost_csv

    path = tmp_path / "h.csv"
    pd.DataFrame({"program": ["A", "A"], "lot": [1, 2], "unit_quantity": [10.9, 20.6],
                  "unit_cost": [100.0, 90.0]}).to_csv(path, index=False)
    with pytest.raises(ValidationError, match="whole numbers"):
        load_cost_csv(path)


# ------------------------------------------------------------- schedules
def _task(uid, name, level=1, summary=False, dur_days=None, pct=0, start=None, finish=None,
          slack=None, links=(), milestone=False, extra="", actual_start=None,
          actual_finish=None, cost=None, fixed=None, constraint=0):
    """A tiny MSPDI task, 8-hour days, lags in days."""
    x = [f"<Task><UID>{uid}</UID><ID>{uid}</ID><Name>{name}</Name>"
         f"<OutlineLevel>{level}</OutlineLevel><Summary>{int(summary)}</Summary>"
         f"<Milestone>{int(milestone)}</Milestone>"]
    if dur_days is not None:
        x.append(f"<Duration>PT{dur_days * 8}H0M0S</Duration><DurationFormat>7</DurationFormat>"
                 f"<RemainingDuration>PT{dur_days * (100 - pct) / 100 * 8}H0M0S"
                 "</RemainingDuration>")
    x.append(f"<PercentComplete>{pct}</PercentComplete>")
    for tag, v in (("Start", start), ("Finish", finish), ("ActualStart", actual_start),
                   ("ActualFinish", actual_finish)):
        if v:
            x.append(f"<{tag}>{v}</{tag}>")
    if slack is not None:
        x.append(f"<TotalSlack>{int(slack * 4800)}</TotalSlack>")
    x.append(f"<ConstraintType>{constraint}</ConstraintType>")
    if cost is not None:
        x.append(f"<Cost>{int(cost * 100)}</Cost>")
    if fixed is not None:
        x.append(f"<FixedCost>{int(fixed * 100)}</FixedCost>")
    for link in links:
        pred, typ, lag = (list(link) + [1, 0])[:3] if isinstance(link, tuple) else (link, 1, 0)
        x.append(f"<PredecessorLink><PredecessorUID>{pred}</PredecessorUID><Type>{typ}</Type>"
                 f"<LinkLag>{int(lag * 4800)}</LinkLag><LagFormat>7</LagFormat>"
                 "</PredecessorLink>")
    return "".join(x) + extra + "</Task>"


def _project(tmp_path, tasks, status=None, finish=None):
    head = ("<?xml version=\"1.0\" encoding=\"UTF-8\"?>\n"
            "<Project xmlns=\"http://schemas.microsoft.com/project\"><Title>T</Title>"
            "<StartDate>2026-01-05T08:00:00</StartDate>")
    if finish:
        head += f"<FinishDate>{finish}</FinishDate>"
    if status:
        head += f"<StatusDate>{status}</StatusDate>"
    head += ("<MinutesPerDay>480</MinutesPerDay><MinutesPerWeek>2400</MinutesPerWeek>"
             "<DaysPerMonth>20</DaysPerMonth><Tasks><Task><UID>0</UID><ID>0</ID><Name>T</Name>"
             "<OutlineLevel>0</OutlineLevel><Summary>1</Summary></Task>")
    path = tmp_path / "ims.xml"
    path.write_text(head + "".join(tasks) + "</Tasks></Project>", encoding="utf-8")
    from cost_core.schedule.mspdi import read_mspdi
    return read_mspdi(path)


def test_check_12_fails_when_a_constraint_holds_the_finish(tmp_path):
    from cost_core.schedule.dcma import dcma_14_point

    s = _project(tmp_path, [
        _task(1, "A", dur_days=10, start="2026-01-05T08:00:00", finish="2026-01-16T17:00:00",
              slack=0, cost=10),
        _task(2, "B", dur_days=10, links=[1], start="2026-01-19T08:00:00",
              finish="2026-01-30T17:00:00", slack=0, cost=10),
        _task(3, "Launch", dur_days=0, milestone=True, links=[2], start="2026-02-13T17:00:00",
              finish="2026-02-13T17:00:00", slack=0, constraint=3,
              extra="<ConstraintDate>2026-02-13T17:00:00</ConstraintDate>")])
    row = dcma_14_point(s).table.set_index("check").loc[12]
    assert row["passed"] == False  # noqa: E712
    assert "holds the finish" in row["note"]


def test_a_lag_already_used_up_is_not_counted_again(tmp_path):
    from cost_core.schedule import critical_path

    s = _project(tmp_path, [
        _task(1, "A", dur_days=10, pct=100, start="2026-01-05T08:00:00",
              finish="2026-01-16T17:00:00", actual_start="2026-01-05T08:00:00",
              actual_finish="2026-01-16T17:00:00", cost=50),
        _task(2, "B", dur_days=10, links=[(1, 1, 10)], start="2026-03-03T08:00:00",
              finish="2026-03-16T17:00:00", cost=100)],
        status="2026-03-02T17:00:00", finish="2026-03-16T17:00:00")
    cpm = critical_path(s.to_project()).set_index("activity")
    assert cpm.loc["T2", "early_start"] == pytest.approx(0.0)
    assert cpm["early_finish"].max() == pytest.approx(0.5)


def test_a_start_link_from_a_summary_starts_with_its_first_child(tmp_path):
    from cost_core.schedule import critical_path

    s = _project(tmp_path, [
        _task(1, "Phase 1", summary=True), _task(2, "A", level=2, dur_days=10),
        _task(3, "B", level=2, dur_days=10, links=[2]),
        _task(4, "C", dur_days=15, links=[(1, 3, 0)]), _task(5, "D", dur_days=2, links=[3, 4])])
    cpm = critical_path(s.to_project())
    assert cpm["early_finish"].max() * 20 == pytest.approx(22.0)


def test_link_checks_count_the_files_own_links(tmp_path):
    from cost_core.schedule.dcma import dcma_14_point

    tasks = [_task(1, "Design", summary=True)]
    tasks += [_task(10 + i, f"D{i}", level=2, dur_days=5, links=[9 + i] if i > 1 else [])
              for i in range(1, 6)]
    tasks.append(_task(2, "Build", summary=True, links=[1]))
    kinds = {2: (21, 1, 3), 3: (22, 3, 0), 4: (23, 3, 0), 5: (24, 1, 0)}
    tasks += [_task(20 + i, f"B{i}", level=2, dur_days=5,
                    links=[kinds[i]] if i in kinds else []) for i in range(1, 6)]
    t = dcma_14_point(_project(tmp_path, tasks)).table.set_index("check")
    assert t.loc[3, "base"] == 9 and t.loc[3, "count"] == 1 and not t.loc[3, "passed"]
    assert t.loc[4, "value"] == pytest.approx(7 / 9) and not t.loc[4, "passed"]


def test_a_task_with_no_links_is_counted_even_if_it_finishes_last(tmp_path):
    from cost_core.schedule.dcma import dcma_14_point

    tasks = [_task(i, f"T{i}", dur_days=2, links=[i - 1] if i > 1 else []) for i in range(1, 21)]
    tasks.append(_task(99, "Orphan", dur_days=40, start="2026-01-12T08:00:00",
                       finish="2026-03-06T17:00:00"))
    r = dcma_14_point(_project(tmp_path, tasks))
    assert "Orphan" in r.tasks[1] and not r.table.set_index("check").loc[1, "passed"]


def test_float_checks_need_float_in_the_file(tmp_path):
    from cost_core import plain
    from cost_core.schedule.dcma import dcma_14_point

    s = _project(tmp_path, [_task(1, "A", dur_days=5, cost=10),
                            _task(2, "B", dur_days=5, links=[1], cost=10),
                            _task(3, "C", dur_days=5, links=[2], cost=10)])
    r = dcma_14_point(s)
    t = r.table.set_index("check")
    assert t.loc[6, "passed"] is None and t.loc[7, "passed"] is None
    assert "sound enough" not in " ".join(plain.dcma(r, s))


def test_a_shared_task_name_in_a_jcl_spec_is_refused(tmp_path):
    import json
    from cost_core.schedule import ScheduleError
    from cost_core.schedule.spec import load_project

    _project(tmp_path, [_task(1, "Bus", summary=True),
                        _task(2, "Design review", level=2, dur_days=5),
                        _task(3, "Payload", summary=True),
                        _task(4, "Design review", level=2, dur_days=5)])
    spec = tmp_path / "spec.json"
    spec.write_text(json.dumps({"mspdi": "ims.xml", "risks": [
        {"name": "Slip", "probability": 1.0, "activities": ["Design review"], "delay": 1}]}))
    with pytest.raises(ScheduleError, match="are named 'Design review'"):
        load_project(spec)


def test_remaining_fixed_cost_never_exceeds_the_files_remaining_cost(tmp_path):
    from cost_core.schedule.jcl import point_estimate

    s = _project(tmp_path, [_task(1, "Buy", dur_days=20, pct=50, cost=100, fixed=80,
                                  start="2026-01-05T08:00:00",
                                  actual_start="2026-01-05T08:00:00",
                                  extra="<FixedCostAccrual>1</FixedCostAccrual>"
                                        "<RemainingCost>1000</RemainingCost>")])
    assert point_estimate(s.to_project())[1] == pytest.approx(100.0)


def test_an_opportunity_larger_than_the_activity_floors_it_at_zero():
    from cost_core.schedule.jcl import Activity, Project, Risk, simulate

    p = Project([Activity("A", 2, burn_rate=10),
                 Activity("B", 4, predecessors=["A"], burn_rate=1)],
                risks=[Risk("Early", 1.0, ["A"], delay=-3)])
    r = simulate(p, n_iter=10)
    assert (r.cost >= 0).all() and r.cost[0] == pytest.approx(4.0)


# ------------------------------------------------------------------- EVM
def _evm(rows):
    from cost_core.evm import EvmData
    return EvmData.from_frame(pd.DataFrame(rows))


@pytest.mark.parametrize("labels", [["Oct-25", "Nov-25", "Dec-25", "Jan-26", "Feb-26", "Mar-26"],
                                    ["Oct 2025", "Nov 2025", "Dec 2025", "Jan 2026",
                                     "Feb 2026", "Mar 2026"]])
def test_month_year_period_labels_sort_in_time_order(labels):
    d = _evm([{"period": p, "wbs": w, "bcws": 10, "bcwp": 8 if i < 4 else None,
               "acwp": 10 if i < 4 else None} for w in "AB" for i, p in enumerate(labels)])
    r = d.metrics().iloc[-1]
    assert r["period"] == labels[3] and r["spi"] == pytest.approx(0.8)
    assert r["sv_t"] == pytest.approx(-0.8)


def test_numbered_period_labels_sort_as_numbers():
    d = _evm([{"period": f"Month {i + 1}", "bcws": 10, "bcwp": 8 if i < 4 else None,
               "acwp": 10 if i < 4 else None} for i in range(10)])
    assert d.metrics().iloc[-1]["period"] == "Month 4"


def test_tcpi_warnings_fire_once_cost_passes_the_budget_or_the_eac():
    d = _evm({"period": [1, 2, 3, 4], "bcws": [25] * 4, "bcwp": [20, 20, 20, None],
              "acwp": [40, 40, 40, None], "eac": [None, None, 110, None]})
    f = d.flags().set_index("flag")
    assert f.loc["TCPI to BAC above 1.10", "raised"]
    assert f.loc["TCPI to the EAC exceeds the CPI by more than 0.10", "raised"]
    assert len(f) == 6


def test_a_finished_accounts_eac_still_counts():
    a = pd.DataFrame({"period": [1, 2], "wbs": "A", "bcws": [10, 10], "bcwp": [10, 10],
                      "acwp": [12, 12], "eac": [24, 24]})
    b = pd.DataFrame({"period": [1, 2, 3, 4], "wbs": "B", "bcws": [20] * 4,
                      "bcwp": [15, 15, 15, None], "acwp": [20, 20, 20, None],
                      "eac": [None, None, 100, None]})
    from cost_core.evm import EvmData
    assert EvmData.from_frame(pd.concat([a, b])).eac[-1] == pytest.approx(124.0)


def test_an_account_that_stops_reporting_is_noted():
    d = _evm({"period": [1, 2, 3, 4] * 2, "wbs": ["A"] * 4 + ["B"] * 4, "bcws": [10] * 8,
              "bcwp": [10, 10, 10, None, 10, 10, None, None],
              "acwp": [10, 10, 10, None, 10, 10, None, None]})
    assert any("Account B reports nothing after 2" in n for n in d.notes)


def test_earned_schedule_counts_the_periods_that_planned_nothing():
    from cost_core.evm import earned_schedule
    assert earned_schedule([0, 0, 0, 0, 50, 100], 0) == pytest.approx(4.0)
    assert earned_schedule([10, 30, 60, 100], 0) == pytest.approx(0.0)


def test_the_eac_warning_uses_the_same_ieacs_as_the_summary():
    d = _evm({"period": [1, 2, 3, 4], "bcws": [40, 40, 10, 10], "bcwp": [45, 40, None, None],
              "acwp": [45, 40, None, None], "eac": [None, 98, None, None]})
    summary = d.summary().set_index("measure")["value"]
    flag = d.flags().set_index("flag").loc["Contractor EAC below every independent EAC"]
    assert (98 < float(summary["independent EACs: lowest"])) == bool(flag["raised"])


def test_a_negative_first_period_quotes_zero_as_the_prior_value():
    from cost_core.evm.checks import data_checks
    d = _evm({"period": [1, 2, 3, 4], "bcws": [10, 20, 30, 40], "bcwp": [-5, 25, 30, None],
              "acwp": [4, 20, 28, None]})
    assert "(cumulative 0.00 to -5.00)" in data_checks(d).iloc[0]["detail"]


def test_ipmdar_earned_schedule_is_blank_where_the_baseline_was_invented(tmp_path):
    from test_ipmdar import cpd_tables, write_folder
    from cost_core.evm.ipmdar import read_ipmdar

    r = read_ipmdar(write_folder(cpd_tables(status=14, time_phased=False), tmp_path / "c")
                    ).metrics().iloc[-1]
    assert np.isnan(r["es"]) and np.isnan(r["spi_t"])


def test_ipmdar_bac_is_reconciled_with_the_pmb(tmp_path):
    from test_ipmdar import cpd_tables, write_folder
    from cost_core.evm.ipmdar import read_ipmdar

    t = cpd_tables()
    t["BCWS_ToComplete"] = [x for x in t["BCWS_ToComplete"] if x["ControlAccountID"] != "CA2"]
    notes = read_ipmdar(write_folder(t, tmp_path / "s")).notes
    assert any("BAC of" in n and "PMB's" in n for n in notes)


def test_a_dollar_variance_threshold_given_alone_decides_alone(tmp_path, capsys):
    from cost_core.examples import example_path
    cli.main(["evm", "--data", str(example_path("evm")), "--cv-dollars", "50",
              "--sv-dollars", "50", "--out", str(tmp_path / "o"), "--iters", "1000"])
    capsys.readouterr()
    assert len(pd.read_csv(tmp_path / "o" / "variance_reports.csv")) == 3


# ------------------------------------------------------------- cost risk
def _risk_workbook(path, elements, risks=(), pairs=(), settings=(("Seed", 1),),
                   sheet_names=("Elements", "Risks", "Correlation", "Settings"), extra=None):
    with pd.ExcelWriter(path) as xl:
        pd.DataFrame(list(elements), columns=["Element", "Point Estimate", "Low",
                                              "Most Likely", "High"]
                     ).to_excel(xl, sheet_name=sheet_names[0], index=False)
        pd.DataFrame(list(risks), columns=["Risk", "Probability", "Low", "Most Likely", "High"]
                     ).to_excel(xl, sheet_name=sheet_names[1], index=False)
        pd.DataFrame(list(pairs), columns=["Element A", "Element B", "Correlation"]
                     ).to_excel(xl, sheet_name=sheet_names[2], index=False)
        pd.DataFrame(list(settings), columns=["Setting", "Value"]
                     ).to_excel(xl, sheet_name=sheet_names[3], index=False)
        if extra:
            pd.DataFrame({"x": [1]}).to_excel(xl, sheet_name=extra, index=False)
    return path


def test_the_s_curve_labels_costs_in_their_units():
    from cost_core.reporting.charts import _cost_labels

    assert _cost_labels("millions")[0](2000.0) == "$2.00B"
    assert _cost_labels("$K")[0](18_500.0) == "$18.5M"
    assert _cost_labels("")[0](18.5) == "18.5"
    assert _cost_labels(None)[0](2000.0) == "$2K"


def test_a_credit_entered_below_zero_is_simulated_below_zero():
    from cost_core.monte_carlo import CostElement, DiscreteRisk, RiskModel, simulate_risk_model

    model = RiskModel(elements=[
        CostElement("Hardware", {"type": "triangular", "left": 90, "mode": 100, "right": 130}),
        CostElement("Credit", {"type": "triangular", "left": -14, "mode": -10, "right": -6})],
        risks=[DiscreteRisk("Reuse", 1.0, {"type": "uniform", "low": -12, "high": -4})],
        correlation=np.eye(2))
    r = simulate_risk_model(model, n_iter=20_000, seed=1)
    assert r.element_samples[:, 1].mean() == pytest.approx(-10.0, abs=0.1)
    assert r.risk_samples[:, 0].mean() == pytest.approx(-8.0, abs=0.1)


def test_an_element_with_no_range_is_carried_at_its_point_estimate(tmp_path):
    from cost_core.costrisk import read_workbook

    p = _risk_workbook(tmp_path / "r.xlsx", [("Hardware", 100, 90, 100, 130),
                                             ("Fee", 50, None, 40, None),
                                             ("Spares", 12, 8, 8, 8)])
    inp = read_workbook(p)
    fee, spares = inp.model.elements[1], inp.model.elements[2]
    assert fee.distribution == {"type": "fixed", "value": 50.0}
    assert spares.distribution == {"type": "fixed", "value": 8.0}
    assert any("'Spares' has its Low, Most Likely and High all at 8" in n for n in inp.notes)
    assert any("1 element (Fee) has no range" in n for n in inp.notes)


def test_own_p80s_that_add_to_less_than_the_total_are_described_as_such(tmp_path):
    from cost_core import plain
    from cost_core.costrisk import analyse, read_workbook

    p = _risk_workbook(tmp_path / "r.xlsx", [("Airframe", 100, 95, 100, 110),
                                             ("Avionics", 60, 57, 60, 66)],
                       risks=[(f"Risk {i}", 0.15, 8, 10, 12) for i in range(1, 6)],
                       settings=[("Units", "millions"), ("Iterations", 20000), ("Seed", 1)])
    text = " ".join(plain.cost_risk(analyse(read_workbook(p)), "$M"))
    assert "less than the P80 of the whole" in text and "more than the P80" not in text
    assert "goes mostly" not in text and "the largest part of" in text


def test_near_certainty_is_not_rounded_to_a_hundred_percent():
    from cost_core import plain

    assert plain._chance(0.996) == "more than a 99% chance"
    assert plain._pct(0.996) == "over 99%"


def test_a_risk_register_sheet_is_read_and_an_unread_sheet_is_named(tmp_path):
    from cost_core.costrisk import read_workbook

    p = _risk_workbook(tmp_path / "r.xlsx", [("A", 100, 90, 100, 130), ("B", 50, 45, 50, 70)],
                       risks=[("Slip", 0.5, 5, 10, 20)], pairs=[("A", "B", 0.6)],
                       sheet_names=("Elements", "Risk Register", "Correlations", "Settings"),
                       extra="Scratch")
    inp = read_workbook(p)
    assert len(inp.model.risks) == 1 and len(inp.pairs) == 1
    assert any("'Scratch' was not read" in n for n in inp.notes)


def test_a_correlation_pair_entered_twice_is_refused(tmp_path):
    from cost_core.costrisk import CostRiskError, read_workbook

    p = _risk_workbook(tmp_path / "r.xlsx", [("A", 100, 90, 100, 130), ("B", 50, 45, 50, 70)],
                       pairs=[("A", "B", 0.2), ("B", "A", 0.8)])
    with pytest.raises(CostRiskError, match="already paired on row 2"):
        read_workbook(p)


def test_a_settings_sheet_with_no_heading_row_keeps_its_first_setting(tmp_path):
    from cost_core.costrisk import read_workbook
    from cost_core.xlspec import read_spec

    p = tmp_path / "r.xlsx"
    with pd.ExcelWriter(p) as xl:
        pd.DataFrame({"Element": ["A", "B"], "Point Estimate": [100, 50], "Low": [90, 45],
                      "Most Likely": [100, 50], "High": [140, 70]}
                     ).to_excel(xl, sheet_name="Elements", index=False)
        pd.DataFrame([("Default Correlation", 0.6), ("Iterations", 5000)]).to_excel(
            xl, sheet_name="Settings", index=False, header=False)
    assert read_workbook(p).model.default_correlation == pytest.approx(0.6)

    q = tmp_path / "aoa.xlsx"
    with pd.ExcelWriter(q) as xl:
        pd.DataFrame([("Base Year", 2026), ("Discount Rate", 0.02), ("Inflation Rate", 0.02)]
                     ).to_excel(xl, sheet_name="Settings", index=False, header=False)
        pd.DataFrame({"Alternative": ["A"], "Line": ["Dev"], "Phase": ["RDT&E"],
                      "Total": [600], "Start Year": [2027], "Years": [3]}
                     ).to_excel(xl, sheet_name="Lines", index=False)
    assert read_spec(q)["base_year"] == 2026


def test_workbook_rows_that_would_be_dropped_are_refused(tmp_path):
    from cost_core.xlspec import SpecError, read_spec

    b = tmp_path / "b.xlsx"
    with pd.ExcelWriter(b) as xl:
        pd.DataFrame({"Candidate": ["A"], "Option": ["Full"], "Value": [10], 2027: [300]}
                     ).to_excel(xl, sheet_name="Candidates", index=False)
        pd.DataFrame({"Year": [2027, 2027], "Budget": [400, 300]}
                     ).to_excel(xl, sheet_name="Budget", index=False)
    with pytest.raises(SpecError, match="2027 is already on the sheet"):
        read_spec(b)

    settings = pd.DataFrame({"Setting": ["Base Year", "Discount Rate", "Inflation Rate"],
                             "Value": [2026, 0.02, 0.02]})
    both = tmp_path / "both.xlsx"
    with pd.ExcelWriter(both) as xl:
        settings.to_excel(xl, sheet_name="Settings", index=False)
        pd.DataFrame({"Alternative": ["A"], "Line": ["Dev"], "Phase": ["RDT&E"],
                      "Total": [600], "Start Year": [2027], "Years": [3],
                      "Annual Amount": [100], "First Year": [2030], "Last Year": [2039]}
                     ).to_excel(xl, sheet_name="Lines", index=False)
    with pytest.raises(SpecError, match="a Total and an Annual Amount"):
        read_spec(both)

    ph = tmp_path / "ph.xlsx"
    with pd.ExcelWriter(ph) as xl:
        settings.to_excel(xl, sheet_name="Settings", index=False)
        pd.DataFrame({"Alternative": ["A"], "Line": ["Disposal"], "Phase": ["Disposal"]}
                     ).to_excel(xl, sheet_name="Lines", index=False)
        pd.DataFrame({"Alternative": ["A", "A"], "Line": ["Disposal", "Disposal"],
                      2040: [40, None], 2041: [None, 40]}
                     ).to_excel(xl, sheet_name="Phased", index=False)
    with pytest.raises(SpecError, match="already has row 2"):
        read_spec(ph)


def test_a_jcl_spec_survives_the_workbook_round_trip(tmp_path):
    from cost_core.xlspec import SpecError, read_spec, write_workbook

    spec = {"name": "t", "units": "$M", "activities": [
        {"id": "design", "duration": 12},
        {"id": "fsw", "duration": 9, "predecessors": [["design", 0.333333333, "FF"]]}],
        "risks": [{"name": "r1", "probability": 0.3, "activities": ["fsw"],
                   "delay": {"type": "uniform", "low": 1.0, "high": 3.0}}]}
    back = read_spec(write_workbook(spec, "jcl", tmp_path / "j.xlsx"))
    assert back["risks"][0]["delay"] == {"type": "uniform", "low": 1.0, "high": 3.0}
    assert back["activities"][1]["predecessors"][0][1] == pytest.approx(0.333333333, abs=1e-12)

    spec["activities"][0]["duration_uncertainty"] = {"type": "pert", "left": 0.9, "mode": 1.0,
                                                     "right": 1.5, "lambda_": 6}
    with pytest.raises(SpecError, match="lambda_ 6 has no workbook form"):
        write_workbook(spec, "jcl", tmp_path / "k.xlsx")


def test_the_cost_risk_assumptions_record_the_seed_used(tmp_path):
    from cost_core.costrisk import analyse, read_workbook

    p = _risk_workbook(tmp_path / "r.xlsx", [("A", 100, 90, 100, 130), ("B", 50, 45, 50, 70)])
    assert analyse(read_workbook(p), n_iter=2000, seed=7).assumptions["seed"] == 7


def test_the_cost_risk_workbook_formats_money_in_its_units(tmp_path):
    from openpyxl import load_workbook
    from cost_core.costrisk import analyse, read_workbook
    from cost_core.reporting.excel_report import cost_risk_workbook

    p = _risk_workbook(tmp_path / "r.xlsx", [("A", 18.5, 17, 18.5, 25), ("B", 5, 4, 5, 7)],
                       settings=[("Units", "millions"), ("Seed", 1)])
    out = cost_risk_workbook(analyse(read_workbook(p), n_iter=2000), tmp_path / "report.xlsx")
    formats = {c.number_format for row in load_workbook(out)["Summary"].iter_rows()
               for c in row if isinstance(c.value, float)}
    assert '"$"#,##0.0"M"' in formats


def test_an_ampersand_in_the_marking_prints_as_one(tmp_path):
    from openpyxl import Workbook
    from cost_core.reporting.output import configured, mark_workbook

    wb = Workbook()
    with configured(marking="R&D ONLY"):
        mark_workbook(wb)
    assert wb.active.oddHeader.center.text == "R&&D ONLY"


# ------------------------------------------------- AoA, portfolio, inflate
def _alt(name, cost, spec=None, effectiveness=None):
    from cost_core.aoa import Alternative, CostLine
    return Alternative(name, [CostLine("Buy", "Procurement", {2027: float(cost)}, spec)],
                       effectiveness)


def _evaluate(alts, **kw):
    from cost_core.aoa import evaluate
    from cost_core.ingest import InflationTable
    infl = InflationTable.from_rate(0.0, base_year=2026, first_year=2026, last_year=2030)
    return evaluate(alts, base_year=2026, inflation=infl, discount_rate=0.0,
                    **{"basis": "by", "n_iter": 20000, "seed": 1, **kw})


def test_two_options_with_one_name_are_refused():
    from cost_core.portfolio import Candidate, Option, PortfolioError

    with pytest.raises(PortfolioError, match="used twice"):
        Candidate("Radar", [Option("Fund", {2027: 60}, 3), Option("Fund", {2027: 40}, 13)])


def test_mandatory_no_means_no(tmp_path):
    import json
    from cost_core.portfolio import PortfolioError
    from cost_core.portfolio.spec import load_portfolio

    spec = {"budget": {"2027": 100}, "candidates": [
        {"name": "Radar", "mandatory": "no",
         "options": [{"name": "Full", "value": 2, "cost_by_year": {"2027": 60}}]}]}
    p = tmp_path / "p.json"
    p.write_text(json.dumps(spec))
    assert load_portfolio(p)[0].candidates[0].mandatory is False
    spec["candidates"][0]["mandatory"] = "maybe"
    p.write_text(json.dumps(spec))
    with pytest.raises(PortfolioError, match="true or false"):
        load_portfolio(p)


def test_the_aoa_summary_does_not_contradict_itself():
    from cost_core import plain

    tri = {"type": "triangular", "left": 0.9, "mode": 0.95, "right": 1.6}
    r = _evaluate([_alt("Risky", 100, tri, 0.6), _alt("Steady", 113, None, 0.7)])
    text = " ".join(plain.aoa(r, "$M"))
    assert "most likely" not in text
    assert "'Risky' costs more on average than 'Steady'" in text
    assert "Its 50% cost is lower" in text


def test_a_tie_for_cheapest_is_shared():
    r = _evaluate([_alt("A", 100), _alt("B", 100)], n_iter=2000)
    assert list(r.summary["p_cheapest"]) == [0.5, 0.5]


def test_an_unscored_alternative_is_not_on_the_frontier():
    r = _evaluate([_alt("Cheap", 100, None, 0.8), _alt("Unscored", 300),
                   _alt("Middle", 200, None, 0.5)], n_iter=2000)
    s = r.summary.set_index("alternative")
    assert not s.loc["Unscored", "on_frontier"] and s.loc["Cheap", "on_frontier"]


def test_an_alternatives_draws_do_not_depend_on_its_place_in_the_list():
    tri = {"type": "triangular", "left": 0.8, "mode": 1.0, "right": 1.6}
    a, b, c = _alt("A", 100, tri), _alt("B", 110, tri), _alt("C", 90, tri)
    r1, r2 = _evaluate([a, b, c], n_iter=2000), _evaluate([c, a, b], n_iter=2000)
    assert np.array_equal(r1.draws["A"], r2.draws["A"])


def test_ce_core_toml_fills_in_what_the_input_file_leaves_out(tmp_path, monkeypatch, capsys):
    import json
    from cost_core.examples import example_path

    monkeypatch.chdir(tmp_path)
    (tmp_path / "ce-core.toml").write_text("seed = 11\niterations = 3000\n", encoding="utf-8")
    spec = json.loads(example_path("aoa").read_text())
    for key in ("seed", "n_iter"):
        spec.pop(key, None)
    (tmp_path / "aoa.json").write_text(json.dumps(spec))
    cli.main(["aoa", "--spec", "aoa.json", "--out", "a"])
    record = json.loads((tmp_path / "a" / "assumptions.json").read_text())
    assert record["seed"] == 11 and record["n_iter"] == 3000

    p = _risk_workbook(tmp_path / "r.xlsx", [("A", 100, 90, 100, 130), ("B", 50, 45, 50, 70)],
                       settings=[("Iterations", 2000)])
    cli.main(["cost-risk", "--data", str(p), "--out", "c"])
    capsys.readouterr()
    from openpyxl import load_workbook
    ws = load_workbook(tmp_path / "c" / "report.xlsx")["Assumptions"]
    got = {r[0]: r[1] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert got["seed"] == 11 and got["iterations"] == 2000


def test_inflate_reads_its_own_output_and_hyphenated_names(tmp_path, capsys):
    index = tmp_path / "idx.csv"
    pd.DataFrame({"index_name": "i", "fiscal_year": range(2025, 2031),
                  "index_value": [1.02 ** k for k in range(6)]}).to_csv(index, index=False)
    data = tmp_path / "d.csv"
    pd.DataFrame({"fy": [2027, 2028], "cost-basis": [10.0, 20.0]}).to_csv(data, index=False)
    cli.main(["inflate", "--data", str(data), "--index", str(index), "--from", "BY2025",
              "--to", "TY", "--amount-col", "cost-basis", "--out", str(tmp_path / "o1.csv")])
    cli.main(["inflate", "--data", str(tmp_path / "o1.csv"), "--index", str(index), "--from",
              "BY2025", "--to", "TY", "--amount-col", "cost-basis",
              "--out", str(tmp_path / "o2.csv")])
    out = capsys.readouterr().out
    assert "cost-basis  cost_basis_then_year_2" in out
    again = pd.read_csv(tmp_path / "o2.csv")
    assert list(again["cost_basis_then_year"]) == list(again["cost_basis_then_year_2"])


def test_inflate_prints_the_amount_it_was_given(tmp_path, capsys):
    index = tmp_path / "idx.csv"
    pd.DataFrame({"index_name": "i", "fiscal_year": range(2025, 2031),
                  "index_value": [1.02 ** k for k in range(6)]}).to_csv(index, index=False)
    cli.main(["inflate", "--amount", "1234567.5", "--index", str(index), "--from", "BY2025",
              "--to", "TY", "--year", "2027"])
    assert "1,234,567.5 in BY2025 dollars" in capsys.readouterr().out


def test_the_portfolio_report_records_what_it_ran_with(tmp_path, capsys):
    from openpyxl import load_workbook
    from cost_core.examples import example_path

    pytest.importorskip("pulp")
    cli.main(["portfolio", "--spec", str(example_path("portfolio")), "--out", str(tmp_path)])
    text = capsys.readouterr().out
    assert "as history says" not in text
    ws = load_workbook(tmp_path / "report.xlsx")["Assumptions"]
    got = {r[0] for r in ws.iter_rows(min_row=2, values_only=True)}
    assert {"delta", "growth", "growth correlation", "iterations", "seed"} <= got
