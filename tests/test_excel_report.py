# Copyright 2026 Michael Fowler
# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0

"""The Excel workbook every EVM, JCL, schedule check, AoA and portfolio run
writes. The EVM metrics are Excel formulas, so where LibreOffice is installed
the workbook is recalculated by it and every formula checked against the
numbers Python computed."""
import shutil
import subprocess

import openpyxl
import pytest

from cost_core import cli
from cost_core.evm import EvmData
from cost_core.examples import example_path

TOPICS = ["evm", "schedule", "jcl", "aoa", "portfolio"]


@pytest.fixture(scope="module")
def demos(tmp_path_factory):
    root = tmp_path_factory.mktemp("demos")
    for topic in TOPICS:
        if topic == "portfolio":
            pytest.importorskip("pulp")
        cli.main(["demo", topic, "--out", str(root / topic)])
    return root


@pytest.mark.parametrize("topic", TOPICS)
def test_every_run_writes_a_workbook_summary_first_and_assumptions_last(demos, topic):
    wb = openpyxl.load_workbook(demos / topic / "report.xlsx")
    assert wb.sheetnames[0] == "Summary" and wb.sheetnames[-1] == "Assumptions"
    summary = [c.value for row in wb["Summary"].iter_rows(max_col=1) for c in row if c.value]
    assert "What this means" in summary and "Key numbers" in summary
    for ws in wb.worksheets:
        assert ws.page_setup.fitToWidth == 1  # prints one page wide
        for chart in ws._charts:
            assert all(s.smooth is False for s in chart.series)  # no invented curves


def test_the_evm_workbook_shows_its_arithmetic_as_formulas(demos):
    ws = openpyxl.load_workbook(demos / "evm" / "report.xlsx")["Metrics"]
    heads = [c.value for c in ws[3]]
    row = {h: ws.cell(row=4, column=j + 1).value for j, h in enumerate(heads)}
    assert row["CPI"] == '=IF(E4=0,"",D4/E4)' and row["CV"] == "=D4-E4"
    assert row["TCPI (to BAC)"].startswith("=IF($C$1-E4=0") and row["SPI(t)"].startswith("=")
    assert isinstance(row["BCWP (cum)"], float) and isinstance(row["ES"], float)
    wb = openpyxl.load_workbook(demos / "evm" / "report.xlsx")
    assert {"Warning signs", "Forecast", "Accounts"} <= set(wb.sheetnames)


def _calc_available():
    exe = shutil.which("soffice")
    if not exe:
        return None
    return exe


@pytest.mark.skipif(_calc_available() is None, reason="LibreOffice is not installed")
def test_the_evm_formulas_give_what_python_computed(demos, tmp_path):
    src = demos / "evm" / "report.xlsx"
    done = subprocess.run(
        [_calc_available(), "--headless", "--norestore",
         f"-env:UserInstallation=file://{tmp_path}/profile",
         "--convert-to", "xlsx", "--outdir", str(tmp_path / "out"), str(src)],
        capture_output=True, text=True, timeout=240)
    out = tmp_path / "out" / "report.xlsx"
    if not out.exists():
        pytest.skip(f"LibreOffice could not open spreadsheets here: {done.stderr.strip()[-120:]}")
    ws = openpyxl.load_workbook(out, data_only=True)["Metrics"]
    m = EvmData.read(example_path("evm")).metrics()
    heads = [c.value for c in ws[3]]
    cols = {"CV": "cv", "SV": "sv", "CPI": "cpi", "SPI": "spi", "% complete": "pct_complete",
            "TCPI (to BAC)": "tcpi_bac", "IEAC (CPI)": "ieac_cpi", "SPI(t)": "spi_t"}
    for i, r in enumerate(m.itertuples(index=False), start=4):
        for head, name in cols.items():
            got = ws.cell(row=i, column=heads.index(head) + 1).value
            assert got == pytest.approx(getattr(r, name), rel=1e-12), (i, head)
