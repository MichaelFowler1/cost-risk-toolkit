# Copyright 2026 Michael Fowler
# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0

"""What makes cost-core usable on day one: settings, markings, house slide
templates, ce-core open, ce-core inflate, the EVM monthly checks and the
cost-risk reserve allocation."""

import io
import json
import shutil
import zipfile
from pathlib import Path

import numpy as np
import pandas as pd
import pytest

from cost_core import cli, opener, settings
from cost_core.examples import example_path
from cost_core.reporting import output


@pytest.fixture(autouse=True)
def clean_output():
    output.configure(None, None)
    yield
    output.configure(None, None)


def run(argv, capsys):
    cli.main([str(a) for a in argv])
    return capsys.readouterr().out


def refused(argv, caplog, capsys):
    with pytest.raises(SystemExit):
        cli.main([str(a) for a in argv])
    return caplog.text + capsys.readouterr().err


# --------------------------------------------------------------- settings
def test_the_folder_file_wins_over_the_home_file(tmp_path):
    home, here = tmp_path / "home", tmp_path / "here"
    home.mkdir(), here.mkdir()
    (home / "ce-core.toml").write_text('marking = "HOME"\nseed = 4\n', encoding="utf-8")
    (here / "ce-core.toml").write_text('marking = "HERE"\n', encoding="utf-8")
    loaded = settings.load(cwd=here, home=home)
    assert loaded["marking"] == ("HERE", str(here / "ce-core.toml"))
    assert loaded["seed"] == (4, str(home / "ce-core.toml"))
    assert loaded["fiscal_year_start"] == (10, "built-in default")


@pytest.mark.parametrize("text,message", [
    ('marking = "CUI"\nmarkng = "x"\n', "'markng' is not a setting"),
    ("seed = 1.5\n", "seed is a whole number"),
    ("seed = -1\n", "seed is 0 or more"),
    ("iterations = 500\n", "iterations is 1,000 or more"),
    ("fiscal_year_start = 13\n", "fiscal_year_start is a month"),
    ("marking = 5\n", "marking is text, in quotes"),
    ("marking = CUI\n", "isn't valid TOML"),
])
def test_a_bad_settings_file_says_which_line(tmp_path, text, message):
    (tmp_path / "ce-core.toml").write_text(text, encoding="utf-8")
    with pytest.raises(settings.SettingsError, match=message):
        settings.load(cwd=tmp_path, home=tmp_path / "nohome")


def test_a_slide_template_is_relative_to_the_file_that_names_it(tmp_path):
    (tmp_path / "ce-core.toml").write_text('slide_template = "house.pptx"\n', encoding="utf-8")
    value = settings.values(cwd=tmp_path, home=tmp_path / "nohome")["slide_template"]
    assert Path(value) == tmp_path / "house.pptx"


def test_settings_fill_in_options_but_a_flag_wins(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(Path, "home", lambda: tmp_path / "nohome")
    (tmp_path / "ce-core.toml").write_text('units = "millions"\nseed = 9\n', encoding="utf-8")
    out = run(["evm", "--data", example_path("evm"), "--out", "o"], capsys)
    assert "costs in $M" in out
    assert json.loads((tmp_path / "o" / "assumptions.json").read_text())["seed"] == 9
    out = run(["evm", "--data", example_path("evm"), "--out", "o2", "--units", "thousands"],
              capsys)
    assert "costs in $K" in out


def test_the_settings_command_shows_values_and_where_they_came_from(tmp_path, monkeypatch,
                                                                      capsys):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(Path, "home", lambda: tmp_path / "nohome")
    out = run(["settings", "--write"], capsys)
    assert (tmp_path / "ce-core.toml").is_file() and "Wrote" in out
    (tmp_path / "ce-core.toml").write_text('marking = "CUI"\n', encoding="utf-8")
    out = run(["settings"], capsys)
    assert "CUI" in out and str(tmp_path / "ce-core.toml") in out and "built-in default" in out


def test_version_and_about(capsys):
    from cost_core import __version__

    with pytest.raises(SystemExit):
        cli.main(["--version"])
    assert __version__ in capsys.readouterr().out
    out = run(["--about"], capsys)
    assert f"cost-core {__version__}" in out and "numpy" in out and "never sends" in out


# ---------------------------------------------------------------- marking
def test_the_marking_is_on_every_sheet_and_every_slide(tmp_path, capsys):
    import openpyxl
    pytest.importorskip("pptx")
    from pptx import Presentation

    run(["cost-risk", "--data", example_path("cost-risk"), "--iters", 2000,
         "--out", tmp_path, "--marking", "CUI//TEST"], capsys)
    wb = openpyxl.load_workbook(tmp_path / "report.xlsx")
    for ws in wb.worksheets:
        assert ws.oddHeader.center.text == ws.oddFooter.center.text == "CUI//TEST"
    assert wb["Summary"]["A3"].value == "CUI//TEST"
    for slide in Presentation(tmp_path / "brief.pptx").slides:
        marks = [s for s in slide.shapes if s.name == "Marking"]
        assert [m.text_frame.text for m in marks] == ["CUI//TEST", "CUI//TEST"]


def test_no_marking_means_nothing_is_stamped(tmp_path, capsys):
    import openpyxl

    run(["cost-risk", "--data", example_path("cost-risk"), "--iters", 2000, "--out", tmp_path],
        capsys)
    ws = openpyxl.load_workbook(tmp_path / "report.xlsx")["Summary"]
    assert not ws.oddHeader.center.text and ws["A3"].value is None


def _template(path: Path, width_in=10.0, drop_layouts=False) -> Path:
    from pptx import Presentation
    from pptx.util import Inches

    prs = Presentation()
    prs.slide_width, prs.slide_height = Inches(width_in), Inches(7.5)
    prs.slides.add_slide(prs.slide_layouts[0]).shapes.title.text = "SAMPLE"
    if drop_layouts:
        for layout in prs.slide_layouts:
            layout.name = "Custom " + layout.name
    buf = io.BytesIO()
    prs.save(buf)
    if path.suffix == ".pptx":
        path.write_bytes(buf.getvalue())
        return path
    src = zipfile.ZipFile(buf)
    with zipfile.ZipFile(path, "w") as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.presentation.main+xml",
                                    b"presentationml.template.main+xml")
            out.writestr(item, data)
    return path


def test_a_house_template_is_used_by_layout_name_and_everything_fits(tmp_path, capsys):
    pytest.importorskip("pptx")
    from pptx import Presentation

    template = _template(tmp_path / "house.potx", width_in=10.0)
    run(["cost-risk", "--data", example_path("cost-risk"), "--iters", 2000,
         "--out", tmp_path / "o", "--template", template], capsys)
    prs = Presentation(tmp_path / "o" / "brief.pptx")
    assert prs.slide_width == template_width(template)
    layouts = [s.slide_layout.name for s in prs.slides]
    assert layouts[0] == "Title Slide" and set(layouts[1:]) == {"Title Only"}
    assert all(s.shapes.title.text != "SAMPLE" for s in prs.slides)
    for slide in prs.slides:
        for shape in slide.shapes:
            assert shape.left + shape.width <= prs.slide_width + 10
            assert shape.top + shape.height <= prs.slide_height + 10
    pics = [s for slide in prs.slides for s in slide.shapes if s.shape_type == 13]
    assert pics and all(p._element.nvPicPr.cNvPr.get("descr") for p in pics)


def template_width(path):
    from pptx.util import Inches
    return Inches(10.0)


def test_a_template_without_the_named_layouts_falls_back_and_says_so(tmp_path, capsys, caplog):
    pytest.importorskip("pptx")
    template = _template(tmp_path / "odd.pptx", drop_layouts=True)
    run(["cost-risk", "--data", example_path("cost-risk"), "--iters", 2000,
         "--out", tmp_path / "o", "--template", template], capsys)
    assert "has no 'title slide' or 'title only' layout" in caplog.text
    assert (tmp_path / "o" / "brief.pptx").is_file()


def test_a_missing_template_is_refused(tmp_path, caplog, capsys):
    msg = refused(["cost-risk", "--data", example_path("cost-risk"), "--out", tmp_path,
                   "--template", tmp_path / "nope.potx"], caplog, capsys)
    assert "slide template" in msg and "isn't there" in msg


# ------------------------------------------------------------------- open
@pytest.mark.parametrize("topic,command", [
    ("evm", "evm"), ("schedule", "schedule-check"), ("jcl", "jcl"), ("aoa", "aoa"),
    ("portfolio", "portfolio"), ("cost-risk", "cost-risk")])
def test_open_knows_every_example_by_its_contents(tmp_path, topic, command):
    src = example_path(topic)
    renamed = tmp_path / f"renamed{src.suffix}"   # the name must not matter
    shutil.copy(src, renamed)
    assert opener.plan(renamed).argv[0] == command


def test_open_knows_the_excel_templates(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    for topic, command in (("evm", "evm"), ("jcl", "jcl"), ("aoa", "aoa"),
                           ("portfolio", "portfolio")):
        run(["template", topic, "--out", f"{topic}.xlsx"], capsys)
        assert opener.plan(tmp_path / f"{topic}.xlsx").argv[0] == command


def test_open_knows_an_ipmdar_folder_and_zip(tmp_path):
    from test_ipmdar import cpd_tables, write_folder

    folder = write_folder(cpd_tables(), tmp_path / "cpd")
    assert opener.plan(folder).argv[:2] == ["evm", "--ipmdar"]
    zipped = shutil.make_archive(str(tmp_path / "cpd"), "zip", folder)
    assert opener.plan(zipped).argv[:2] == ["evm", "--ipmdar"]


@pytest.mark.parametrize("name,content,message", [
    ("notes.txt", "hello", "doesn't read .txt files"),
    ("lots.csv", "lot,units,cost\n1,10,100\n", "looks like production lots"),
    ("other.csv", "a,b\n1,2\n", "don't match anything cost-core reads"),
    ("plan.mpp", "x", "save it as XML from Microsoft Project"),
    ("book.xml", "<Workbook/>", "not a Microsoft Project schedule"),
])
def test_open_says_what_it_looked_for(tmp_path, name, content, message):
    path = tmp_path / name
    path.write_text(content, encoding="utf-8")
    with pytest.raises(opener.OpenError, match=message):
        opener.plan(path)


def test_open_runs_the_command_and_puts_results_beside_the_file(tmp_path, capsys):
    data = tmp_path / "my data.csv"
    shutil.copy(example_path("evm"), data)
    out = run(["open", data, "--marking", "CUI"], capsys)
    assert "looks like earned value data" in out
    assert (tmp_path / "my data results" / "report.xlsx").is_file()


# ---------------------------------------------------------------- inflate
def test_factors_are_ratios_of_the_index():
    from cost_core.inflate import factors, parse_basis

    index = {2024: 1.0, 2025: 1.03, 2026: 1.0609}
    by24, ty, by26 = parse_basis("by2024"), parse_basis("ty"), parse_basis("BY2026")
    assert factors(index, [2026], by24, ty)[0] == pytest.approx(1.0609)
    assert factors(index, [2025], ty, by26)[0] == pytest.approx(1.0609 / 1.03)
    assert factors(index, [2025], by24, by26)[0] == pytest.approx(1.0609)
    assert factors(index, [2025], ty, ty)[0] == 1.0


def test_dates_become_us_government_fiscal_years():
    from cost_core.inflate import fiscal_year

    fy = fiscal_year(["2025-09-30", "2025-10-01", "2026-03-15"], start_month=10)
    assert list(fy) == [2025, 2026, 2026]
    assert list(fiscal_year(["FY2027", "2028"])) == [2027, 2028]
    assert list(fiscal_year(["2025-10-01"], start_month=1)) == [2025]


def test_a_converted_table_keeps_what_it_started_with(tmp_path):
    from cost_core.inflate import convert, parse_basis

    frame = pd.DataFrame({"fy": [2025, 2026], "cost": [100.0, 200.0]})
    out = convert(frame, {2024: 1.0, 2025: 1.02, 2026: 1.0404}, parse_basis("by2024"),
                  parse_basis("ty"), start_month=10)
    assert list(out["cost"]) == [100.0, 200.0]
    assert out["cost_then_year"].tolist() == pytest.approx([102.0, 208.08])


@pytest.mark.parametrize("kwargs,message", [
    ({"src": "by2019"}, "covers FY2024 to FY2026 and was asked for FY2019"),
    ({"src": "by 20x6"}, "isn't a dollar basis"),
])
def test_inflate_refuses_what_it_cannot_do(kwargs, message):
    from cost_core.inflate import InflateError, factors, parse_basis

    with pytest.raises(InflateError, match=message):
        factors({2024: 1.0, 2025: 1.02, 2026: 1.04}, [2025], parse_basis(kwargs["src"]),
                parse_basis("ty"))


def test_a_table_of_several_indices_needs_a_name(tmp_path):
    from cost_core.inflate import InflateError, read_index

    path = tmp_path / "idx.csv"
    pd.DataFrame({"index_name": ["RDTE", "RDTE", "PROC", "PROC"],
                  "fiscal_year": [2025, 2026, 2025, 2026],
                  "index_value": [1.0, 1.02, 1.0, 1.03]}).to_csv(path, index=False)
    with pytest.raises(InflateError, match="choose one with --index-name"):
        read_index(path)
    assert read_index(path, "PROC")[0][2026] == 1.03


def test_one_amount_prints_the_arithmetic(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    run(["template", "inflate"], capsys)
    out = run(["inflate", "--index", "my_index.csv", "--amount", 100, "--year", 2030,
               "--from", "by2026", "--to", "ty"], capsys)
    assert f"{100 * 1.02 ** 4:,.4f}" in out and "= 100 x" in out


def test_then_year_needs_a_year_for_one_amount(tmp_path, monkeypatch, capsys, caplog):
    monkeypatch.chdir(tmp_path)
    run(["template", "inflate"], capsys)
    msg = refused(["inflate", "--index", "my_index.csv", "--amount", 100, "--from", "by2026",
                   "--to", "ty"], caplog, capsys)
    assert "give --year" in msg


def test_the_inflate_template_is_marked_as_invented(tmp_path, monkeypatch, capsys):
    monkeypatch.chdir(tmp_path)
    run(["template", "inflate"], capsys)
    assert pd.read_csv("my_index.csv")["index_name"].str.contains("ILLUSTRATIVE").all()


# ------------------------------------------------------------- EVM checks
def messy_evm():
    f = pd.read_csv(example_path("evm"))
    done = f.dropna(subset=["acwp"])
    status, periods = done["period"].max(), sorted(done["period"].unique())
    a1, a2, a3 = f["wbs"].unique()
    f.loc[(f.wbs == a1) & (f.period == periods[4]), "bcwp"] = -5.0
    bac3 = f.loc[f.wbs == a3, "bcws"].sum()
    earned = f.loc[(f.wbs == a3) & (f.period < status), "bcwp"].sum()
    f.loc[(f.wbs == a3) & (f.period == status), "bcwp"] = bac3 - earned
    f.loc[(f.wbs == a2) & (f.period == status), "bcwp"] = 0.0
    return f


def test_the_data_checks_find_what_was_planted():
    from cost_core.evm import EvmData
    from cost_core.evm.checks import data_checks

    checks = data_checks(EvmData.from_frame(messy_evm()))
    assert set(checks["check"]) == {"Earned value went down", "Cost with no earned value",
                                    "Complete but still charging"}
    assert checks.set_index("check").loc["Earned value went down", "period"] == "2026-02"


def test_clean_data_raises_no_checks():
    from cost_core.evm import EvmData
    from cost_core.evm.checks import data_checks

    assert data_checks(EvmData.read(example_path("evm"))).empty


def test_variance_thresholds_in_percent_dollars_or_both():
    from cost_core.evm import EvmData
    from cost_core.evm.checks import variance_breaches

    data = EvmData.read(example_path("evm"))
    # CV -10.6% and -13.1%, SV -31% on the first two; SE is -8.8% and -5.0%.
    assert list(variance_breaches(data)["account"]) == ["1.1 Air vehicle", "1.2 Software"]
    assert variance_breaches(data, cv_pct=50, sv_pct=50).empty
    # With both a percent and a dollar threshold, both have to be broken.
    both = variance_breaches(data, cv_pct=10, sv_pct=None, cv_dollars=200)
    assert list(both["account"]) == ["1.1 Air vehicle"]
    dollars = variance_breaches(data, cv_pct=None, sv_pct=None, cv_dollars=120)
    assert set(dollars["account"]) == {"1.1 Air vehicle", "1.2 Software"}


def test_the_evm_command_writes_and_reports_the_checks(tmp_path, capsys):
    import openpyxl

    path = tmp_path / "messy.csv"
    messy_evm().to_csv(path, index=False)
    out = run(["evm", "--data", path, "--out", tmp_path / "o", "--cv-pct", 11], capsys)
    assert "Data checks (3)" in out and "owe a variance analysis report" in out
    assert len(pd.read_csv(tmp_path / "o" / "data_checks.csv")) == 3
    sheets = openpyxl.load_workbook(tmp_path / "o" / "report.xlsx").sheetnames
    assert {"Data checks", "Variance reports"} <= set(sheets)


# ------------------------------------------------------ reserve allocation
@pytest.fixture(scope="module")
def cost_risk():
    from cost_core.costrisk import analyse, read_workbook
    return analyse(read_workbook(example_path("cost-risk")), n_iter=20000)


def test_the_allocation_adds_up_to_the_p80(cost_risk):
    alloc = cost_risk.allocation(0.8)
    assert alloc["allocated_p80"].sum() == pytest.approx(cost_risk.sim.p80)
    assert alloc["reserve"].sum() == pytest.approx(cost_risk.sim.p80 - cost_risk.point_estimate)
    assert alloc["share_of_reserve"].sum() == pytest.approx(1.0)
    # Each element's own P80 adds to more than the whole's: percentiles don't add.
    assert alloc["own_p80"].sum() > cost_risk.sim.p80
    assert set(alloc["kind"]) == {"element", "risk"}


def test_the_biggest_driver_gets_the_biggest_reserve(cost_risk):
    alloc = cost_risk.allocation(0.8)
    assert alloc.iloc[0]["component"] == cost_risk.drivers().iloc[0]["component"]
    fixed = alloc.set_index("component").loc["8.0 Initial spares"]
    assert fixed["reserve"] == pytest.approx(0.0, abs=1e-9)


def test_the_allocation_level_is_a_fraction(cost_risk):
    from cost_core.costrisk import CostRiskError

    with pytest.raises(CostRiskError, match="between 0 and 1"):
        cost_risk.allocation(80)
    assert cost_risk.allocation(0.5)["allocated_p50"].sum() == pytest.approx(cost_risk.sim.p50)


# ------------------------------------------------ found by the bug sweep
@pytest.mark.parametrize("name,content", [("junk.pptx", b"PK\x03\x04junk"),
                                          ("text.potx", b"hello"), ("book.xlsx", None)])
def test_a_template_that_is_not_powerpoint_is_refused_before_work_starts(tmp_path, name,
                                                                         content, caplog,
                                                                         capsys):
    pytest.importorskip("pptx")
    path = tmp_path / name
    if content is None:
        import openpyxl
        openpyxl.Workbook().save(path)
    else:
        path.write_bytes(content)
    msg = refused(["cost-risk", "--data", example_path("cost-risk"), "--out", tmp_path / "o",
                   "--template", path], caplog, capsys)
    assert "isn't a PowerPoint file (.pptx) or template (.potx)" in msg
    assert not (tmp_path / "o").exists()


def test_a_broken_settings_file_does_not_block_version_about_or_settings(tmp_path, monkeypatch,
                                                                         capsys, caplog):
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(Path, "home", lambda: tmp_path / "nohome")
    (tmp_path / "ce-core.toml").write_text('seed = "x"\n', encoding="utf-8")
    assert "cost-core" in run(["--about"], capsys)
    msg = refused(["settings"], caplog, capsys)
    assert "seed is a whole number" in msg
    msg = refused(["demo", "evm"], caplog, capsys)
    assert "seed is a whole number" in msg


def test_inflate_refuses_an_empty_table(tmp_path, caplog, capsys):
    data = tmp_path / "empty.csv"
    pd.DataFrame({"fy": [], "amount": []}).to_csv(data, index=False)
    msg = refused(["inflate", "--index", example_path("inflate").with_name("inflate_index.csv"),
                   "--data", data, "--from", "by2026", "--to", "ty"], caplog, capsys)
    assert "has no rows to convert" in msg


@pytest.mark.parametrize("name,content,message", [
    ("bad.json", b"{nope", "isn't valid JSON"),
    ("u16.csv", "period,bcws\n1,2\n".encode("utf-16"), "CSV UTF-8"),
])
def test_open_explains_unreadable_text_files(tmp_path, name, content, message):
    path = tmp_path / name
    path.write_bytes(content)
    with pytest.raises(opener.OpenError, match=message):
        opener.plan(path)


def test_excel_csv_utf8_with_a_byte_order_mark_reads(tmp_path):
    path = tmp_path / "bom.csv"
    pd.read_csv(example_path("evm")).to_csv(path, index=False, encoding="utf-8-sig")
    assert opener.plan(path).argv[0] == "evm"


def test_every_demo_runs_with_a_marking_in_the_settings(tmp_path, monkeypatch, capsys):
    # Found by the sweep: the marking went on to inflate, which takes none.
    monkeypatch.chdir(tmp_path)
    monkeypatch.setattr(Path, "home", lambda: tmp_path / "nohome")
    (tmp_path / "ce-core.toml").write_text('marking = "CUI"\n', encoding="utf-8")
    for topic in ("inflate", "cost-risk"):
        run(["demo", topic], capsys)
    written = list((tmp_path / "ce-core-demo" / "inflate").glob("*.csv"))
    assert [p.name for p in written] == ["inflate_example then-year.csv"]
