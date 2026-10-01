# Copyright 2026 Michael Fowler
# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0

"""ce-core phase: an estimate spread over fiscal years and inflated, checked
against hand arithmetic. Every number here is invented."""

import numpy as np
import pandas as pd
import pytest

from cost_core import cli
from cost_core.phasing import (RAYLEIGH_PEAK, PhaseError, phase, profile_weights,
                               read_workbook, write_workbook)


def index_frame(rates, first=2020, last=2045, base=2026):
    return pd.DataFrame([{"index_name": name, "fiscal_year": y,
                          "index_value": (1 + r) ** (y - base)}
                         for name, r in rates.items() for y in range(first, last + 1)])


def workbook(path, lines, settings=(("Base Year", 2026),), index=None):
    cols = ["Line", "Appropriation", "Amount", "Start", "Years", "Profile", "Peak", "Index"]
    frame = pd.DataFrame(lines).reindex(columns=cols)
    with pd.ExcelWriter(path) as xl:
        frame.to_excel(xl, sheet_name="Phasing", index=False)
        pd.DataFrame(list(settings), columns=["Setting", "Value"]).to_excel(
            xl, sheet_name="Settings", index=False)
        if index is not None:
            index.to_excel(xl, sheet_name="Index", index=False)
    return path


def one(**kw):
    return [{"Line": "A", "Appropriation": "RDT&E", "Amount": 100.0, "Start": 2027,
             "Years": 2, "Profile": "uniform", **kw}]


def test_the_template_reads_back(tmp_path):
    inp = read_workbook(write_workbook(tmp_path / "p.xlsx"))
    assert inp.base_year == 2026 and inp.units == "millions" and len(inp.lines) == 5
    assert inp.lines[0].profile == f"rayleigh (peak {RAYLEIGH_PEAK:.2f})"
    assert inp.lines[2].years == 5 and inp.lines[2].weights.tolist() == pytest.approx(
        [0.10, 0.20, 0.25, 0.25, 0.20])


def test_then_year_is_base_year_times_the_index_ratio(tmp_path):
    p = workbook(tmp_path / "p.xlsx", one(), index=index_frame({"Two": 0.02}),
                 settings=[("Base Year", 2026), ("Index", "Two")])
    r = phase(read_workbook(p))
    assert r.long["base_year"].tolist() == [50.0, 50.0]
    assert r.long["then_year"].tolist() == pytest.approx([50 * 1.02, 50 * 1.02 ** 2])
    assert r.totals()["then_year"] == pytest.approx(50 * 1.02 + 50 * 1.0404)


def test_every_line_lands_on_its_amount(tmp_path):
    r = phase(read_workbook(write_workbook(tmp_path / "p.xlsx")))
    by_line = r.long.groupby("line", sort=False)["base_year"].sum()
    assert by_line.tolist() == pytest.approx([120.0, 30.0, 260.0, 18.0, 45.0], abs=1e-12)


@pytest.mark.parametrize("years, peak, top", [(10, 0.45, 5), (10, 0.15, 2), (5, None, 2)])
def test_a_rayleigh_peaks_where_it_says(years, peak, top):
    w = profile_weights("rayleigh", years, peak)
    assert w.sum() == pytest.approx(1.0) and int(np.argmax(w)) + 1 == top
    assert (w > 0).all()


def test_percentages_must_match_the_years_and_are_scaled_with_a_note(tmp_path):
    with pytest.raises(PhaseError, match="3 percentages for 2 years"):
        read_workbook(workbook(tmp_path / "a.xlsx", one(Profile="20;30;50")))
    inp = read_workbook(workbook(tmp_path / "b.xlsx", one(Profile="1;3")))
    assert inp.lines[0].weights.tolist() == [0.25, 0.75]
    assert any("add up to 4" in n for n in inp.notes)


def test_an_end_year_can_stand_in_for_years(tmp_path):
    lines = one(Years=None)
    lines[0]["End"] = 2030
    frame = pd.DataFrame(lines)
    p = tmp_path / "e.xlsx"
    with pd.ExcelWriter(p) as xl:
        frame.to_excel(xl, sheet_name="Phasing", index=False)
        pd.DataFrame([("Base Year", 2026)], columns=["Setting", "Value"]).to_excel(
            xl, sheet_name="Settings", index=False)
    assert read_workbook(p).lines[0].years == 4
    frame["Years"] = 2
    with pd.ExcelWriter(p) as xl:
        frame.to_excel(xl, sheet_name="Phasing", index=False)
        pd.DataFrame([("Base Year", 2026)], columns=["Setting", "Value"]).to_excel(
            xl, sheet_name="Settings", index=False)
    with pytest.raises(PhaseError, match="ends in FY2028, not FY2030"):
        read_workbook(p)


def test_each_line_can_use_its_own_index(tmp_path):
    lines = one(Index="Proc") + [{"Line": "B", "Appropriation": "O&M", "Amount": 100.0,
                                  "Start": 2027, "Years": 1, "Profile": "uniform"}]
    p = workbook(tmp_path / "p.xlsx", lines,
                 index=index_frame({"Proc": 0.02, "OM": 0.03}),
                 settings=[("Base Year", 2026), ("Index", "OM")])
    r = phase(read_workbook(p))
    assert r.long["index"].tolist() == ["Proc", "Proc", "OM"]
    assert r.long["then_year"].iloc[2] == pytest.approx(103.0)


@pytest.mark.parametrize("lines, settings, message", [
    (one(Index="Nope"), [("Base Year", 2026), ("Index", "Proc")], "there's no index 'Nope'"),
    (one(), [("Base Year", 2026)], "holds 2 indices"),
    (one(Start=2044, Years=4), [("Base Year", 2026), ("Index", "Proc")],
     "Extend the table"),
])
def test_index_mistakes_are_refused(tmp_path, lines, settings, message):
    p = workbook(tmp_path / "p.xlsx", lines, settings=settings,
                 index=index_frame({"Proc": 0.02, "OM": 0.03}))
    with pytest.raises(PhaseError, match=message):
        phase(read_workbook(p))


def test_a_constant_rate_or_no_index_at_all(tmp_path):
    p = workbook(tmp_path / "r.xlsx", one(),
                 settings=[("Base Year", "BY2026"), ("Inflation Rate", "2.5%")])
    r = phase(read_workbook(p))
    assert r.long["then_year"].tolist() == pytest.approx([51.25, 50 * 1.025 ** 2])
    bare = phase(read_workbook(workbook(tmp_path / "b.xlsx", one())))
    assert not bare.has_then_year and bare.long["then_year"].isna().all()
    assert any("base-year dollars only" in n for n in bare.inputs.notes)


def test_a_base_year_is_required(tmp_path):
    with pytest.raises(PhaseError, match="Base Year on the Settings sheet"):
        read_workbook(workbook(tmp_path / "p.xlsx", one(), settings=[("Units", "$M")]))


def test_the_command_writes_every_output(tmp_path, capsys):
    from openpyxl import load_workbook

    p = write_workbook(tmp_path / "p.xlsx")
    cli.main(["phase", "--data", str(p), "--out", str(tmp_path / "o"), "--marking", "CUI"])
    out = capsys.readouterr().out
    assert "inflation adds" in out and "invented illustrative" in out
    for f in ("report.xlsx", "brief.pptx", "phasing_long.csv", "base_year.csv",
              "then_year.csv", "by_appropriation.csv", "assumptions.json", "phasing.png"):
        assert (tmp_path / "o" / f).exists(), f
    wb = load_workbook(tmp_path / "o" / "report.xlsx")
    assert {"Then-year by FY", "Base-year by FY", "By appropriation", "Every line-year",
            "Index used", "Assumptions"} <= set(wb.sheetnames)
    assert wb["Summary"].oddHeader.center.text == "CUI"


def test_an_index_file_and_base_year_flag_override_the_workbook(tmp_path, capsys):
    idx = tmp_path / "idx.csv"
    index_frame({"Published": 0.03}, base=2025).to_csv(idx, index=False)
    p = workbook(tmp_path / "p.xlsx", one(), settings=[("Base Year", 2026)])
    cli.main(["phase", "--data", str(p), "--out", str(tmp_path / "o"), "--index", str(idx),
              "--base-year", "2025"])
    capsys.readouterr()
    long = pd.read_csv(tmp_path / "o" / "phasing_long.csv")
    assert long["then_year"].tolist() == pytest.approx([50 * 1.03 ** 2, 50 * 1.03 ** 3])


def test_demo_template_and_open_know_the_command(tmp_path, capsys):
    cli.main(["demo", "phase", "--out", str(tmp_path / "demo")])
    assert (tmp_path / "demo" / "report.xlsx").exists()
    cli.main(["template", "phase", "--out", str(tmp_path / "mine.xlsx")])
    assert "ce-core phase --data" in capsys.readouterr().out
    from cost_core import opener

    assert opener.plan(tmp_path / "mine.xlsx").argv[0] == "phase"
