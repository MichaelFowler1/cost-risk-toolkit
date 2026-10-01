# Copyright 2026 Michael Fowler
# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0

"""Templates show what to fill in: yellow input cells, required headings
dark blue and optional ones light blue, a note on every heading, drop-downs
where the answers are fixed, and a colour key. And the styling changes
nothing the commands read."""

import pandas as pd
import pytest
from openpyxl import load_workbook

from cost_core import cli, templates
from cost_core.template_style import (BLANK_ROWS, INPUT, OPTIONAL, REQUIRED, SETTINGS,
                                      TABLES, kinds)

KINDS = list(kinds())


def _rgb(cell):
    return (cell.fill.fgColor.rgb or "")[-6:]


@pytest.mark.parametrize("kind", KINDS)
def test_headings_say_required_or_optional_and_inputs_are_yellow(tmp_path, kind):
    path = templates.write(kind, tmp_path / cli.TEMPLATE_FILES[kind])
    wb = load_workbook(path)
    for sheet, table in TABLES.get(kind, {}).items():
        if sheet not in wb.sheetnames:
            continue
        ws = wb[sheet]
        for cell in ws[1]:
            if cell.value is None:
                continue
            spec = table.get(str(cell.value).strip()) or table.get("*") or (False,)
            assert _rgb(cell) == (REQUIRED if spec[0] else OPTIONAL), (sheet, cell.value)
            assert cell.comment is not None and cell.comment.text.startswith(
                "Required." if spec[0] else "Optional.")
            below = ws.cell(row=ws.max_row, column=cell.column)
            assert _rgb(ws.cell(row=2, column=cell.column)) == INPUT
            assert _rgb(below) == INPUT
        assert ws.max_row >= BLANK_ROWS + 2


@pytest.mark.parametrize("kind", KINDS)
def test_the_key_and_the_settings_are_marked(tmp_path, kind):
    path = templates.write(kind, tmp_path / cli.TEMPLATE_FILES[kind])
    wb = load_workbook(path)
    if "Instructions" in wb.sheetnames:
        assert wb["Instructions"]["A1"].value == "How to fill in this workbook"
    if kind in SETTINGS and "Settings" in wb.sheetnames:
        ws = wb["Settings"]
        named = [r for r in range(2, ws.max_row + 1) if ws[f"A{r}"].value is not None]
        assert named and all(_rgb(ws[f"B{r}"]) == INPUT for r in named)


def test_fixed_answers_get_a_drop_down(tmp_path):
    wb = load_workbook(templates.write("cost-risk", tmp_path / "e.xlsx"))
    ws = wb["Elements"]
    lists = [dv.formula1 for dv in ws.data_validations.dataValidation]
    assert any("triangular" in f for f in lists)
    strict = [dv for dv in ws.data_validations.dataValidation if "triangular" in dv.formula1]
    assert strict[0].showErrorMessage


@pytest.mark.parametrize("kind", ["cost-risk", "phase", "evm", "cer"])
def test_styling_changes_nothing_the_command_reads(tmp_path, kind):
    styled = templates.write(kind, tmp_path / "styled" / cli.TEMPLATE_FILES[kind])
    plain = templates._write(kind, tmp_path / "plain" / cli.TEMPLATE_FILES[kind])
    a = pd.read_excel(styled, sheet_name=None)
    b = pd.read_excel(plain, sheet_name=None)
    for name in b:
        if name == "Instructions":
            continue
        pd.testing.assert_frame_equal(a[name].dropna(how="all"), b[name].dropna(how="all"))


def test_a_csv_template_is_left_alone(tmp_path):
    path = templates.write("inflate", tmp_path / "idx.csv")
    assert path.suffix == ".csv" and len(pd.read_csv(path)) > 0
