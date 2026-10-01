# Copyright 2026 Michael Fowler
# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0

"""ce-core cer: a CER study read from Excel, checked against hand arithmetic.

Every number here is invented.
"""

import numpy as np
import pandas as pd
import pytest
from scipy import stats

from cost_core import cli
from cost_core.cer.study import CerError, analyse, read_workbook, write_workbook


def workbook(path, data, estimate=None, settings=None):
    with pd.ExcelWriter(path) as xl:
        pd.DataFrame(data).to_excel(xl, sheet_name="Data", index=False)
        if estimate is not None:
            pd.DataFrame(estimate).to_excel(xl, sheet_name="Estimate", index=False)
        if settings is not None:
            pd.DataFrame(settings, columns=["Setting", "Value"]).to_excel(
                xl, sheet_name="Settings", index=False)
    return path


def line_data(n=10, seed=3):
    rng = np.random.default_rng(seed)
    x = np.linspace(10, 100, n)
    return {"Program": [f"P{i}" for i in range(n)], "Cost": 5 + 2 * x + rng.normal(0, 8, n),
            "Weight": x}


def test_the_template_reads_back_with_its_settings(tmp_path):
    inp = read_workbook(write_workbook(tmp_path / "cer.xlsx"))
    assert inp.response == "Cost" and inp.predictors == ("Weight (lb)", "Power (kW)")
    assert (inp.form.value, inp.method, inp.level) == ("log_log", "mupe", 0.8)
    assert inp.units == "millions" and inp.dollar_year == "BY2026"
    assert len(inp.used) == 12 and list(inp.estimate["name"]) == ["Next radar", "Large radar"]


def test_log_log_ols_matches_least_squares_on_the_logs(tmp_path):
    inp = read_workbook(write_workbook(tmp_path / "cer.xlsx"))
    study = analyse(inp, method="ols")
    d = inp.used
    design = np.column_stack([np.ones(len(d)), np.log(d["Weight (lb)"]), np.log(d["Power (kW)"])])
    beta, *_ = np.linalg.lstsq(design, np.log(d["Cost"]), rcond=None)
    got = study.coefficients()["estimate"].to_numpy()
    assert got == pytest.approx(beta, rel=1e-9)


def test_the_prediction_interval_is_the_textbook_one(tmp_path):
    data = line_data()
    p = workbook(tmp_path / "line.xlsx", data, estimate={"Program": ["New"], "Weight": [64.0]},
                 settings=[("Form", "linear"), ("Method", "OLS"), ("Confidence", 0.9)])
    est = analyse(read_workbook(p)).estimates.iloc[0]
    x, y = np.asarray(data["Weight"]), np.asarray(data["Cost"])
    n = len(x)
    b1, b0 = np.polyfit(x, y, 1)
    s = np.sqrt(np.sum((y - (b0 + b1 * x)) ** 2) / (n - 2))
    half = stats.t.ppf(0.95, n - 2) * s * np.sqrt(1 + 1 / n + (64 - x.mean()) ** 2
                                                   / np.sum((x - x.mean()) ** 2))
    assert est["estimate"] == pytest.approx(b0 + b1 * 64, rel=1e-9)
    assert (est["lower"], est["upper"]) == pytest.approx((b0 + b1 * 64 - half,
                                                          b0 + b1 * 64 + half), rel=1e-9)


def test_rows_left_out_are_named_with_the_reason(tmp_path):
    data = line_data()
    data["Use"] = ["yes"] * 8 + ["no", 1.0]
    data["Note"] = [""] * 8 + ["data suspect", ""]
    data["Weight"] = list(data["Weight"][:-1]) + [None]
    inp = read_workbook(workbook(tmp_path / "d.xlsx", data))
    assert len(inp.used) == 8
    assert list(inp.excluded["why"]) == ["data suspect", "no Weight given"]
    assert "2 of 10 rows are left out" in inp.notes[0]


def test_a_year_column_is_not_taken_as_a_driver(tmp_path):
    data = {**line_data(), "FY": list(range(2010, 2020))}
    inp = read_workbook(workbook(tmp_path / "d.xlsx", data))
    assert inp.predictors == ("Weight",)
    assert any("'FY' is not used as a driver" in n for n in inp.notes)
    named = read_workbook(workbook(tmp_path / "e.xlsx", data,
                                   settings=[("Drivers", "Weight; FY"), ("Form", "linear")]))
    assert named.predictors == ("Weight", "FY")


@pytest.mark.parametrize("change, message", [
    (lambda d: d.__setitem__("Cost", ["lots"] + list(d["Cost"][1:])), "Data sheet, row 2"),
    (lambda d: d.__setitem__("Weight", [-1.0] + list(d["Weight"][1:])),
     "log-log CER needs every value above zero"),
    (lambda d: d.__setitem__("Use", ["maybe"] + ["yes"] * 9), "Use is yes or no"),
])
def test_bad_data_is_refused_with_its_row(tmp_path, change, message):
    data = line_data()
    change(data)
    with pytest.raises(CerError, match=message):
        read_workbook(workbook(tmp_path / "d.xlsx", data))


def test_too_few_programs_and_a_missing_estimate_driver_are_refused(tmp_path):
    small = {"Program": ["A", "B"], "Cost": [10.0, 20.0], "Weight": [1.0, 2.0]}
    with pytest.raises(CerError, match="2 programs can't fit 2 parameters"):
        analyse(read_workbook(workbook(tmp_path / "s.xlsx", small)))
    with pytest.raises(CerError, match="Estimate sheet has no column for 'Weight'"):
        read_workbook(workbook(tmp_path / "e.xlsx", line_data(),
                               estimate={"Program": ["New"], "Mass": [5.0]}))


def test_extrapolation_and_a_weak_driver_are_said(tmp_path):
    from cost_core import plain

    study = analyse(read_workbook(write_workbook(tmp_path / "cer.xlsx")))
    assert list(study.estimates["outside_data"]) == ["", "yes"]
    rng = np.random.default_rng(8)
    data = line_data(seed=8)
    data["Colour"] = rng.uniform(1, 9, 10)
    weak = analyse(read_workbook(workbook(tmp_path / "w.xlsx", data,
                                          settings=[("Form", "linear")])))
    assert "Colour's effect can't be told apart from zero" in " ".join(plain.cer(weak, ""))


def test_a_csv_of_the_data_alone_fits(tmp_path):
    p = tmp_path / "d.csv"
    pd.DataFrame(line_data()).to_csv(p, index=False)
    study = analyse(read_workbook(p, drivers=["Weight"]), form="linear")
    assert study.cer.predictors == ("Weight",) and study.estimates.empty


def test_the_command_writes_every_output(tmp_path, capsys):
    from openpyxl import load_workbook

    p = write_workbook(tmp_path / "cer.xlsx")
    cli.main(["cer", "--data", str(p), "--out", str(tmp_path / "o"), "--marking", "CUI"])
    out = capsys.readouterr().out
    assert "Next radar" in out and "outside the data" in out
    for f in ("report.xlsx", "brief.pptx", "methods.csv", "coefficients.csv", "estimates.csv",
              "diagnostics.csv", "data.csv", "assumptions.json", "cer_fit.png"):
        assert (tmp_path / "o" / f).exists(), f
    wb = load_workbook(tmp_path / "o" / "report.xlsx")
    assert {"Estimates", "Methods compared", "Coefficients", "Diagnostics", "Data",
            "Log-space bias", "Assumptions"} <= set(wb.sheetnames)
    assert wb["Summary"].oddHeader.center.text == "CUI"


def test_flags_override_the_workbook(tmp_path, capsys):
    import json

    p = write_workbook(tmp_path / "cer.xlsx")
    cli.main(["cer", "--data", str(p), "--out", str(tmp_path / "o"), "--method", "ols",
              "--confidence", "0.9", "--drivers", "Weight (lb)"])
    capsys.readouterr()
    record = json.loads((tmp_path / "o" / "assumptions.json").read_text())
    assert record["method"] == "OLS" and record["interval"] == "90% prediction interval"
    assert record["drivers"] == "Weight (lb)"


def test_demo_template_and_open_know_the_command(tmp_path, capsys, caplog):
    cli.main(["demo", "cer", "--out", str(tmp_path / "demo")])
    assert (tmp_path / "demo" / "report.xlsx").exists()
    cli.main(["template", "cer", "--out", str(tmp_path / "mine.xlsx")])
    assert "ce-core cer --data" in capsys.readouterr().out
    with pytest.raises(SystemExit):
        cli.main(["template", "cer", "--out", str(tmp_path / "mine.csv")])
    from cost_core import opener

    assert opener.plan(tmp_path / "mine.xlsx").argv[0] == "cer"


def test_the_chart_draws_one_driver_with_its_band(tmp_path):
    pytest.importorskip("matplotlib")
    from cost_core.reporting.charts import plot_cer_diagnostics

    study = analyse(read_workbook(workbook(tmp_path / "l.xlsx", line_data(),
                                           settings=[("Form", "linear")])))
    assert plot_cer_diagnostics(study.cer, tmp_path / "c.png", level=0.9,
                                units="$K").exists()


def test_a_text_cell_stops_the_command_with_its_row(tmp_path, capsys, caplog):
    data = line_data()
    data["Cost"] = ["lots"] + list(data["Cost"][1:])
    p = workbook(tmp_path / "d.xlsx", data)
    with pytest.raises(SystemExit):
        cli.main(["cer", "--data", str(p), "--out", str(tmp_path / "o")])
    assert "Data sheet, row 2: the Cost 'lots' is not a number" in (
        caplog.text + capsys.readouterr().err)


# ------------------------------------------------- into cost risk, lognormal
def risk_workbook(path, elements, settings=(("Iterations", 100000), ("Seed", 3))):
    from cost_core import costrisk

    return costrisk.write_workbook(
        path, elements=pd.DataFrame(elements),
        risks=pd.DataFrame(columns=["Risk", "Probability", "Low", "Most Likely", "High",
                                    "Element"]),
        pairs=pd.DataFrame(columns=["Element A", "Element B", "Correlation"]),
        settings=pd.DataFrame(list(settings), columns=["Setting", "Value"]))


@pytest.mark.parametrize("share, setting", [(0.8, None), (0.9, "90%")])
def test_a_lognormal_element_puts_low_and_high_at_its_range(tmp_path, share, setting):
    from cost_core import costrisk

    settings = [("Iterations", 100000), ("Seed", 3)] + (
        [("Lognormal Range", setting)] if setting else [])
    p = risk_workbook(tmp_path / "r.xlsx", {
        "Element": ["Radar", "Software"], "Point Estimate": [100.0, 40.0],
        "Low": [80.0, 35.0], "Most Likely": [None, 40.0], "High": [125.0, 60.0],
        "Distribution": ["lognormal", ""]}, settings)
    r = costrisk.analyse(costrisk.read_workbook(p))
    tail = (1 - share) / 2 * 100
    got = np.percentile(r.sim.element_samples[:, 0], [tail, 50, 100 - tail])
    assert got == pytest.approx([80.0, 100.0, 125.0], rel=0.01)
    assert r.assumptions["lognormal range"] == f"{share:.0%} between Low and High"


def test_a_lognormal_needs_a_positive_range(tmp_path):
    from cost_core import costrisk

    for low, high, message in ((None, 120.0, "needs both a Low and a High"),
                               (0.0, 120.0, "0 < Low < High")):
        p = risk_workbook(tmp_path / "r.xlsx", {
            "Element": ["Radar"], "Point Estimate": [100.0], "Low": [low],
            "Most Likely": [None], "High": [high], "Distribution": ["lognormal"]})
        with pytest.raises(costrisk.CostRiskError, match=message):
            costrisk.read_workbook(p)


def test_an_off_centre_point_estimate_is_noted(tmp_path):
    from cost_core import costrisk

    p = risk_workbook(tmp_path / "r.xlsx", {
        "Element": ["Radar"], "Point Estimate": [110.0], "Low": [80.0],
        "Most Likely": [None], "High": [125.0], "Distribution": ["lognormal"]})
    assert any("centred on 100" in n for n in costrisk.read_workbook(p).notes)


def test_cer_estimates_carry_into_cost_risk(tmp_path, capsys):
    from openpyxl import load_workbook
    from cost_core import costrisk

    cli.main(["cer", "--data", str(write_workbook(tmp_path / "cer.xlsx")), "--out",
              str(tmp_path / "o")])
    capsys.readouterr()
    rows = pd.read_csv(tmp_path / "o" / "cost_risk_rows.csv")
    sheet = load_workbook(tmp_path / "o" / "report.xlsx")["For cost risk"]
    assert "default Lognormal Range" in sheet["A1"].value
    est = pd.read_csv(tmp_path / "o" / "estimates.csv").iloc[0]
    r = costrisk.analyse(costrisk.read_workbook(risk_workbook(tmp_path / "r.xlsx",
                                                              rows.iloc[:1])))
    got = np.percentile(r.sim.element_samples[:, 0], [10, 50, 90])
    assert got == pytest.approx([est["lower"], est["estimate"], est["upper"]], rel=0.01)
    cli.main(["cer", "--data", str(tmp_path / "cer.xlsx"), "--out", str(tmp_path / "p"),
              "--confidence", "0.9"])
    capsys.readouterr()
    note = load_workbook(tmp_path / "p" / "report.xlsx")["For cost risk"]["A1"].value
    assert "set Lognormal Range to 0.9" in note


@pytest.mark.parametrize("form", ["log_log", "linear"])
def test_mupe_matches_a_statsmodels_gamma_glm(form):
    """MUPE solves the same estimating equations as a Gamma GLM with Pearson
    scale, so statsmodels is an independent check on the numbers."""
    sm = pytest.importorskip("statsmodels.api")
    from cost_core.cer import fit_cer
    from cost_core.cer.study import EXAMPLE_DATA

    d = EXAMPLE_DATA.rename(columns={"Weight (lb)": "W", "Power (kW)": "P"})
    cer = fit_cer(d, "Cost", ["W", "P"], form=form, method="mupe")
    if form == "log_log":
        X, link = np.column_stack([np.ones(len(d)), np.log(d.W), np.log(d.P)]), \
            sm.families.links.Log()
    else:
        X, link = np.column_stack([np.ones(len(d)), d.W, d.P]), sm.families.links.Identity()
    glm = sm.GLM(d["Cost"], X, family=sm.families.Gamma(link=link)).fit(
        scale="X2", tol=1e-14, maxiter=1000)
    # Both are iterative and agree to MUPE's own stopping rule: 7e-7 relative
    # on this machine, 1.0e-6 on a GitHub Linux runner. 1e-5 is ten times that,
    # still five significant figures, and nowhere near a different answer.
    assert cer.result.theta == pytest.approx(glm.params.to_numpy(), rel=1e-5, abs=1e-6)
    assert np.sqrt(np.diag(cer.result.cov)) == pytest.approx(glm.bse.to_numpy(), rel=1e-5)
    assert cer.result.sigma == pytest.approx(np.sqrt(glm.scale), rel=1e-5)


# ------------------------------------------------- 2.7.0 pre-release sweep
def test_a_driver_with_one_value_is_refused(tmp_path):
    data = {**line_data(), "Weight": [200.0] * 10}
    with pytest.raises(CerError, match="same value for every program"):
        analyse(read_workbook(workbook(tmp_path / "d.xlsx", data)))


def test_drivers_in_lockstep_are_refused_and_near_lockstep_is_noted(tmp_path):
    data = line_data()
    with pytest.raises(CerError, match="'Weight' and 'Double' move in lockstep on the log"):
        analyse(read_workbook(workbook(tmp_path / "a.xlsx",
                                       {**data, "Double": [2 * w for w in data["Weight"]]})))
    rng = np.random.default_rng(4)
    near = [w * (1 + rng.normal(0, 0.01)) for w in data["Weight"]]
    s = analyse(read_workbook(workbook(tmp_path / "b.xlsx", {**data, "Near": near})))
    assert any("variance inflation factor" in n for n in s.notes)


def test_a_perfect_fit_is_called_out(tmp_path):
    from cost_core import plain

    data = {**line_data(), "Cost": [w * 0.1 for w in line_data()["Weight"]]}
    s = analyse(read_workbook(workbook(tmp_path / "d.xlsx", data)))
    assert "Real cost data never do" in plain.cer(s, "")[0]


def test_risk_distributions_are_said_to_be_triangular(tmp_path):
    from cost_core import costrisk

    p = tmp_path / "r.xlsx"
    with pd.ExcelWriter(p) as xl:
        pd.DataFrame({"Element": ["A"], "Point Estimate": [100.0], "Low": [90.0],
                      "Most Likely": [100.0], "High": [130.0]}).to_excel(
            xl, sheet_name="Elements", index=False)
        pd.DataFrame({"Risk": ["R"], "Probability": [0.5], "Low": [1], "Most Likely": [2],
                      "High": [3], "Distribution": ["lognormal"]}).to_excel(
            xl, sheet_name="Risks", index=False)
    assert any("Distribution column isn't read" in n for n in costrisk.read_workbook(p).notes)
