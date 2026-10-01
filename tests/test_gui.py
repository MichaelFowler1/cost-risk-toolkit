# Copyright 2026 Michael Fowler
# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0

"""The window's decisions, tested without a screen: the command each job
runs, where results go, what the output said, the key numbers, where a
problem is. Every number here is invented."""

import os

import pytest

from cost_core import cli, gui


@pytest.fixture
def ran(tmp_path_factory):
    """Run a demo once and keep its results folder."""
    made = {}

    def run(topic):
        if topic not in made:
            out = tmp_path_factory.mktemp(topic) / "out"
            cli.main(["demo", topic, "--out", str(out)])
            made[topic] = out
        return made[topic]

    return run


# ------------------------------------------------------------- commands
def test_every_job_runs_the_command_line_with_the_file(tmp_path):
    book = tmp_path / "my estimate.xlsx"
    argv = gui.run_argv(gui.task("cost-risk"), book, "millions", "CUI")
    assert argv == ["cost-risk", "--data", str(book), "--out",
                    str(tmp_path / "my estimate results"), "--units", "millions",
                    "--marking", "CUI"]
    assert "--units" not in gui.run_argv(gui.task("cost-risk"), book, "as entered")
    assert "--units" not in gui.run_argv(gui.task("jcl"), book, "millions")
    assert gui.run_argv(gui.task("open"), book)[:2] == ["open", str(book)]
    assert gui.run_argv(gui.task("schedule"), book)[:2] == ["schedule-check", "--mspdi"]


def test_inflate_and_aoa_pass_their_extra_choices(tmp_path):
    book = tmp_path / "amounts.csv"
    argv = gui.run_argv(gui.task("inflate"), book, marking="CUI",
                        extras={"index": "idx.csv", "from": "by2025", "to": "ty"})
    assert argv[argv.index("--index") + 1] == "idx.csv"
    assert argv[argv.index("--from") + 1] == "by2025" and "--marking" not in argv
    aoa = gui.run_argv(gui.task("aoa"), book, extras={"status_quo": "Keep it"})
    assert aoa[aoa.index("--status-quo") + 1] == "Keep it"
    assert "--status-quo" not in gui.run_argv(gui.task("aoa"), book, extras={})


@pytest.mark.parametrize("key", [t.key for t in gui.TASKS])
def test_every_job_s_arguments_are_ones_the_command_takes(tmp_path, capsys, key):
    t = gui.task(key)
    extras = {"index": str(tmp_path / "i.csv"), "status_quo": "Keep"}
    argv = gui.run_argv(t, tmp_path / "missing.xlsx", "millions", "CUI", extras=extras)
    with pytest.raises(SystemExit) as stop:
        cli.main(argv)
    err = capsys.readouterr().err
    assert "unrecognized arguments" not in err and "invalid choice" not in err
    assert stop.value.code != 2, err


def test_each_job_with_an_example_or_template_has_one():
    for t in gui.TASKS:
        if t.demo:
            assert t.demo in cli.DEMOS
        if t.template:
            assert t.template in cli.TEMPLATE_FILES


# ---------------------------------------------------------------- places
def test_results_go_beside_the_workbook_and_never_clash(tmp_path):
    out = gui.results_dir(tmp_path / "est.xlsx")
    assert out == tmp_path / "est results"
    assert gui.next_free(out) == tmp_path / "est results 2"
    (tmp_path / "est results 2").mkdir()
    assert gui.next_free(out) == tmp_path / "est results 3"


def test_a_report_open_in_excel_is_seen(tmp_path):
    assert gui.locked(tmp_path / "nothing") == []
    (tmp_path / "report.xlsx").write_bytes(b"x")
    (tmp_path / "brief.pptx").write_bytes(b"x")
    assert gui.locked(tmp_path) == []
    (tmp_path / "~$report.xlsx").write_bytes(b"x")
    assert gui.locked(tmp_path) == ["report.xlsx"]


def test_preferences_survive_and_a_broken_file_is_forgotten(tmp_path):
    path = tmp_path / "gui.json"
    assert gui.load_prefs(path) == {}
    prefs = {"units": "millions"}
    for f in ("a.xlsx", "b.xlsx", "a.xlsx"):
        gui.add_recent(prefs, "cer", f)
    assert prefs["recent"]["cer"] == ["a.xlsx", "b.xlsx"]
    for i in range(10):
        gui.add_recent(prefs, "cer", f"{i}.xlsx")
    assert len(prefs["recent"]["cer"]) == 6
    gui.save_prefs(prefs, path)
    assert gui.load_prefs(path) == prefs
    path.write_text("{not json", encoding="utf-8")
    assert gui.load_prefs(path) == {}


# --------------------------------------------------------------- reading
OUTPUT = """Cost risk: 9 elements, 10,000 draws
What this means:
  - The estimate of 156 is exceeded 98% of the time.
  - To be 80% sure, fund 214,
    a reserve of 58.
note: 1 element has no range.

Wrote report.xlsx
"""


def test_the_plain_words_and_notes_are_read_from_the_output():
    assert gui.summary(OUTPUT) == ["The estimate of 156 is exceeded 98% of the time.",
                                   "To be 80% sure, fund 214, a reserve of 58."]
    assert gui.notes(OUTPUT) == ["1 element has no range."]


def test_errors_lose_the_prefix_and_point_at_the_buttons():
    out = ("[ERROR] Cost risk failed: Elements sheet, row 3: the most likely 'x' is not a "
           "number.\n  To see this command work first: ce-core demo cost-risk\n")
    found = gui.errors(out)
    assert found[0].startswith("Elements sheet, row 3")
    assert "press 'See an example'" in found[0] and "ce-core" not in found[0]
    assert "open in Excel" in gui.errors("PermissionError: [Errno 13] Permission denied")[0]
    assert gui.errors("Traceback\nboom\n") == ["Traceback", "boom"]


@pytest.mark.parametrize("topic, first", [
    ("cost-risk", "Point estimate"), ("cer", None), ("phase", "Base-year total"),
    ("evm", "CPI"), ("jcl", "Plan's joint confidence"), ("aoa", "Cheapest at P50"),
    ("portfolio", "Programs funded"), ("schedule", "Checks passed"), ("inflate", "Before")])
def test_each_result_gives_its_key_numbers(ran, topic, first):
    out = ran(topic)
    assert gui.result_kind(out) == topic
    tiles, chart = gui.answer(out)
    assert len(tiles) >= 2 and all(isinstance(v, str) and v for _, v in tiles)
    if first:
        assert tiles[0][0] == first
    assert (chart is not None) == (gui.SIGNATURES.get(topic, (None,))[0] is not None)
    found = gui.written(out)
    assert "folder" in found and "report" in found


def test_a_folder_with_nothing_recognisable_gives_no_tiles(tmp_path):
    assert gui.answer(tmp_path) == ([], None)
    assert gui.answer(tmp_path / "missing") == ([], None)


# ------------------------------------------------------ finding problems
def test_a_problem_names_its_place():
    assert gui.problem_location("Elements sheet, row 7 (Radar): Low is above High.") == \
        ("Elements", 7, "Elements sheet, row 7 (Radar): Low is above High.")
    assert gui.problem_location("No Settings sheet.") is None


def _broken_estimate(tmp_path):
    from openpyxl import load_workbook

    path = gui.make_template(gui.task("cost-risk"), tmp_path / "est.xlsx")
    wb = load_workbook(path)
    ws = wb["Elements"]
    heads = [c.value for c in ws[1]]
    ws.cell(row=3, column=heads.index("Most Likely") + 1).value = "about 40"
    wb.save(path)
    return path, heads.index("Most Likely") + 1


def test_check_finds_text_in_a_number_once_and_marks_it(tmp_path):
    from openpyxl import load_workbook

    path, col = _broken_estimate(tmp_path)
    before = path.read_bytes()
    found = gui.check_workbook(gui.task("cost-risk"), path)
    assert len(found["problems"]) == 1 and "row 3" in found["problems"][0]
    marked = gui.mark_problem(path, found["problems"][0])
    assert path.read_bytes() == before
    assert marked.name == "est - problem marked.xlsx"
    ws = load_workbook(marked)["Elements"]
    cell = ws.cell(row=3, column=col)
    assert cell.comment is not None and "about 40" in cell.comment.text
    assert cell.border.left.color.rgb.endswith("C00000")
    assert ws.cell(row=4, column=col).comment is None


def test_a_message_naming_no_column_marks_the_row(tmp_path):
    from openpyxl import load_workbook

    path, _ = _broken_estimate(tmp_path)
    marked = gui.mark_problem(path, "Elements sheet, row 4: something is off.")
    ws = load_workbook(marked)["Elements"]
    filled = [c for c in ws[4] if c.value is not None or c.column <= 6]
    assert all(c.fill.fgColor.rgb.endswith("F8CBAD") for c in filled[:6])
    assert gui.mark_problem(path, "No sheet named.") is None
    assert gui.mark_problem(path, "Nowhere sheet, row 2: x") is None


def test_check_finds_a_blank_required_cell(tmp_path):
    from openpyxl import load_workbook

    path = gui.make_template(gui.task("cer"), tmp_path / "cer.xlsx")
    wb = load_workbook(path)
    ws = wb["Data"]
    cost = [c.value for c in ws[1]].index("Cost") + 1
    ws.cell(row=2, column=cost).value = None
    wb.save(path)
    problems = gui.check_workbook(gui.task("cer"), path)["problems"]
    assert any("row 2: Cost is blank" in p for p in problems)


@pytest.mark.parametrize("key", [t.key for t in gui.TASKS if t.template])
def test_every_fresh_template_checks_clean(tmp_path, key):
    from cost_core.templates import companion_path

    t = gui.task(key)
    path = gui.make_template(t, tmp_path / gui.template_path(t).name)
    if key == "inflate":
        path = companion_path(path)
    assert gui.check_workbook(t, path)["problems"] == []


def test_the_status_quo_choices_come_from_the_workbook(tmp_path):
    path = gui.make_template(gui.task("aoa"), tmp_path / "aoa.xlsx")
    names = gui.alternatives(path)
    assert len(names) >= 2 and all(isinstance(n, str) and n for n in names)
    assert gui.alternatives(tmp_path / "missing.xlsx") == []


def test_the_glossary_covers_the_words_the_window_uses():
    terms = " ".join(t for t, _ in gui.GLOSSARY)
    for word in ("P50", "CER", "SIR", "CPI", "JCL", "Status quo", "Appropriation"):
        assert word in terms


# ---------------------------------------------------------------- window
def test_the_window_lays_out_every_job(tmp_path, ran, monkeypatch):
    tk = pytest.importorskip("tkinter")
    monkeypatch.setenv("APPDATA", str(tmp_path))
    try:
        root = tk.Tk()
    except tk.TclError:
        pytest.skip("no display")
    try:
        app = gui.App(root)
        for t in gui.TASKS:
            app.choose(t.key)
            root.update()
        app.choose("cost-risk")
        app.finish(0, OUTPUT, ran("cost-risk"), "", None)
        root.update()
        assert len(app.tiles.winfo_children()) == 4
        assert {"report", "brief", "folder"} <= set(app.results)
        assert "fund 214" in app.text.get("1.0", "end")
        app.finish(1, "[ERROR] boom", tmp_path / "x", "", None)
        assert app.results == {} and "boom" in app.text.get("1.0", "end")
        app.glossary()
        app.settings()
        root.update()
    finally:
        root.destroy()
    assert not (tmp_path / "cost-core" / "gui.json").exists() or \
        os.path.getsize(tmp_path / "cost-core" / "gui.json") > 0
