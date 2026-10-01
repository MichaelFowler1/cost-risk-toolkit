# Copyright 2026 Michael Fowler
# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0

"""Economic analysis on an AoA: each alternative against the status quo,
checked by hand. Every number here is invented."""

import json

import numpy as np
import pandas as pd
import pytest

from cost_core import cli
from cost_core.aoa import Alternative, AoAError, CostLine, evaluate
from cost_core.ingest import InflationTable


def alt(name, lines):
    return Alternative(name, [CostLine(n, phase, {int(y): float(v) for y, v in by.items()})
                              for n, phase, by in lines])


SQ = alt("Status quo", [("O&S", "O&S", {y: 100 for y in range(2027, 2032)})])
NEW = alt("New", [("Buy", "Procurement", {2027: 150}),
                  ("O&S", "O&S", {y: 50 for y in range(2027, 2032)})])


def run(alts, rate=0.0, status_quo="Status quo"):
    infl = InflationTable.from_rate(0.02, base_year=2026, first_year=2020, last_year=2070)
    return evaluate(alts, base_year=2026, inflation=infl, discount_rate=rate, n_iter=200,
                    seed=1, status_quo=status_quo).economic()


def test_at_a_zero_rate_the_metrics_are_plain_sums():
    e = run([SQ, NEW]).summary.set_index("alternative").loc["New"]
    assert (e["investment_pv"], e["savings_pv"], e["net_savings_pv"]) == pytest.approx(
        (150.0, 250.0, 100.0))
    assert e["sir"] == pytest.approx(250 / 150)
    # Differences -100, +50, +50, +50, +50: back to zero by the end of FY2029.
    assert e["break_even_year"] == 2029 and e["payback_year"] == 2029


def test_discounting_follows_the_pv_year():
    r = 0.05
    e = run([SQ, NEW], rate=r).summary.set_index("alternative").loc["New"]
    d = {y: 1 / (1 + r) ** (y - 2026) for y in range(2027, 2032)}
    assert e["investment_pv"] == pytest.approx(150 * d[2027])
    assert e["savings_pv"] == pytest.approx(sum(50 * d[y] for y in d))
    assert e["uniform_annual_cost"] == pytest.approx(
        (150 * d[2027] + sum(50 * d[y] for y in d)) / sum(d.values()))


def test_the_irr_zeroes_the_stream_of_differences():
    e = run([SQ, NEW]).summary.set_index("alternative").loc["New"]
    flows = np.array([-100.0, 50, 50, 50, 50])
    assert np.sum(flows / (1 + e["irr"]) ** np.arange(5)) == pytest.approx(0.0, abs=1e-8)
    assert 0.34 < e["irr"] < 0.36


def test_an_irr_that_isnt_unique_is_left_blank():
    late = alt("Late", [("Buy", "Procurement", {2027: 150}),
                        ("O&S", "O&S", {**{y: 50 for y in range(2027, 2031)}, 2031: 300})])
    e = run([SQ, late]).summary.set_index("alternative").loc["Late"]
    assert e["irr"] is None or np.isnan(e["irr"])


def test_no_extra_investment_has_no_sir():
    from cost_core import plain

    cheap = alt("Cheaper to run", [("O&S", "O&S", {y: 80 for y in range(2027, 2032)})])
    econ = run([SQ, cheap])
    e = econ.summary.set_index("alternative").loc["Cheaper to run"]
    assert np.isnan(e["sir"]) and e["net_savings_pv"] == pytest.approx(100.0)
    assert "needs no more investment" in " ".join(plain.economic(econ, ""))


def test_unequal_lives_are_named_and_compared_per_year():
    from cost_core import plain

    longer = alt("Longer", [("Buy", "Procurement", {2027: 150}),
                            ("O&S", "O&S", {y: 50 for y in range(2027, 2037)})])
    econ = run([SQ, longer])
    assert any("runs to FY2036 and the status quo to FY2031" in n for n in econ.notes)
    assert "service lives differ" in " ".join(plain.economic(econ, ""))


def test_an_unknown_status_quo_is_refused():
    with pytest.raises(AoAError, match="isn't one of the alternatives"):
        run([SQ, NEW], status_quo="Nobody")


def test_the_example_and_the_command(tmp_path, capsys):
    from openpyxl import load_workbook

    cli.main(["demo", "aoa", "--out", str(tmp_path / "d")])
    out = capsys.readouterr().out
    assert "Against the status quo, 'Keep current system'" in out
    assert "pays for itself" in out
    wb = load_workbook(tmp_path / "d" / "report.xlsx")
    assert {"Economic analysis", "Savings by year"} <= set(wb.sheetnames)
    econ = pd.read_csv(tmp_path / "d" / "economic.csv").set_index("alternative")
    assert econ.loc["Upgrade in place", "sir"] > 1 > econ.loc["New development", "sir"]


def test_the_flag_and_the_workbook_name_the_status_quo(tmp_path, capsys):
    from cost_core.examples import example_path
    from cost_core.xlspec import read_spec, write_workbook

    spec = json.loads(example_path("aoa").read_text())
    spec.pop("status_quo")
    p = tmp_path / "aoa.json"
    p.write_text(json.dumps(spec))
    cli.main(["aoa", "--spec", str(p), "--out", str(tmp_path / "a")])
    assert not (tmp_path / "a" / "economic.csv").exists()
    cli.main(["aoa", "--spec", str(p), "--out", str(tmp_path / "b"), "--status-quo",
              "Buy commercial"])
    capsys.readouterr()
    assert (tmp_path / "b" / "economic.csv").exists()
    spec["status_quo"] = "Keep current system"
    back = read_spec(write_workbook(spec, "aoa", tmp_path / "aoa.xlsx"))
    assert back["status_quo"] == "Keep current system"
