# Copyright 2026 Michael Fowler
# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0

"""
costrisk.py - Cost risk analysis from an estimate kept in Excel.

Most estimates live in a workbook: one row per WBS element with its point
estimate, and, once someone has asked "how sure are we?", a low, a most
likely and a high beside it. This module reads that workbook, runs it
through the correlated Monte Carlo engine in :mod:`cost_core.monte_carlo`,
and returns what a cost risk briefing needs: the S-curve, the confidence
table, the drivers and what ignoring correlation would have cost.

The workbook has up to four sheets, and only the first is required:

``Elements``
    Element, Point Estimate, Low, Most Likely, High, and optionally
    Distribution (triangular, the default, pert or uniform). An element
    with no range is carried at its point estimate with no uncertainty.
``Risks``
    Risk, Probability, Low, Most Likely, High (the cost if it happens), and
    optionally Element (which element it belongs to, for the report).
``Correlation``
    Element A, Element B, Correlation. Pairs not listed take the default.
``Settings``
    Units, Default Correlation, Iterations, Seed.

``ce-core template cost-risk`` writes one to fill in, and ``ce-core demo
cost-risk`` runs the invented example that ships with the package.
"""

from __future__ import annotations

import re
import warnings
import zipfile
from dataclasses import dataclass, field
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd

from cost_core.monte_carlo import (
    CorrelationImpact,
    CorrelationWarning,
    CostElement,
    DiscreteRisk,
    RiskModel,
    RiskModelError,
    correlation_impact,
    validate_correlation,
)


class CostRiskError(ValueError):
    """The workbook cannot be read as a cost risk model. The message says
    which sheet, row or column, in words a first-time user can act on."""


#: The headings the template writes, and the other names each is known by.
ELEMENT_COLUMNS: Dict[str, Tuple[str, ...]] = {
    "element": ("element", "wbs_element", "wbs", "name", "item", "description"),
    "point": ("point_estimate", "point", "estimate", "pe", "base", "base_estimate", "cost"),
    "low": ("low", "min", "minimum", "optimistic", "best_case"),
    "mode": ("most_likely", "mode", "ml", "likely"),
    "high": ("high", "max", "maximum", "pessimistic", "worst_case"),
    "distribution": ("distribution", "shape", "dist"),
}
RISK_COLUMNS: Dict[str, Tuple[str, ...]] = {
    "risk": ("risk", "name", "risk_name", "title", "description"),
    "probability": ("probability", "chance", "likelihood", "pof", "p"),
    "low": ELEMENT_COLUMNS["low"],
    "mode": ELEMENT_COLUMNS["mode"] + ("impact", "cost_impact"),
    "high": ELEMENT_COLUMNS["high"],
    "element": ("element", "affects", "wbs_element", "wbs"),
}
PAIR_COLUMNS: Dict[str, Tuple[str, ...]] = {
    "a": ("element_a", "a", "first", "from", "element_1"),
    "b": ("element_b", "b", "second", "to", "element_2"),
    "rho": ("correlation", "rho", "r", "value"),
}

DISTRIBUTIONS = ("triangular", "pert", "uniform")

#: The correlation between any two elements the workbook doesn't pair. 0.3 is
#: the usual default in cost risk practice when nothing better is known; the
#: engine's own default elsewhere in the library stays as it was.
DEFAULT_CORRELATION = 0.3

#: What the settings sheet may hold, with the defaults when it does not.
SETTINGS_DEFAULTS = {"units": "", "default_correlation": DEFAULT_CORRELATION,
                     "iterations": 20_000, "seed": 0}

#: Confidence levels the table reports, as percentiles.
CONFIDENCE_LEVELS = tuple(range(5, 100, 5))


def _norm(name) -> str:
    return re.sub(r"[^a-z0-9]+", "_", str(name).strip().lower()).strip("_")


def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and np.isnan(v)) or \
        (isinstance(v, str) and not v.strip())


def excel_rows(frame: pd.DataFrame) -> List[int]:
    """The Excel row each record of a sheet came from. The heading is row 1,
    and blank rows dropped along the way still count, so a message can send
    someone to the right line."""
    return [int(i) + 2 for i in frame.index]


def open_workbook(path: Path, error=ValueError, hint: str = "") -> Dict[str, pd.DataFrame]:
    """Every sheet of a workbook, or ``error`` saying in plain words why not."""
    try:
        return pd.read_excel(path, sheet_name=None)
    except (ValueError, KeyError, zipfile.BadZipFile) as e:
        raise error(
            f"{path.name} can't be read as an Excel workbook: it may be damaged, still "
            "downloading, password-protected, or an old .xls. Open it in Excel and save "
            f"it again as .xlsx{hint}. (Details: {e})") from None


def _columns(frame: pd.DataFrame, wanted: Dict[str, Tuple[str, ...]], sheet: str,
             required: Sequence[str]) -> Dict[str, str]:
    """Map each wanted field to the sheet's heading for it."""
    have = {_norm(c): c for c in frame.columns}
    found = {}
    for key, aliases in wanted.items():
        for alias in aliases:
            if alias in have:
                found[key] = have[alias]
                break
    missing = [k for k in required if k not in found]
    if missing:
        names = ", ".join(repr(wanted[k][0].replace("_", " ").title()) for k in missing)
        raise CostRiskError(
            f"The {sheet} sheet has no column for {names}. Its headings are "
            f"{', '.join(repr(str(c)) for c in frame.columns) or 'empty'}; "
            "ce-core template cost-risk writes a sheet with the headings it reads.")
    return found


def _number(value, sheet: str, row: int, what: str) -> Optional[float]:
    """A cell as a number, None when blank; a clear error when it is text."""
    if value is None or (isinstance(value, float) and np.isnan(value)):
        return None
    if isinstance(value, str):
        text = value.strip().replace(",", "").replace("$", "")
        if not text:
            return None
        try:
            return float(text)
        except ValueError:
            raise CostRiskError(f"{sheet} sheet, row {row}: the {what} {value!r} is not a "
                                "number.") from None
    return float(value)


def _probability(value, row: int) -> float:
    """0.3, 30% typed as text, or a cell formatted as a percentage."""
    if isinstance(value, str) and value.strip().endswith("%"):
        p = _number(value.strip()[:-1], "Risks", row, "probability")
        p = None if p is None else p / 100.0
    else:
        p = _number(value, "Risks", row, "probability")
    if p is None:
        raise CostRiskError(f"Risks sheet, row {row}: the probability is blank.")
    if not 0.0 <= p <= 1.0:
        raise CostRiskError(
            f"Risks sheet, row {row}: a probability of {p:g} is not between 0 and 1. "
            "Write 30% as 0.3, or as 30% in a cell formatted as a percentage.")
    return p


def _three_point(low, mode, high, sheet: str, row: int, label: str,
                 dist: str = "triangular", point: Optional[float] = None) -> Dict:
    """A distribution spec from a low, most likely and high."""
    if mode is None:
        mode = point
    given = [v is not None for v in (low, mode, high)]
    if not any(given):
        raise CostRiskError(f"{sheet} sheet, row {row} ({label}): no cost is given.")
    if given == [False, True, False] or (low == mode == high):
        return {"type": "fixed", "value": float(mode if mode is not None else low)}
    if dist == "uniform":
        if low is None or high is None:
            raise CostRiskError(f"{sheet} sheet, row {row} ({label}): a uniform range "
                                "needs both a Low and a High.")
        if high < low:
            raise CostRiskError(f"{sheet} sheet, row {row} ({label}): the High ({high:,g}) "
                                f"is below the Low ({low:,g}).")
        return {"type": "uniform", "low": float(low), "high": float(high)}
    if not all(given):
        blank = [n for n, g in zip(("Low", "Most Likely", "High"), given) if not g]
        raise CostRiskError(f"{sheet} sheet, row {row} ({label}): the {' and '.join(blank)} "
                            "is blank. Give all three, or leave Low and High blank to "
                            "carry it with no uncertainty.")
    if not low <= mode <= high:
        raise CostRiskError(f"{sheet} sheet, row {row} ({label}): the numbers must run "
                            f"Low <= Most Likely <= High, and they are {low:,g}, {mode:,g}, "
                            f"{high:,g}.")
    return {"type": dist, "left": float(low), "mode": float(mode), "right": float(high)}


@dataclass
class CostRiskInput:
    """What the workbook says: the model and how to run it."""

    model: RiskModel
    elements: pd.DataFrame
    risks: pd.DataFrame
    pairs: pd.DataFrame
    units: str = ""
    n_iter: int = 20_000
    seed: int = 0
    notes: List[str] = field(default_factory=list)


def _read_elements(frame: pd.DataFrame) -> Tuple[List[CostElement], pd.DataFrame]:
    frame = frame.dropna(how="all")
    cols = _columns(frame, ELEMENT_COLUMNS, "Elements", ["element"])
    if not ({"point", "mode"} & set(cols)):
        raise CostRiskError("The Elements sheet needs a 'Point Estimate' or a 'Most Likely' "
                            "column; ce-core template cost-risk writes one with both.")
    elements, rows = [], []
    for i, rec in zip(excel_rows(frame), frame.to_dict("records")):
        name = rec.get(cols["element"])
        if name is None or (isinstance(name, float) and np.isnan(name)) or not str(name).strip():
            raise CostRiskError(f"Elements sheet, row {i}: the element has no name.")
        name = str(name).strip()
        get = {k: _number(rec.get(cols[k]), "Elements", i, k.replace("mode", "most likely"))
               for k in ("point", "low", "mode", "high") if k in cols}
        dist = str(rec.get(cols["distribution"]) or "triangular").strip().lower() \
            if "distribution" in cols else "triangular"
        if dist in ("", "nan", "tri", "triangle"):
            dist = "triangular"
        if dist not in DISTRIBUTIONS:
            raise CostRiskError(f"Elements sheet, row {i} ({name}): the distribution {dist!r} "
                                f"is not one of {', '.join(DISTRIBUTIONS)}.")
        point = get.get("point")
        spec = _three_point(get.get("low"), get.get("mode"), get.get("high"), "Elements", i,
                            name, dist, point)
        if point is None:
            point = get.get("mode")
        if point is None:
            raise CostRiskError(f"Elements sheet, row {i} ({name}): give a Point Estimate or "
                                "a Most Likely.")
        elements.append(CostElement(name, spec, point_estimate=float(point)))
        rows.append({"element": name, "point_estimate": point, "low": get.get("low"),
                     "most_likely": get["mode"] if get.get("mode") is not None else point,
                     "high": get.get("high"),
                     "distribution": spec["type"]})
    if not elements:
        raise CostRiskError("The Elements sheet has no rows. Put one WBS element on each row.")
    names = [e.name for e in elements]
    dup = sorted({n for n in names if names.count(n) > 1})
    if dup:
        raise CostRiskError(f"The Elements sheet lists {', '.join(map(repr, dup))} more than "
                            "once. Each element needs its own name, for the drivers chart and "
                            "the correlation pairs.")
    return elements, pd.DataFrame(rows)


def _read_risks(frame: Optional[pd.DataFrame], names: Sequence[str]
                ) -> Tuple[List[DiscreteRisk], pd.DataFrame]:
    empty = pd.DataFrame(columns=["risk", "probability", "low", "most_likely", "high", "element"])
    if frame is None:
        return [], empty
    frame = frame.dropna(how="all")
    if frame.empty:
        return [], empty
    cols = _columns(frame, RISK_COLUMNS, "Risks", ["risk", "probability"])
    risks, rows = [], []
    for i, rec in zip(excel_rows(frame), frame.to_dict("records")):
        name = str(rec.get(cols["risk"]) or "").strip()
        if not name or name == "nan":
            raise CostRiskError(f"Risks sheet, row {i}: the risk has no name.")
        p = _probability(rec.get(cols["probability"]), i)
        get = {k: _number(rec.get(cols[k]), "Risks", i, k.replace("mode", "most likely"))
               for k in ("low", "mode", "high") if k in cols}
        spec = _three_point(get.get("low"), get.get("mode"), get.get("high"), "Risks", i, name)
        where = rec.get(cols["element"]) if "element" in cols else None
        where = None if where is None or str(where).strip() in ("", "nan") else str(where).strip()
        if where is not None and where not in names:
            raise CostRiskError(f"Risks sheet, row {i} ({name}): the element {where!r} is not "
                                "on the Elements sheet.")
        risks.append(DiscreteRisk(name, p, spec, affects=where))
        rows.append({"risk": name, "probability": p, "low": get.get("low"),
                     "most_likely": get.get("mode"), "high": get.get("high"),
                     "element": where})
    names_r = [r.name for r in risks]
    dup = sorted({n for n in names_r if names_r.count(n) > 1})
    if dup:
        raise CostRiskError(f"The Risks sheet lists {', '.join(map(repr, dup))} more than once.")
    return risks, pd.DataFrame(rows)


def _read_pairs(frame: Optional[pd.DataFrame], names: Sequence[str], default: float
                ) -> Tuple[np.ndarray, pd.DataFrame, List[str]]:
    k = len(names)
    matrix = np.full((k, k), float(default))
    np.fill_diagonal(matrix, 1.0)
    pairs = pd.DataFrame(columns=["element_a", "element_b", "correlation"])
    if frame is None or frame.dropna(how="all").empty:
        return matrix, pairs, []
    frame = frame.dropna(how="all")
    cols = _columns(frame, PAIR_COLUMNS, "Correlation", ["a", "b", "rho"])
    index = {n: j for j, n in enumerate(names)}
    rows = []
    for i, rec in zip(excel_rows(frame), frame.to_dict("records")):
        if any(_blank(rec.get(cols[c])) for c in ("a", "b")):
            raise CostRiskError(f"Correlation sheet, row {i}: name both elements of the "
                                "pair, or delete the row.")
        a, b = (str(rec.get(cols[c])).strip() for c in ("a", "b"))
        for n in (a, b):
            if n not in index:
                raise CostRiskError(f"Correlation sheet, row {i}: {n!r} is not on the "
                                    "Elements sheet (names have to match exactly).")
        if a == b:
            raise CostRiskError(f"Correlation sheet, row {i}: {a!r} is paired with itself.")
        rho = _number(rec.get(cols["rho"]), "Correlation", i, "correlation")
        if rho is None:
            continue
        if not -1.0 <= rho <= 1.0:
            raise CostRiskError(f"Correlation sheet, row {i}: a correlation of {rho:g} is not "
                                "between -1 and 1.")
        matrix[index[a], index[b]] = matrix[index[b], index[a]] = rho
        rows.append({"element_a": a, "element_b": b, "correlation": rho})
    notes = []
    with warnings.catch_warnings():
        # Said below in plain words instead.
        warnings.simplefilter("ignore", CorrelationWarning)
        fixed, repair = validate_correlation(matrix, list(names))
    if repair:
        notes.append("The correlations entered can't all hold at once, so the nearest set "
                     "that can was used; the Correlation sheet of report.xlsx shows it.")
    return fixed, pd.DataFrame(rows, columns=pairs.columns), notes


#: The Settings sheet's labels, as normalised headings, and what each sets.
SETTING_NAMES = {"units": "units", "default_correlation": "default_correlation",
                 "iterations": "iterations", "iters": "iterations", "draws": "iterations",
                 "simulations": "iterations", "seed": "seed", "random_seed": "seed",
                 "description": None, "notes": None}


def _read_settings(frame: Optional[pd.DataFrame]) -> Dict:
    out = dict(SETTINGS_DEFAULTS)
    if frame is None or frame.empty:
        return out
    frame = frame.dropna(how="all")
    for row, rec in zip(excel_rows(frame), frame.itertuples(index=False)):
        if len(rec) < 2 or _blank(rec[0]):
            continue
        label = _norm(rec[0])
        if label not in SETTING_NAMES:
            # A misspelt setting would otherwise be ignored without a word, and
            # the default used in its place.
            raise CostRiskError(
                f"Settings sheet, row {row}: {str(rec[0]).strip()!r} is not a setting. The "
                "settings are Units, Default Correlation, Iterations and Seed.")
        key, value = SETTING_NAMES[label], rec[1]
        if key is None or _blank(value):
            continue
        if key == "units":
            out[key] = str(value).strip()
            continue
        v = _number(value, "Settings", row, str(rec[0]).strip())
        if key in ("iterations", "seed"):
            if v != int(v) or v < 0:
                raise CostRiskError(f"Settings sheet, row {row}: {str(rec[0]).strip()} is a "
                                    f"whole number, 0 or more, not {v:g}.")
            v = int(v)
        out[key] = v
    return out


def read_workbook(path) -> CostRiskInput:
    """Read a cost risk workbook (or a CSV of elements alone)."""
    path = Path(path)
    if path.suffix.lower() == ".csv":
        sheets = {"elements": pd.read_csv(path)}
    else:
        raw = open_workbook(path, CostRiskError, ", or give the elements alone as a .csv")
        sheets = {_norm(k): v for k, v in raw.items()}
        if "elements" not in sheets:
            first = next(iter(raw), None)
            if first is None:
                raise CostRiskError(f"{path.name} has no sheets.")
            sheets["elements"] = raw[first]
    settings = _read_settings(sheets.get("settings"))
    elements, element_rows = _read_elements(sheets["elements"])
    names = [e.name for e in elements]
    risks, risk_rows = _read_risks(sheets.get("risks"), names)
    default = settings["default_correlation"]
    if not -1.0 < default < 1.0:
        raise CostRiskError(f"Settings sheet: a default correlation of {default:g} is not "
                            "between -1 and 1.")
    matrix, pairs, notes = _read_pairs(sheets.get("correlation"), names, default)
    for e in elements:
        d = e.distribution
        lo, hi = d.get("left", d.get("low")), d.get("right", d.get("high"))
        if lo is not None and not lo <= e.point_estimate <= hi:
            notes.append(f"{e.name!r} has a point estimate of {e.point_estimate:,g}, outside "
                         f"its own range of {lo:,g} to {hi:,g}: check the units, or the "
                         "range.")
    fixed = [e.name for e in elements if e.distribution["type"] == "fixed"]
    if fixed:
        notes.append(f"{len(fixed)} element{'s' if len(fixed) > 1 else ''} "
                     f"({', '.join(fixed[:4])}{', ...' if len(fixed) > 4 else ''}) "
                     f"{'have' if len(fixed) > 1 else 'has'} no range, so "
                     f"{'they are' if len(fixed) > 1 else 'it is'} counted at the point "
                     "estimate with no uncertainty.")
    try:
        model = RiskModel(elements=elements, risks=risks, correlation=matrix,
                          default_correlation=default, name=path.stem)
    except RiskModelError as e:
        raise CostRiskError(str(e)) from None
    return CostRiskInput(model=model, elements=element_rows, risks=risk_rows, pairs=pairs,
                         units=settings["units"], n_iter=settings["iterations"],
                         seed=settings["seed"], notes=notes)


@dataclass
class CostRiskResult:
    """A cost risk analysis: the simulation with correlation, the same model
    without it, and the input it came from."""

    inputs: CostRiskInput
    impact: CorrelationImpact
    units: str = ""

    @property
    def sim(self):
        return self.impact.correlated

    @property
    def point_estimate(self) -> float:
        return float(self.sim.point_estimate)

    def confidence_table(self, levels: Sequence[int] = CONFIDENCE_LEVELS) -> pd.DataFrame:
        """The cost at each confidence level and the reserve it needs over
        the point estimate."""
        cost = np.percentile(self.sim.totals, levels)
        point = self.point_estimate
        return pd.DataFrame({"confidence": np.asarray(levels) / 100.0, "cost": cost,
                             "reserve": cost - point,
                             "reserve_pct": cost / point - 1.0 if point else np.nan})

    def drivers(self) -> pd.DataFrame:
        return self.sim.tornado()

    def element_table(self) -> pd.DataFrame:
        """Each element as entered, beside its simulated mean, P50 and P80."""
        s = self.sim.element_samples
        out = self.inputs.elements.copy()
        out["mean"] = s.mean(axis=0)
        out["p50"] = np.percentile(s, 50, axis=0)
        out["p80"] = np.percentile(s, 80, axis=0)
        return out

    def risk_table(self) -> pd.DataFrame:
        out = self.inputs.risks.copy()
        if len(out):
            out["mean_if_it_happens"] = [r.frozen().mean() for r in self.inputs.model.risks]
            out["expected_cost"] = [r.expected_value for r in self.inputs.model.risks]
        return out

    def correlation_matrix(self) -> pd.DataFrame:
        names = self.inputs.model.element_names
        return pd.DataFrame(self.inputs.model.correlation, index=names, columns=names)

    @property
    def assumptions(self) -> Dict:
        m = self.inputs.model
        return {"iterations": self.sim.n_iter, "seed": self.inputs.seed,
                "method": "Gaussian copula over the element distributions; discrete "
                          "risks drawn independently of the elements and of each other",
                "elements": len(m.elements), "discrete risks": len(m.risks),
                "default correlation": m.default_correlation,
                "correlation pairs entered": len(self.inputs.pairs),
                "point estimate": "sum of the element point estimates (risks excluded)",
                "units": self.units or "as entered",
                "p80 settled": "yes" if self.sim.is_converged else
                               "no: run more iterations before quoting the P80"}


def analyse(inputs: CostRiskInput, n_iter: Optional[int] = None, seed: Optional[int] = None,
            units: Optional[str] = None) -> CostRiskResult:
    """Run the workbook's model, with and without its correlation."""
    n = int(n_iter or inputs.n_iter)
    if n < 1000:
        raise CostRiskError(f"{n} iterations is too few to read a P80 from; use 1,000 or more.")
    seed = inputs.seed if seed is None else seed
    if seed < 0:
        raise CostRiskError(f"The seed is a whole number, 0 or more, not {seed}.")
    with warnings.catch_warnings():
        warnings.simplefilter("ignore", CorrelationWarning)
        impact = correlation_impact(inputs.model, n_iter=n, seed=seed)
    if not impact.correlated.std > 0:
        m = inputs.model
        raise CostRiskError(
            f"There's no uncertainty to simulate: none of the {len(m.elements)} elements has "
            f"a range{' and none of the risks can vary' if m.risks else ' and there are no risks'}"
            f", so every confidence level is the point estimate, {m.point_estimate:,g}. Give "
            "the elements a Low and a High (the Instructions sheet says how), or add the "
            "risks.")
    return CostRiskResult(inputs=inputs, impact=impact,
                          units=inputs.units if units is None else units)


# ------------------------------------------------------------- the template
INSTRUCTIONS = [
    ("What this is", "A cost estimate with its uncertainty. ce-core cost-risk --data "
     "<this file> simulates it and says how much it takes to be 50%, 80% or 90% sure, "
     "which elements and risks drive that, and what the correlation between elements adds."),
    ("Elements sheet", "One row per WBS element. Element is its name. Point Estimate is the "
     "number in the estimate. Low, Most Likely and High are the range you believe: Low the "
     "least it could plausibly cost, High the most. Leave Low and High blank for an element "
     "with no uncertainty (it is carried at its point estimate)."),
    ("Distribution", "Optional. triangular (the default) spreads the chance evenly along "
     "the three points; pert bunches it toward the most likely; uniform treats every value "
     "between Low and High alike and ignores the most likely."),
    ("Risks sheet", "Things that may or may not happen. Probability is the chance it does "
     "(0.3 or 30%). Low, Most Likely and High are what it costs if it does. Element is "
     "optional and only labels the report. Delete the rows if there are no risks."),
    ("Correlation sheet", "Pairs of elements that tend to overrun together, and how strongly "
     "(0 none, 1 in lockstep; 0.2 to 0.3 is typical, 0.5 or more for elements driven by the "
     "same thing). Every pair not listed takes the Default Correlation on the Settings sheet. "
     "Leaving correlation out makes the P80 too low."),
    ("Settings sheet", "Units labels the money (thousands, millions or any word). Default "
     "Correlation applies to every pair not listed. Iterations and Seed control the "
     "simulation; the same seed gives the same answer."),
    ("Money", "Use one unit and one dollar basis (for example constant FY2026 $M) throughout."),
]

#: The invented example: a ground station upgrade, in $M. Every number is made up.
EXAMPLE_ELEMENTS = pd.DataFrame(
    [("1.0 Program management", 12.0, 11.0, 12.0, 18.0, "triangular"),
     ("2.0 Systems engineering", 18.5, 17.0, 18.5, 28.0, "triangular"),
     ("3.0 Antenna subsystem", 42.0, 37.0, 42.0, 72.0, "triangular"),
     ("4.0 Software", 36.0, 31.0, 36.0, 90.0, "triangular"),
     ("5.0 Facilities", 22.0, 20.0, 22.0, 31.0, "pert"),
     ("6.0 Integration and test", 15.0, 13.0, 15.0, 30.0, "triangular"),
     ("7.0 Training and support", 6.5, 6.0, 6.5, 10.0, "pert"),
     ("8.0 Initial spares", 4.0, None, None, None, "")],
    columns=["Element", "Point Estimate", "Low", "Most Likely", "High", "Distribution"])
EXAMPLE_RISKS = pd.DataFrame(
    [("Antenna vendor requalification", 0.25, 4.0, 6.0, 10.0, "3.0 Antenna subsystem"),
     ("Cyber accreditation rework", 0.40, 2.0, 3.0, 6.0, "4.0 Software"),
     ("Site permitting delay", 0.15, 1.5, 3.0, 5.0, "5.0 Facilities")],
    columns=["Risk", "Probability", "Low", "Most Likely", "High", "Element"])
EXAMPLE_PAIRS = pd.DataFrame(
    [("4.0 Software", "6.0 Integration and test", 0.6),
     ("1.0 Program management", "2.0 Systems engineering", 0.5),
     ("2.0 Systems engineering", "4.0 Software", 0.4)],
    columns=["Element A", "Element B", "Correlation"])
EXAMPLE_SETTINGS = pd.DataFrame(
    [("Units", "millions"), ("Default Correlation", 0.3), ("Iterations", 20000), ("Seed", 1)],
    columns=["Setting", "Value"])


def write_workbook(out, elements: pd.DataFrame = EXAMPLE_ELEMENTS,
                   risks: pd.DataFrame = EXAMPLE_RISKS, pairs: pd.DataFrame = EXAMPLE_PAIRS,
                   settings: pd.DataFrame = EXAMPLE_SETTINGS) -> Path:
    """Write a workbook laid out the way :func:`read_workbook` reads it."""
    from openpyxl.styles import Alignment, Font

    out = Path(out)
    with pd.ExcelWriter(out, engine="openpyxl") as xl:
        elements.to_excel(xl, sheet_name="Elements", index=False)
        risks.to_excel(xl, sheet_name="Risks", index=False)
        pairs.to_excel(xl, sheet_name="Correlation", index=False)
        settings.to_excel(xl, sheet_name="Settings", index=False)
        pd.DataFrame(INSTRUCTIONS, columns=["Sheet or topic", "What to put there"]) \
            .to_excel(xl, sheet_name="Instructions", index=False)
        widths = {"Elements": (32, 16, 12, 14, 12, 14), "Risks": (34, 13, 10, 14, 10, 30),
                  "Correlation": (32, 32, 13), "Settings": (22, 14),
                  "Instructions": (22, 100)}
        for name, cols in widths.items():
            ws = xl.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for j, w in enumerate(cols):
                ws.column_dimensions[chr(ord("A") + j)].width = w
        ws = xl.sheets["Risks"]
        for row in ws.iter_rows(min_row=2, min_col=2, max_col=2):
            row[0].number_format = "0%"
        for row in xl.sheets["Instructions"].iter_rows(min_row=2):
            row[0].alignment = Alignment(vertical="top")
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
    return out
