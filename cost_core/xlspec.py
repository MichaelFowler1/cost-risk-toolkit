# Copyright 2026 Michael Fowler
# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0

"""
xlspec.py - Excel workbooks for the JCL, AoA and portfolio specs.

The three spec-driven commands were written against JSON, which suits a
script and suits nobody who lives in Excel. This module reads a workbook
into exactly the dictionary the JSON file would have held, so every loader
takes either and the engines never know the difference; and it writes a
spec back out as a workbook, which is what ``ce-core template`` hands out.

Each workbook has a Settings sheet (Setting, Value) for the scalar fields and
one sheet per list, with an Instructions sheet saying what goes where:

JCL
    Activities (ID, Name, Duration, Duration Low/Most Likely/High in months,
    Predecessors, Fixed Cost, Cost Low/Most Likely/High, Burn Rate) and
    Risks (Risk, Probability, Activities, Delay Low/Most Likely/High, Cost).
AoA
    Alternatives (Alternative, Effectiveness, Correlation), Lines (one cost
    line per row: spread as Total, Start Year, Years, Profile; or level as
    Annual Amount, First Year, Last Year; or by year on the Phased sheet),
    and Phased (Alternative, Line, then one column per year).
Portfolio
    Candidates (one funding option per row, with a column per year of cost),
    Budget (Year, Budget) and Exclusive (Candidates, separated by ;).

Lists inside a cell (predecessors, the activities a risk hits, what a
candidate requires) are separated by semicolons.
"""

from __future__ import annotations

import json
import re
from pathlib import Path
from typing import Any, Callable, Dict, List, Optional, Sequence, Tuple

import numpy as np
import pandas as pd

from cost_core.costrisk import CostRiskError, _norm, _number, excel_rows, open_workbook


class SpecError(ValueError):
    """A workbook that can't be read as a spec; the message names the sheet,
    row and column."""


def read_spec(path) -> Dict[str, Any]:
    """A spec file as a dictionary: JSON as it is, a workbook converted.

    Which of the three a workbook holds is told by its sheets.
    """
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix in (".xlsx", ".xlsm"):
        return read_workbook(path)
    if suffix in (".xls", ".xlsb", ".ods", ".csv"):
        raise SpecError(f"{path.name}: a spec is an .xlsx workbook or a .json file. Open it "
                        "in Excel and save it as .xlsx (Excel Workbook).")
    try:
        return json.loads(path.read_text(encoding="utf-8"))
    except UnicodeDecodeError:
        raise SpecError(f"{path.name} isn't a JSON spec or an .xlsx workbook. If it's a "
                        "spreadsheet, save it from Excel as .xlsx.") from None
    except json.JSONDecodeError as e:
        raise SpecError(f"{path.name} is not valid JSON ({e}). For a spreadsheet, save it "
                        "as .xlsx.") from None


# ------------------------------------------------------------------ helpers
def _blank(v) -> bool:
    return v is None or (isinstance(v, float) and np.isnan(v)) or \
        (isinstance(v, str) and not v.strip())


def _num(v, sheet: str, row: int, what: str) -> Optional[float]:
    try:
        return _number(v, sheet, row, what)
    except CostRiskError as e:
        raise SpecError(str(e)) from None


def _text(v) -> str:
    if _blank(v):
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _items(v) -> List[str]:
    """``a; b; c`` as a list."""
    return [t.strip() for t in re.split(r"[;\n]", _text(v)) if t.strip()]


def _year(c) -> Optional[int]:
    """A heading that is a year (2027, "2027", "FY2027"), else None."""
    m = re.fullmatch(r"(?:fy)?\s*(\d{4})(?:\.0)?", str(c).strip().lower())
    return int(m.group(1)) if m and 1900 < int(m.group(1)) < 2300 else None


class _Sheet:
    """One sheet's rows, with its headings matched loosely."""

    def __init__(self, name: str, frame: Optional[pd.DataFrame]):
        self.name = name
        self.frame = pd.DataFrame() if frame is None else frame.dropna(how="all")
        self.cols = {_norm(c): c for c in self.frame.columns}

    def __bool__(self):
        return not self.frame.empty

    def has(self, *names: str) -> Optional[str]:
        for n in names:
            if _norm(n) in self.cols:
                return self.cols[_norm(n)]
        return None

    def need(self, *names: str) -> str:
        col = self.has(*names)
        if col is None:
            raise SpecError(f"The {self.name} sheet has no {names[0]!r} column. Its headings "
                            f"are {', '.join(repr(str(c)) for c in self.frame.columns)}.")
        return col

    def rows(self):
        """(Excel row number, record) for each row."""
        yield from zip(excel_rows(self.frame), self.frame.to_dict("records"))

    def years(self) -> Dict[int, str]:
        return {y: c for c in self.frame.columns if (y := _year(c)) is not None}


def _get(sheet: _Sheet, rec, row: int, *names, number=True, default=None):
    col = sheet.has(*names)
    if col is None or _blank(rec.get(col)):
        return default
    return _num(rec[col], sheet.name, row, names[0]) if number else _text(rec[col])


def _dist(low, mode, high, dist: str, where: str, scale: float = 1.0):
    """A distribution spec from three points, as factors of ``scale``."""
    if low is None and mode is None and high is None:
        return None
    if scale != 1.0:
        # 16.8 months over 12 is 1.4000000000000001 in binary; rounding keeps
        # a workbook and the JSON it came from drawing exactly the same numbers.
        low, mode, high = (None if v is None else round(v / scale, 12) * scale
                           for v in (low, mode, high))
    dist = (dist or "triangular").strip().lower()
    if dist in ("tri", "triangle", "nan"):
        dist = "triangular"
    if dist == "uniform":
        if low is None or high is None:
            raise SpecError(f"{where}: a uniform range needs both a Low and a High.")
        return {"type": "uniform", "low": round(low / scale, 12), "high": round(high / scale, 12)}
    if dist not in ("triangular", "pert"):
        raise SpecError(f"{where}: the distribution {dist!r} is not triangular, pert or "
                        "uniform.")
    if mode is None:
        mode = scale
    if low is None or high is None:
        raise SpecError(f"{where}: give both a Low and a High, or neither.")
    if not low <= mode <= high:
        raise SpecError(f"{where}: the numbers must run Low <= Most Likely <= High, and they "
                        f"are {low:g}, {mode:g}, {high:g}.")
    return {"type": dist, "left": round(low / scale, 12), "mode": round(mode / scale, 12),
            "right": round(high / scale, 12)}


def _trio(spec, scale: float = 1.0) -> Tuple[Any, Any, Any, str]:
    """The reverse of :func:`_dist`, for writing a template."""
    if spec is None:
        return None, None, None, ""
    t = spec.get("type")
    if t in ("triangular", "pert"):
        return (spec["left"] * scale, spec["mode"] * scale, spec["right"] * scale,
                "" if t == "triangular" else t)
    if t == "uniform":
        return spec["low"] * scale, None, spec["high"] * scale, "uniform"
    raise SpecError(f"A {t!r} distribution has no three-point form for a workbook; keep "
                    "this spec as JSON.")


def _r(v, nd: int = 10):
    """Round away the float noise a factor times a duration leaves behind."""
    return None if v is None else round(float(v), nd)


# ----------------------------------------------------------------- settings
#: Each spec's scalar fields: (heading, key, type).
SETTINGS: Dict[str, List[Tuple[str, str, Callable]]] = {
    "jcl": [("Name", "name", str), ("Units", "units", str),
            ("Standing Army", "standing_army", float),
            ("Duration Correlation", "duration_correlation", float),
            ("Cost Correlation", "cost_correlation", float),
            ("Cross Correlation", "cross_correlation", float),
            ("Confidence", "confidence", float), ("Iterations", "n_iter", int),
            ("Seed", "seed", int)],
    "aoa": [("Base Year", "base_year", int), ("Discount Rate", "discount_rate", float),
            ("Inflation Rate", "inflation_rate", float),
            ("Inflation CSV", "inflation_csv", str), ("Units", "units", str),
            ("Basis", "basis", str), ("PV Year", "pv_year", int),
            ("Iterations", "n_iter", int), ("Seed", "seed", int)],
    "portfolio": [("Units", "units", str), ("Delta", "delta", float),
                  ("Growth Low", "growth_low", float),
                  ("Growth Most Likely", "growth_mode", float),
                  ("Growth High", "growth_high", float),
                  ("Growth Distribution", "growth_dist", str),
                  ("Growth Correlation", "growth_correlation", float),
                  ("Frontier Scales", "frontier_scales", str),
                  ("Iterations", "n_iter", int), ("Seed", "seed", int)],
}


def _read_settings(kind: str, frame: Optional[pd.DataFrame]) -> Dict[str, Any]:
    known = {_norm(h): (key, cast) for h, key, cast in SETTINGS[kind]}
    known.update({"description": ("description", str), "n_iter": ("n_iter", int)})
    out: Dict[str, Any] = {}
    if frame is None:
        return out
    frame = frame.dropna(how="all")
    for i, rec in zip(excel_rows(frame), frame.itertuples(index=False)):
        if len(rec) < 2 or _blank(rec[0]) or _blank(rec[1]):
            continue
        label = _norm(rec[0])
        if label not in known:
            raise SpecError(f"Settings sheet, row {i}: {str(rec[0])!r} is not a setting here. "
                            f"The settings are {', '.join(h for h, _, _ in SETTINGS[kind])}.")
        key, cast = known[label]
        if cast is str:
            out[key] = _text(rec[1])
        else:
            v = _num(rec[1], "Settings", i, str(rec[0]))
            if key in ("seed", "n_iter") and (v != int(v) or v < (0 if key == "seed" else 1)):
                raise SpecError(f"Settings sheet, row {i}: {str(rec[0]).strip()} is a whole "
                                f"number, {'0' if key == 'seed' else '1'} or more, not {v:g}.")
            out[key] = int(v) if cast is int else v
    return out


def _settings_frame(kind: str, spec: Dict[str, Any]) -> pd.DataFrame:
    rows = [(h, spec[key]) for h, key, _ in SETTINGS[kind] if key in spec]
    if "description" in spec:
        rows.append(("Description", spec["description"]))
    return pd.DataFrame(rows, columns=["Setting", "Value"])


# ---------------------------------------------------------------------- JCL
_LINK = re.compile(r"^(?P<id>.+?)(?:\s+(?P<type>FS|SS|FF|SF))?(?:\s*(?P<lag>[+-]\s*\d+(?:\.\d+)?))?$",
                   re.IGNORECASE)


def _link(text: str, ids: Sequence[str], where: str):
    """``design``, ``design -4``, ``bus SS +2`` as the forms Activity reads."""
    if text in ids:
        return text
    m = _LINK.match(text)
    if not m or m.group("id").strip() not in ids:
        raise SpecError(f"{where}: the predecessor {text!r} is not an activity ID. Write an "
                        "ID, optionally with a link type and a lag in months: design, "
                        "design -4, bus SS +2.")
    lag = float(m.group("lag").replace(" ", "")) if m.group("lag") else 0.0
    kind = (m.group("type") or "FS").upper()
    if kind == "FS":
        return [m.group("id").strip(), lag]
    return [m.group("id").strip(), lag, kind]


def _link_text(p) -> str:
    if isinstance(p, str):
        return p
    if isinstance(p, dict):
        p = [p["id"], p.get("lag", 0), p.get("type", "FS")]
    p = list(p)
    kind = p[2].upper() if len(p) > 2 else "FS"
    lag = float(p[1]) if len(p) > 1 else 0.0
    out = p[0] + ("" if kind == "FS" else f" {kind}")
    return out + (f" {lag:+g}" if lag else "")


def _read_jcl(sheets: Dict[str, pd.DataFrame]) -> Dict[str, Any]:
    spec = _read_settings("jcl", sheets.get("settings"))
    acts = _Sheet("Activities", sheets.get("activities"))
    if not acts:
        raise SpecError("The Activities sheet has no rows.")
    id_col = acts.need("ID", "Activity ID")
    ids = [_text(r[id_col]) for _, r in acts.rows()]
    dup = sorted({i for i in ids if ids.count(i) > 1})
    if dup or "" in ids:
        raise SpecError("Every activity needs an ID of its own on the Activities sheet"
                        + (f"; {', '.join(map(repr, dup))} appear more than once." if dup else "."))
    spec["activities"] = []
    for row, rec in acts.rows():
        aid = _text(rec[id_col])
        where = f"Activities sheet, row {row} ({aid})"
        g = lambda *n: _get(acts, rec, row, *n)  # noqa: E731
        ml = g("Duration Most Likely", "Most Likely Duration")
        duration = g("Duration") if g("Duration") is not None else ml
        if duration is None:
            raise SpecError(f"{where}: no Duration.")
        a: Dict[str, Any] = {"id": aid}
        name = _get(acts, rec, row, "Name", number=False)
        if name:
            a["name"] = name
        a["duration"] = duration
        preds = [_link(t, ids, where) for t in _items(rec.get(acts.has("Predecessors") or ""))]
        if preds:
            a["predecessors"] = preds
        dist = _get(acts, rec, row, "Duration Distribution", number=False, default="")
        du = _dist(g("Duration Low"), ml, g("Duration High"), dist, where, duration or 1.0)
        if du and duration == 0:
            raise SpecError(f"{where}: a zero duration can't carry a range.")
        if du:
            a["duration_uncertainty"] = du
        fixed = g("Fixed Cost")
        cml = g("Cost Most Likely")
        if fixed is None:
            fixed = cml
        if not fixed and (g("Cost Low") is not None or g("Cost High") is not None):
            # Otherwise the range would be dropped without a word.
            raise SpecError(f"{where}: there's a Cost Low or High but no Fixed Cost (or Cost "
                            "Most Likely) for it to range around.")
        if fixed:
            a["fixed_cost"] = fixed
            cu = _dist(g("Cost Low"), cml, g("Cost High"),
                       _get(acts, rec, row, "Cost Distribution", number=False, default=""),
                       where, fixed)
            if cu:
                a["fixed_cost_uncertainty"] = cu
        burn = g("Burn Rate")
        if burn:
            a["burn_rate"] = burn
        spec["activities"].append(a)
    risks = _Sheet("Risks", sheets.get("risks"))
    if risks:
        spec["risks"] = []
        name_col = risks.need("Risk", "Name")
        for row, rec in risks.rows():
            where = f"Risks sheet, row {row}"
            r: Dict[str, Any] = {"name": _text(rec[name_col])}
            if not r["name"]:
                raise SpecError(f"{where}: the risk has no name.")
            p = rec.get(risks.need("Probability"))
            if isinstance(p, str) and p.strip().endswith("%"):
                p = _num(p.strip()[:-1], "Risks", row, "probability") / 100.0
            else:
                p = _num(p, "Risks", row, "probability")
            if p is None or not 0 <= p <= 1:
                raise SpecError(f"{where}: the probability must be between 0 and 1 (0.3 or "
                                "30%).")
            r["probability"] = p
            hit = _items(rec.get(risks.has("Activities", "Activity") or ""))
            bad = [h for h in hit if h not in ids]
            if bad:
                raise SpecError(f"{where}: {', '.join(map(repr, bad))} is not an activity ID.")
            r["activities"] = hit
            g = lambda *n: _get(risks, rec, row, *n)  # noqa: E731
            lo, ml, hi = g("Delay Low"), g("Delay Most Likely", "Delay"), g("Delay High")
            if lo is None and hi is None:
                if ml:
                    r["delay"] = ml
            else:
                if ml is None:
                    raise SpecError(f"{where}: give a Delay Most Likely as well as the Low and "
                                    "High, or a single Delay.")
                r["delay"] = _dist(lo, ml, hi,
                                   _get(risks, rec, row, "Delay Distribution", number=False,
                                        default=""), where)
            cost = g("Cost")
            if cost:
                r["cost"] = cost
            spec["risks"].append(r)
    return spec


def _jcl_frames(spec) -> Dict[str, pd.DataFrame]:
    rows = []
    for a in spec["activities"]:
        d = float(a["duration"])
        dl, dm, dh, dd = _trio(a.get("duration_uncertainty"), d)
        fc = float(a.get("fixed_cost", 0.0))
        cl, cm, ch, cd = _trio(a.get("fixed_cost_uncertainty"), fc)
        rows.append({"ID": a["id"], "Name": a.get("name", ""), "Duration": d,
                     "Duration Low": _r(dl), "Duration Most Likely": _r(dm),
                     "Duration High": _r(dh), "Duration Distribution": dd,
                     "Predecessors": "; ".join(_link_text(p) for p in a.get("predecessors", [])),
                     "Fixed Cost": fc or None, "Cost Low": _r(cl),
                     "Cost Most Likely": _r(cm), "Cost High": _r(ch),
                     "Cost Distribution": cd, "Burn Rate": a.get("burn_rate")})
    risks = []
    for r in spec.get("risks", []):
        delay = r.get("delay", 0.0)
        if isinstance(delay, dict):
            lo, ml, hi, dd = _trio(delay)
        else:
            lo, ml, hi, dd = None, delay or None, None, ""
        risks.append({"Risk": r["name"], "Probability": r["probability"],
                      "Activities": "; ".join(r.get("activities", [])),
                      "Delay Low": lo, "Delay Most Likely": ml, "Delay High": hi,
                      "Delay Distribution": dd, "Cost": r.get("cost")})
    return {"Activities": pd.DataFrame(rows),
            "Risks": pd.DataFrame(risks, columns=["Risk", "Probability", "Activities",
                                                  "Delay Low", "Delay Most Likely",
                                                  "Delay High", "Delay Distribution", "Cost"])}


# ---------------------------------------------------------------------- AoA
def _read_aoa(sheets: Dict[str, pd.DataFrame]) -> Dict[str, Any]:
    spec = _read_settings("aoa", sheets.get("settings"))
    alts = _Sheet("Alternatives", sheets.get("alternatives"))
    lines = _Sheet("Lines", sheets.get("lines"))
    phased = _Sheet("Phased", sheets.get("phased"))
    if not lines:
        raise SpecError("The Lines sheet has no rows: one cost line per row, for each "
                        "alternative.")
    a_col, l_col = lines.need("Alternative"), lines.need("Line", "Cost Line", "Name")
    p_col = lines.need("Phase")
    by_alt: Dict[str, Dict[str, Any]] = {}
    if alts:
        n_col = alts.need("Alternative", "Name")
        for row, rec in alts.rows():
            name = _text(rec[n_col])
            if not name:
                raise SpecError(f"Alternatives sheet, row {row}: no name.")
            alt: Dict[str, Any] = {"name": name}
            for key, heads in (("effectiveness", ("Effectiveness",)),
                               ("correlation", ("Correlation",))):
                v = _get(alts, rec, row, *heads)
                if v is not None:
                    alt[key] = v
            alt["lines"] = []
            by_alt[name] = alt
    table: Dict[Tuple[str, str], Dict[int, float]] = {}
    if phased:
        pa, pl = phased.need("Alternative"), phased.need("Line", "Cost Line", "Name")
        years = phased.years()
        if not years:
            raise SpecError("The Phased sheet needs one column per year, headed 2027, "
                            "2028 and so on.")
        for row, rec in phased.rows():
            key = (_text(rec[pa]), _text(rec[pl]))
            table[key] = {y: v for y, c in years.items()
                          if (v := _num(rec[c], "Phased", row, str(y))) is not None}
    for row, rec in lines.rows():
        alt_name, line = _text(rec[a_col]), _text(rec[l_col])
        where = f"Lines sheet, row {row} ({alt_name} / {line})"
        if not alt_name or not line:
            raise SpecError(f"Lines sheet, row {row}: give the Alternative and the Line.")
        if alt_name not in by_alt:
            if alts:
                raise SpecError(f"{where}: {alt_name!r} is not on the Alternatives sheet.")
            by_alt[alt_name] = {"name": alt_name, "lines": []}
        d: Dict[str, Any] = {"name": line, "phase": _text(rec[p_col])}
        g = lambda *n: _get(lines, rec, row, *n)  # noqa: E731
        total, amount = g("Total"), g("Annual Amount", "Amount")
        if total is not None:
            start, years = g("Start Year", "Start"), g("Years")
            if start is None or years is None:
                raise SpecError(f"{where}: a Total needs a Start Year and a number of Years.")
            d["spread"] = {"total": total, "start": int(start), "years": int(years)}
            prof = _get(lines, rec, row, "Profile", number=False)
            if prof:
                d["spread"]["profile"] = prof
        elif amount is not None:
            first, last = g("First Year"), g("Last Year")
            if first is None or last is None:
                raise SpecError(f"{where}: an Annual Amount needs a First Year and a Last "
                                "Year.")
            d["annual"] = {"amount": amount, "first": int(first), "last": int(last)}
        elif (alt_name, line) in table:
            d["by_year"] = {str(y): v for y, v in table.pop((alt_name, line)).items()}
        else:
            raise SpecError(f"{where}: no cost. Give a Total (with Start Year and Years), an "
                            "Annual Amount (with First and Last Year), or a row on the Phased "
                            "sheet.")
        u = _dist(g("Low Factor", "Low"), g("Most Likely Factor", "Most Likely"),
                  g("High Factor", "High"),
                  _get(lines, rec, row, "Distribution", number=False, default=""), where)
        if u:
            d["uncertainty"] = u
        by_alt[alt_name]["lines"].append(d)
    if table:
        (a, l), = list(table)[:1]
        raise SpecError(f"The Phased sheet has a row for {a} / {l} that is not on the Lines "
                        "sheet.")
    empty = [n for n, a in by_alt.items() if not a["lines"]]
    if empty:
        raise SpecError(f"{', '.join(map(repr, empty))} has no cost lines on the Lines sheet.")
    spec["alternatives"] = list(by_alt.values())
    return spec


def _aoa_frames(spec) -> Dict[str, pd.DataFrame]:
    alts, lines, phased = [], [], []
    for a in spec["alternatives"]:
        alts.append({"Alternative": a["name"], "Effectiveness": a.get("effectiveness"),
                     "Correlation": a.get("correlation")})
        for ln in a["lines"]:
            lo, ml, hi, dist = _trio(ln.get("uncertainty"))
            row = {"Alternative": a["name"], "Line": ln["name"], "Phase": ln["phase"],
                   "Total": None, "Start Year": None, "Years": None, "Profile": None,
                   "Annual Amount": None, "First Year": None, "Last Year": None,
                   "Low Factor": lo, "Most Likely Factor": ml, "High Factor": hi,
                   "Distribution": dist}
            if "spread" in ln:
                s = ln["spread"]
                row.update({"Total": s["total"], "Start Year": s["start"],
                            "Years": s["years"], "Profile": s.get("profile")})
            elif "annual" in ln:
                s = ln["annual"]
                row.update({"Annual Amount": s["amount"], "First Year": s["first"],
                            "Last Year": s["last"]})
            else:
                phased.append({"Alternative": a["name"], "Line": ln["name"]}
                              | {int(y): v for y, v in ln["by_year"].items()})
            lines.append(row)
    ph = pd.DataFrame(phased) if phased else pd.DataFrame(columns=["Alternative", "Line"])
    return {"Alternatives": pd.DataFrame(alts), "Lines": pd.DataFrame(lines), "Phased": ph}


# ---------------------------------------------------------------- portfolio
def _read_portfolio(sheets: Dict[str, pd.DataFrame]) -> Dict[str, Any]:
    raw = _read_settings("portfolio", sheets.get("settings"))
    spec = {k: v for k, v in raw.items() if not k.startswith("growth_") or
            k == "growth_correlation"}
    g3 = [raw.get(k) for k in ("growth_low", "growth_mode", "growth_high")]
    if any(v is not None for v in g3):
        spec["growth"] = _dist(*g3, raw.get("growth_dist", ""), "Settings sheet (growth)")
    if "frontier_scales" in spec:
        try:
            spec["frontier_scales"] = [float(t) for t in _items(spec["frontier_scales"])]
        except ValueError:
            raise SpecError("Settings sheet: Frontier Scales is a list of numbers separated "
                            "by ; (for example 0.9; 1.0; 1.1).") from None
    budget = _Sheet("Budget", sheets.get("budget"))
    if not budget:
        raise SpecError("The Budget sheet has no rows: one row per year, Year and Budget.")
    y_col, b_col = budget.need("Year"), budget.need("Budget", "Amount")
    spec["budget"] = {}
    for row, rec in budget.rows():
        y = _num(rec[y_col], "Budget", row, "year")
        v = _num(rec[b_col], "Budget", row, "budget")
        if y is None or v is None:
            raise SpecError(f"Budget sheet, row {row}: give both the Year and the Budget.")
        spec["budget"][str(int(y))] = v
    cands = _Sheet("Candidates", sheets.get("candidates"))
    if not cands:
        raise SpecError("The Candidates sheet has no rows: one row per funding option.")
    c_col, o_col = cands.need("Candidate", "Program"), cands.need("Option")
    v_col = cands.need("Value")
    years = cands.years()
    if not years:
        raise SpecError("The Candidates sheet needs one cost column per year, headed 2027, "
                        "2028 and so on.")
    by_name: Dict[str, Dict[str, Any]] = {}
    for row, rec in cands.rows():
        name = _text(rec[c_col])
        where = f"Candidates sheet, row {row} ({name})"
        if not name:
            raise SpecError(f"Candidates sheet, row {row}: no candidate name.")
        c = by_name.setdefault(name, {"name": name, "options": []})
        cat = _get(cands, rec, row, "Category", number=False)
        if cat:
            c["category"] = cat
        mand = _get(cands, rec, row, "Mandatory", number=False)
        if mand:
            if mand.lower() not in ("yes", "y", "true", "1", "x", "no", "n", "false", "0"):
                raise SpecError(f"{where}: Mandatory is yes or no, not {mand!r}.")
            if mand.lower() in ("yes", "y", "true", "1", "x"):
                c["mandatory"] = True
        req = _items(rec.get(cands.has("Requires") or ""))
        if req:
            c["requires"] = req
        opt = _text(rec[o_col])
        if not opt:
            raise SpecError(f"{where}: the Option has no name (use Full if there is only one).")
        value = _num(rec[v_col], "Candidates", row, "value")
        if value is None:
            raise SpecError(f"{where}: no Value.")
        cost = {str(y): v for y, col in years.items()
                if (v := _num(rec[col], "Candidates", row, str(y))) is not None}
        c["options"].append({"name": opt, "value": value, "cost_by_year": cost})
    for c in by_name.values():
        for r in c.get("requires", []):
            if r not in by_name:
                raise SpecError(f"Candidates sheet: {c['name']!r} requires {r!r}, which is "
                                "not a candidate.")
    spec["candidates"] = list(by_name.values())
    ex = _Sheet("Exclusive", sheets.get("exclusive"))
    if ex:
        col = ex.need("Candidates")
        spec["exclusive"] = []
        for row, rec in ex.rows():
            group = _items(rec[col])
            bad = [g for g in group if g not in by_name]
            if bad:
                raise SpecError(f"Exclusive sheet, row {row}: {', '.join(map(repr, bad))} is "
                                "not a candidate.")
            if len(group) > 1:
                spec["exclusive"].append(group)
    return spec


def _portfolio_frames(spec) -> Dict[str, pd.DataFrame]:
    years = sorted({int(y) for c in spec["candidates"] for o in c["options"]
                    for y in o["cost_by_year"]} | {int(y) for y in spec["budget"]})
    rows = []
    for c in spec["candidates"]:
        for k, o in enumerate(c["options"]):
            row = {"Candidate": c["name"], "Category": c.get("category") if k == 0 else None,
                   "Mandatory": ("yes" if c.get("mandatory") else None) if k == 0 else None,
                   "Requires": "; ".join(c.get("requires", [])) if k == 0 else None,
                   "Option": o["name"], "Value": o["value"]}
            row |= {y: o["cost_by_year"].get(str(y)) for y in years}
            rows.append(row)
    budget = pd.DataFrame([(int(y), v) for y, v in spec["budget"].items()],
                          columns=["Year", "Budget"])
    ex = pd.DataFrame([("; ".join(g),) for g in spec.get("exclusive", [])],
                      columns=["Candidates"])
    return {"Candidates": pd.DataFrame(rows), "Budget": budget, "Exclusive": ex}


def _portfolio_settings(spec) -> Dict[str, Any]:
    out = {k: v for k, v in spec.items() if k not in ("growth", "frontier_scales")}
    if spec.get("growth"):
        lo, ml, hi, dist = _trio(spec["growth"])
        out |= {"growth_low": lo, "growth_mode": ml, "growth_high": hi}
        if dist:
            out["growth_dist"] = dist
    if spec.get("frontier_scales"):
        out["frontier_scales"] = "; ".join(f"{v:g}" for v in spec["frontier_scales"])
    return out


# ----------------------------------------------------------- read and write
INSTRUCTIONS = {
    "jcl": [
        ("What this is", "A project network with its uncertainty, for a joint cost and "
         "schedule (JCL) analysis: ce-core jcl --spec <this file>."),
        ("Activities sheet", "One row per activity. ID is a short name the other rows refer "
         "to. Duration is in months; Duration Low, Most Likely and High are the range you "
         "believe (leave them blank for none). Predecessors lists the activities it waits "
         "for, separated by ;, each optionally with a link type and a lag in months: "
         "design; fsw -4; bus SS +2."),
        ("Costs", "Fixed Cost is spent whatever the duration, with Cost Low, Most Likely and "
         "High as its range. Burn Rate is spent per month the activity runs, so it grows "
         "when the activity slips."),
        ("Risks sheet", "Events that may happen. Probability is the chance (0.3 or 30%), "
         "Activities the IDs it delays (separated by ;), Delay the months it adds (one "
         "number, or Low, Most Likely and High), Cost what it adds besides the delay."),
        ("Settings sheet", "Standing Army is the cost per month of the whole project "
         "running. The correlations say how much durations, and costs, move together. "
         "Confidence is the joint level to report against (0.7 for 70%)."),
        ("Distribution columns", "Optional: triangular (the default), pert or uniform."),
    ],
    "aoa": [
        ("What this is", "The alternatives for a capability and their costs over the life "
         "cycle, for an analysis of alternatives: ce-core aoa --spec <this file>."),
        ("Alternatives sheet", "One row per alternative. Effectiveness is optional, on any "
         "scale that is the same for all of them. Correlation (default 0.3) is how much its "
         "cost lines overrun together."),
        ("Lines sheet", "One cost line per row, naming its Alternative. Give its cost one of "
         "three ways: a Total spread over Years from a Start Year (Profile uniform, or "
         "bell); an Annual Amount from a First Year to a Last Year; or a row on the Phased "
         "sheet with the amount in each year."),
        ("Uncertainty", "Low, Most Likely and High Factor multiply the line: 0.9, 1.0, 1.4 "
         "means 10% under to 40% over. Leave them blank for a line with no uncertainty."),
        ("Settings sheet", "Base Year is the constant-dollar year. Discount Rate and "
         "Inflation Rate are annual (0.02 for 2%). Basis is by, ty or pv: which dollars to "
         "compare on."),
    ],
    "portfolio": [
        ("What this is", "Candidate programs, their funding options and the budget, for "
         "choosing what to fund: ce-core portfolio --spec <this file>."),
        ("Candidates sheet", "One row per funding option. A candidate with several options "
         "(full rate, minimum rate, defer) takes one row each; at most one is funded. Value "
         "is the benefit score. The year columns hold the cost in each year. Category, "
         "Mandatory (yes to always fund it) and Requires (candidates it depends on, "
         "separated by ;) go on its first row."),
        ("Budget sheet", "The money available in each year."),
        ("Exclusive sheet", "Candidates only one of which may be funded, separated by ; on "
         "a row."),
        ("Settings sheet", "Growth Low, Most Likely and High are the cost growth factors "
         "for the risk analysis (0.92, 1.0, 1.45). Delta is the budget step for the "
         "trade-off curve."),
    ],
}

_FRAMES = {"jcl": _jcl_frames, "aoa": _aoa_frames, "portfolio": _portfolio_frames}
_READERS = {"jcl": _read_jcl, "aoa": _read_aoa, "portfolio": _read_portfolio}


def kind_of(sheets: Sequence[str]) -> str:
    names = {_norm(s) for s in sheets}
    if "activities" in names:
        return "jcl"
    if "lines" in names or "alternatives" in names:
        return "aoa"
    if "candidates" in names:
        return "portfolio"
    raise SpecError("This workbook has no Activities (JCL), Lines (AoA) or Candidates "
                    "(portfolio) sheet; ce-core template jcl, aoa or portfolio writes one.")


def read_workbook(path, kind: Optional[str] = None) -> Dict[str, Any]:
    """A JCL, AoA or portfolio workbook as its spec dictionary."""
    path = Path(path)
    raw = open_workbook(path, SpecError)
    sheets = {_norm(k): v for k, v in raw.items()}
    kind = kind or kind_of(raw)
    return _READERS[kind](sheets)


def write_workbook(spec: Dict[str, Any], kind: str, out) -> Path:
    """Write a spec dictionary as a workbook :func:`read_workbook` reads back."""
    from openpyxl.styles import Alignment, Font

    out = Path(out)
    frames = _FRAMES[kind](spec)
    settings = _portfolio_settings(spec) if kind == "portfolio" else spec
    frames["Settings"] = _settings_frame(kind, settings)
    frames["Instructions"] = pd.DataFrame(INSTRUCTIONS[kind],
                                          columns=["Sheet or topic", "What to put there"])
    with pd.ExcelWriter(out, engine="openpyxl") as xl:
        for name, frame in frames.items():
            frame.to_excel(xl, sheet_name=name, index=False)
            ws = xl.sheets[name]
            ws.freeze_panes = "A2"
            for cell in ws[1]:
                cell.font = Font(bold=True)
            for j, col in enumerate(frame.columns, start=1):
                width = max([len(str(col))] + [len(str(v)) for v in frame[col].dropna()])
                ws.column_dimensions[ws.cell(row=1, column=j).column_letter].width = \
                    min(max(width + 2, 9), 48)
        notes = xl.sheets["Instructions"]
        notes.column_dimensions["B"].width = 100
        for row in notes.iter_rows(min_row=2):
            row[0].alignment = Alignment(vertical="top")
            row[1].alignment = Alignment(wrap_text=True, vertical="top")
    return out
