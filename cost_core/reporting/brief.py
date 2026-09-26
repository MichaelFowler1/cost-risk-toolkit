# Copyright 2026 Michael Fowler
# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0

"""
brief.py - A short PowerPoint briefing from each result.

Much of a cost analyst's week goes on turning numbers into slides, and the
slides are what decision makers see. So every EVM, JCL, schedule check, AoA
and portfolio run also writes ``brief.pptx``: a handful of 16:9 slides that
open, as a briefing does, with the bottom line, then the chart, the numbers
behind it, and the assumptions, which a reviewer will ask for. The words are
the same plain reading the command prints, so the slides say nothing the
analysis does not.

Needs python-pptx (MIT), which comes with the ``plots`` extra along with the
charts the slides carry.
"""

from __future__ import annotations

from datetime import date
from pathlib import Path
from typing import Optional, Sequence, Tuple

import numpy as np
import pandas as pd


def _pptx():
    try:
        from pptx import Presentation
        from pptx.util import Inches, Pt
    except ImportError as exc:  # pragma: no cover - depends on the install
        raise ImportError(
            "The PowerPoint briefing needs python-pptx, which comes with the plots "
            "extra: pip install \"cost-core[plots]\"") from exc
    return Presentation, Inches, Pt


NAVY = (0x1F, 0x4E, 0x79)
GREY = (0x59, 0x59, 0x59)

#: The template's layouts are chosen by these names, as PowerPoint's own
#: layouts and most house templates call them.
TITLE_LAYOUT, HEADING_LAYOUT, BLANK_LAYOUT = "title slide", "title only", "blank"


def _open_template(path: Path):
    """A .pptx as it is; a .potx (a template) as the presentation it describes.
    python-pptx refuses the template's content type, which is the only
    difference between the two."""
    if path.suffix.lower() != ".potx":
        return str(path)
    import io
    import zipfile

    src = zipfile.ZipFile(path)
    buf = io.BytesIO()
    with zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as out:
        for item in src.infolist():
            data = src.read(item.filename)
            if item.filename == "[Content_Types].xml":
                data = data.replace(b"presentationml.template.main+xml",
                                    b"presentationml.presentation.main+xml")
            out.writestr(item, data)
    buf.seek(0)
    return buf


def _drop_slides(prs) -> None:
    """A template's own sample slides don't belong in the briefing."""
    ids = prs.slides._sldIdLst
    for sld in list(ids):
        prs.part.drop_rel(sld.rId)
        ids.remove(sld)


def _layout(prs, name: str):
    for layout in prs.slide_layouts:
        if layout.name.strip().lower() == name:
            return layout
    return None


class Brief:
    """A deck built slide by slide: 16:9 by default, or on the slide master
    set with ``cost_core.reporting.output.configure(slide_template=...)``."""

    def __init__(self, title: str, subtitle: str = ""):
        from cost_core.reporting import output

        Presentation, Inches, Pt = _pptx()
        self.Inches, self.Pt = Inches, Pt
        template = output.slide_template()
        self.notes = []
        if template is None:
            self.prs = Presentation()
            self.prs.slide_width, self.prs.slide_height = Inches(13.333), Inches(7.5)
            self.layouts = {BLANK_LAYOUT: self.prs.slide_layouts[6]}
        else:
            self.prs = Presentation(_open_template(template))
            _drop_slides(self.prs)
            self.layouts = {n: _layout(self.prs, n)
                            for n in (TITLE_LAYOUT, HEADING_LAYOUT, BLANK_LAYOUT)}
            if self.layouts[BLANK_LAYOUT] is None:
                # The fallback: the layout with the fewest placeholders to fill.
                self.layouts[BLANK_LAYOUT] = min(self.prs.slide_layouts,
                                                 key=lambda lay: len(lay.placeholders))
            missing = [n for n in (TITLE_LAYOUT, HEADING_LAYOUT) if self.layouts[n] is None]
            if missing:
                import logging

                self.notes.append(f"{template.name} has no {' or '.join(map(repr, missing))} "
                                  f"layout, so those slides use "
                                  f"{self.layouts[BLANK_LAYOUT].name!r} with cost-core's "
                                  "own heading.")
                logging.getLogger(__name__).warning(self.notes[-1])
        # Every position below is laid out on a 13.333 x 7.5 inch slide and
        # scaled to the one in use, so a 4:3 house template still fits.
        self.sx = self.prs.slide_width / Inches(13.333)
        self.sy = self.prs.slide_height / Inches(7.5)
        self._title_slide(title, subtitle)

    def X(self, inches: float) -> int:
        return int(self.Inches(inches) * self.sx)

    def Y(self, inches: float) -> int:
        return int(self.Inches(inches) * self.sy)

    def _rgb(self, rgb):
        from pptx.dml.color import RGBColor
        return RGBColor(*rgb)

    def _textbox(self, slide, left, top, width, height, text, size=18, bold=False,
                 color=None, italic=False):
        box = slide.shapes.add_textbox(self.X(left), self.Y(top), self.X(width),
                                       self.Y(height))
        tf = box.text_frame
        tf.word_wrap = True
        p = tf.paragraphs[0]
        p.text = text
        p.font.size, p.font.bold, p.font.italic = self.Pt(size), bold, italic
        if color:
            p.font.color.rgb = self._rgb(color)
        return tf

    def _blank(self, heading: str):
        layout = self.layouts.get(HEADING_LAYOUT)
        if layout is not None:
            slide = self.prs.slides.add_slide(layout)
            if slide.shapes.title is not None:
                slide.shapes.title.text = heading   # in the house style
                return slide
        slide = self.prs.slides.add_slide(self.layouts[BLANK_LAYOUT])
        self._textbox(slide, 0.6, 0.35, 12.1, 0.8, heading, size=28, bold=True, color=NAVY)
        return slide

    def _title_slide(self, title, subtitle):
        layout = self.layouts.get(TITLE_LAYOUT)
        if layout is not None:
            slide = self.prs.slides.add_slide(layout)
            if slide.shapes.title is not None:
                slide.shapes.title.text = title
                others = [p for p in slide.placeholders
                          if p.placeholder_format.idx != 0 and p.has_text_frame]
                if subtitle and others:
                    others[0].text = subtitle
                self._textbox(slide, 0.8, 6.6, 11.7, 0.5,
                              f"Prepared {date.today():%d %B %Y} with cost-core", size=12,
                              color=GREY)
                return
        slide = self.prs.slides.add_slide(self.layouts[BLANK_LAYOUT])
        self._textbox(slide, 0.8, 2.6, 11.7, 1.2, title, size=36, bold=True, color=NAVY)
        if subtitle:
            self._textbox(slide, 0.8, 3.8, 11.7, 0.8, subtitle, size=18, color=GREY)
        self._textbox(slide, 0.8, 6.6, 11.7, 0.5,
                      f"Prepared {date.today():%d %B %Y} with cost-core", size=12, color=GREY)

    def bottom_line(self, lines: Sequence[str], numbers: Sequence[Tuple[str, str]] = ()):
        """The first content slide: the reading, and up to six key numbers."""
        slide = self._blank("Bottom line")
        width = 7.6 if numbers else 12.1
        tf = self._textbox(slide, 0.6, 1.3, width, 5.6, "", size=18)
        for i, line in enumerate(lines):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = "•  " + line
            p.font.size = self.Pt(17 if len(lines) <= 4 else 15)
            p.space_after = self.Pt(10)
        for k, (label, value) in enumerate(numbers[:6]):
            top = 1.3 + k * 0.92
            self._textbox(slide, 8.6, top, 4.2, 0.45, value, size=24, bold=True, color=NAVY)
            self._textbox(slide, 8.6, top + 0.42, 4.2, 0.4, label, size=11, color=GREY)
        return slide

    def image(self, heading: str, path, caption: str = ""):
        slide = self._blank(heading)
        path = Path(path)
        if path.exists():
            pic = slide.shapes.add_picture(str(path), self.X(1.4), self.Y(1.2),
                                           height=self.Y(5.6 if caption else 5.9))
            room = self.X(12.1)
            if pic.width > room:   # a narrow template: fit the width instead
                pic.height = int(pic.height * room / pic.width)
                pic.width = room
            pic.left = int((self.prs.slide_width - pic.width) / 2)
            # Alt text, so a screen reader says what the chart is.
            pic._element.nvPicPr.cNvPr.set("descr", f"{heading}. {caption}".strip(". "))
        else:
            self._textbox(slide, 0.6, 3.0, 12, 1, "(chart not available: install the plots "
                          "extra for charts)", size=16, italic=True, color=GREY)
        if caption:
            self._textbox(slide, 0.6, 6.85, 12.1, 0.5, caption, size=12, italic=True, color=GREY)
        return slide

    def table(self, heading: str, df: pd.DataFrame, formats: Optional[dict] = None,
              note: str = "", max_rows: int = 14, colors: Optional[dict] = None):
        """A table slide; long tables are cut to ``max_rows`` with a note.
        ``colors`` maps a cell's text (such as "FAIL") to its colour."""
        slide = self._blank(heading)
        formats = formats or {}
        shown = df.head(max_rows)
        rows, cols = len(shown) + 1, len(shown.columns)
        height = min(0.37 * rows, 5.4)
        shape = slide.shapes.add_table(rows, cols, self.X(0.6), self.Y(1.3),
                                       self.X(12.1), self.Y(height))
        t = shape.table
        for j, col in enumerate(shown.columns):
            t.cell(0, j).text = str(col)
        for i, row in enumerate(shown.itertuples(index=False), start=1):
            for j, value in enumerate(row):
                t.cell(i, j).text = _fmt(value, formats.get(shown.columns[j]))
        # Columns as wide as their longest entry, within the slide's width.
        lengths = [max([len(str(c))] + [len(t.cell(r, j).text) for r in range(1, rows)])
                   for j, c in enumerate(shown.columns)]
        lengths = [min(max(n, 4), 60) for n in lengths]
        total = sum(lengths)
        for j, n in enumerate(lengths):
            t.columns[j].width = int(self.X(12.1) * n / total)
        size = self.Pt(12 if rows <= 10 else 10)
        for r in range(rows):
            for c in range(cols):
                cell = t.cell(r, c)
                for p in cell.text_frame.paragraphs:
                    p.font.size = size
                    if r and colors and cell.text in colors:
                        p.font.bold = True
                        p.font.color.rgb = self._rgb(colors[cell.text])
        extra = f"{len(df) - max_rows} more rows in report.xlsx. " if len(df) > max_rows else ""
        if extra or note:
            self._textbox(slide, 0.6, 6.85, 12.1, 0.5, extra + note, size=12, italic=True,
                          color=GREY)
        return slide

    def assumptions(self, items: Sequence[str]):
        slide = self._blank("Assumptions and method")
        tf = self._textbox(slide, 0.6, 1.3, 12.1, 5.6, "", size=14)
        for i, item in enumerate(items[:12]):
            p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
            p.text = "•  " + item
            p.font.size = self.Pt(14 if len(items) <= 8 else 12)
            p.space_after = self.Pt(6)
        return slide

    def _mark(self) -> None:
        """The marking, centred at the top and bottom of every slide."""
        from pptx.enum.text import PP_ALIGN

        from cost_core.reporting import output

        text = output.marking()
        if not text:
            return
        height = self.Y(0.3)
        for slide in self.prs.slides:
            for top in (0, self.prs.slide_height - height):
                box = slide.shapes.add_textbox(0, top, self.prs.slide_width, height)
                p = box.text_frame.paragraphs[0]
                p.text = text
                p.alignment = PP_ALIGN.CENTER
                p.font.size, p.font.bold = self.Pt(11), True
                box.name = "Marking"

    def save(self, path) -> Path:
        self._mark()
        path = Path(path)
        path.parent.mkdir(parents=True, exist_ok=True)
        self.prs.save(path)
        return path


def _fmt(value, spec: Optional[str]) -> str:
    if value is None or (isinstance(value, float) and not np.isfinite(value)):
        return ""
    if spec and isinstance(value, (int, float, np.integer, np.floating)):
        return format(float(value), spec)
    if isinstance(value, (float, np.floating)):
        return f"{value:,.3g}" if abs(value) < 1000 else f"{value:,.0f}"
    return str(value)


def _money(v: float, units: str) -> str:
    from cost_core.plain import _money as m
    return m(v, units)


# ------------------------------------------------------------------ topics
def evm_brief(data, fc, out_dir, units: str = "") -> Path:
    from cost_core import plain

    out_dir = Path(out_dir)
    status = str(data.periods[data.status - 1])
    b = Brief(f"Earned value: {data.name}", f"Status as of period {status}")
    m = data.metrics().iloc[-1]
    q = lambda a, p: float(np.quantile(a, p))  # noqa: E731
    b.bottom_line(plain.evm(data, fc, units, where="on the Warning signs slide"), [
        (f"most likely final cost (budget {_money(data.bac, units)})", _money(q(fc.eac, 0.5), units)),
        ("to be 80% sure of the final cost", _money(q(fc.eac, 0.8), units)),
        (f"likely finish period (plan: {data.planned_duration})", f"{q(fc.finish, 0.5):.1f}"),
        ("CPI: work done per dollar spent", f"{m.cpi:.2f}"),
        ("SPI(t): schedule efficiency", f"{m.spi_t:.2f}"),
        ("percent complete", f"{m.pct_complete:.0%}"),
    ])
    b.image("Where the program is heading", out_dir / "evm.png",
            "Baseline, earned value and actual cost to date, with the P50 forecast and the "
            "P10 to P90 range of finish and final cost.")
    flags = data.flags()
    raised = flags[flags["raised"]]
    if len(raised):
        b.table("Warning signs", raised[["flag", "detail"]].rename(
            columns={"flag": "Warning", "detail": "Why"}))
    pct = fc.percentiles((0.2, 0.5, 0.7, 0.8, 0.9))
    pct = pct.assign(confidence=pct["confidence"].map(lambda v: f"P{v * 100:.0f}"),
                     eac=pct["eac"].map(lambda v: _money(v, units)),
                     finish=pct["finish"].map(lambda v: f"{v:.1f}"))
    b.table("Final cost and finish, by confidence",
            pct.rename(columns={"confidence": "Confidence", "eac": "Final cost",
                                "finish": "Finish (period)"}),
            note="Simulated from the program's own period-by-period record.")
    from cost_core.evm.checks import thresholds, variance_breaches

    breaches = variance_breaches(data, **thresholds(data))
    if len(breaches):
        b.table("Accounts owing a variance analysis report", pd.DataFrame({
            "Account": breaches["account"],
            "Cost variance": [f"{_money(v, units)} ({p:+.1f}%)"
                              for v, p in zip(breaches["cv"], breaches["cv_pct"])],
            "Schedule variance": [f"{_money(v, units)} ({p:+.1f}%)"
                                  for v, p in zip(breaches["sv"], breaches["sv_pct"])]}),
            note="Cumulative variance against the thresholds; the Variance reports sheet has "
                 "the detail.")
    if data.accounts:
        rows = sorted(((n, a.metrics().iloc[-1]) for n, a in data.accounts.items()),
                      key=lambda t: t[1].cpi)
        b.table("Control accounts, worst CPI first", pd.DataFrame(
            [{"Account": n, "CPI": f"{r.cpi:.2f}", "SPI(t)": f"{r.spi_t:.2f}",
              "Cost variance": _money(r.cv, units), "% complete": f"{r.pct_complete:.0%}"}
             for n, r in rows]))
    b.assumptions(list(data.notes) + list(fc.notes) + [
        f"{fc.n_iter:,} simulated completions, seed {fc.seed}.",
        "Every figure and its formula is in report.xlsx beside this briefing."])
    return b.save(out_dir / "brief.pptx")


def jcl_brief(result, project, confidence: float, out_dir, units: str = "") -> Path:
    from cost_core import plain

    out_dir = Path(out_dir)
    b = Brief(f"Joint cost and schedule confidence: {project.name}",
              f"{result.n_iter:,} simulated projects")
    cost_p = float(np.quantile(result.cost, confidence))
    fin_p = float(np.quantile(result.finish, confidence))
    b.bottom_line(plain.jcl(result, confidence, units), [
        ("chance of meeting the plan on both", plain._pct(result.point_jcl)),
        (f"cost P{confidence * 100:.0f}", _money(cost_p, units)),
        (f"schedule P{confidence * 100:.0f} (months)", f"{fin_p:.1f}"),
        ("the two together", plain._pct(result.joint(cost_p, fin_p))),
    ])
    b.image("Joint confidence", out_dir / "jcl.png",
            f"Each dot is one simulated project; the line is every budget and date pair that "
            f"reaches {confidence:.0%} together.")
    crit = result.criticality().sort_values("rank_corr_finish", ascending=False).head(8)
    b.table("Where the schedule risk is", crit.assign(
        criticality=crit["criticality"].map(lambda v: f"{v:.0%}"),
        duration_mean=crit["duration_mean"].map(lambda v: f"{v:.1f}"),
        rank_corr_finish=crit["rank_corr_finish"].map(lambda v: f"{v:.2f}"),
        rank_corr_cost=crit["rank_corr_cost"].map(lambda v: f"{v:.2f}")).rename(columns={
            "activity": "Activity", "criticality": "On critical path",
            "duration_mean": "Mean duration", "rank_corr_finish": "Drives finish",
            "rank_corr_cost": "Drives cost"}))
    b.assumptions([f"Standing army {project.standing_army:g} per month.",
                   f"Duration correlation {project.duration_correlation:g}, cost correlation "
                   f"{project.cost_correlation:g}.",
                   f"Risks: {', '.join(r.name for r in project.risks) or 'none'}."]
                  + list(result.notes) + ["Every table is in report.xlsx beside this briefing."])
    return b.save(out_dir / "brief.pptx")


def dcma_brief(result, schedule, out_dir) -> Path:
    from cost_core import plain

    b = Brief(f"Schedule health (DCMA 14-point): {schedule.name}",
              f"{len(schedule.detail)} tasks, {len(schedule.links)} links")
    b.bottom_line(plain.dcma(result, schedule), [
        ("checks passed", f"{result.passed} of 14"), ("checks failed", str(result.failed))])
    t = result.table
    shown = pd.DataFrame({
        "#": t["check"], "Check": t["name"],
        "Result": t["passed"].map({True: "pass", False: "FAIL"}).fillna("n/a"),
        "Value": [("" if pd.isna(v) else f"{v:.2f}") for v in t["value"]],
        "Threshold": [f"{v:.2f}" for v in t["threshold"]]})
    b.table("The 14 checks", shown, max_rows=14,
            colors={"FAIL": (0xC0, 0x00, 0x00), "pass": (0x2E, 0x7D, 0x32)})
    # Those that make the schedule's dates untrustworthy first: a broken
    # critical path, missing logic, hard constraints, negative float, leads.
    first = [12, 1, 5, 7, 2]
    order = first + [k for k in sorted(result.tasks) if k not in first]
    worst = [(k, result.tasks[k]) for k in order if result.tasks.get(k)][:6]
    if worst:
        b.table("Tasks to fix first", pd.DataFrame(
            [{"Check": t.loc[t["check"] == k, "name"].iloc[0],
              "Tasks": ", ".join(v[:6]) + (f" and {len(v) - 6} more" if len(v) > 6 else "")}
             for k, v in worst]))
    b.assumptions(list(schedule.notes) + ["Every failing task is listed in report.xlsx."])
    return b.save(Path(out_dir) / "brief.pptx")


def cost_risk_brief(result, out_dir) -> Path:
    from cost_core import plain

    units = plain.units_label(result.units)
    sim = result.sim
    out_dir = Path(out_dir)
    b = Brief("Cost risk analysis", f"{result.inputs.model.name}, {sim.n_iter:,} simulations")
    b.bottom_line(plain.cost_risk(result, units),
                  [("point estimate", _money(result.point_estimate, units)),
                   ("its confidence", f"{sim.point_estimate_percentile:.0f}%"),
                   ("P50", _money(sim.p50, units)), ("P80", _money(sim.p80, units))])
    b.image("How sure: the S-curve", out_dir / "cost_risk_s_curve.png")
    b.image("What drives the uncertainty", out_dir / "cost_risk_drivers.png")
    alloc = result.allocation(0.8)
    b.table("Where the P80 reserve goes", pd.DataFrame({
        "Element or risk": alloc["component"],
        "Point estimate": alloc["point_estimate"].map(lambda v: _money(v, units)),
        "Share of the P80": alloc["allocated_p80"].map(lambda v: _money(v, units)),
        "Reserve": alloc["reserve"].map(lambda v: _money(v, units)),
        "Of the reserve": alloc["share_of_reserve"].map(lambda v: f"{v:.0%}")}),
        note="Shares add up to the P80 of the total. Each element's own P80 would add up to "
             "more: percentiles don't add.")
    conf = result.confidence_table((50, 70, 80, 90))
    b.table("Cost at each confidence level", pd.DataFrame({
        "Confidence": conf["confidence"].map(lambda v: f"{v:.0%}"),
        "Cost": conf["cost"].map(lambda v: _money(v, units)),
        "Reserve over the point estimate": [f"{_money(r, units)} ({p:.0%})"
                                            for r, p in zip(conf["reserve"], conf["reserve_pct"])]}))
    b.assumptions([f"{k}: {v}" for k, v in result.assumptions.items()]
                  + ["The full tables are in report.xlsx beside this briefing."])
    return b.save(out_dir / "brief.pptx")


def aoa_brief(result, out_dir) -> Path:
    from cost_core import plain

    units = str(result.assumptions.get("units") or "")
    b = Brief("Analysis of alternatives: life-cycle cost",
              f"{result.assumptions.get('basis', '').upper()} dollars")
    s = result.summary.sort_values("p50")
    b.bottom_line(plain.aoa(result, units),
                  [(f"{r.alternative}: most likely", _money(r.p50, units)) for r in s.itertuples()])
    b.image("Life-cycle cost S-curves", Path(out_dir) / "aoa_s_curves.png")
    cols = {"alternative": "Alternative", "p50": "P50", "p80": "P80",
            "p_cheapest": "Chance cheapest", "dominated_by": "Dominated by"}
    shown = s[[c for c in cols if c in s.columns]].rename(columns=cols)
    for c in ("P50", "P80"):
        shown[c] = shown[c].map(lambda v: _money(v, units))
    shown["Chance cheapest"] = shown["Chance cheapest"].map(lambda v: f"{v:.0%}")
    b.table("The alternatives", shown.fillna(""))
    b.assumptions([f"{k}: {v}" for k, v in result.assumptions.items()])
    return b.save(Path(out_dir) / "brief.pptx")


def portfolio_brief(result, portfolio, out_dir, units: str = "", risk=None) -> Path:
    from cost_core import plain

    b = Brief("Portfolio: which programs to fund", f"{len(portfolio.candidates)} candidates")
    b.bottom_line(plain.portfolio(result, len(portfolio.candidates), risk, units),
                  [("programs funded", f"{len(result.funded)} of {len(portfolio.candidates)}"),
                   ("total value", f"{result.value:,.1f}")])
    sel = result.selected
    b.table("The recommended portfolio", pd.DataFrame({
        "Program": sel["candidate"], "Option": sel["option"].fillna("not funded"),
        "Value": [f"{v:,.1f}" if o else "" for v, o in zip(sel["value"], sel["option"].notna())],
        "Cost": [_money(c, units) if o else "" for c, o in zip(sel["cost"], sel["option"].notna())]}))
    if risk is not None and len(risk):
        b.table("Chance each year goes over budget", pd.DataFrame({
            "Year": risk["year"].astype(int), "Budget": risk["budget"].map(lambda v: _money(v, units)),
            "Planned": risk["planned"].map(lambda v: _money(v, units)),
            "Chance over": risk["p_over_budget"].map(lambda v: f"{v:.0%}")}))
    b.assumptions(["Value scores and costs as given in the spec.",
                   "The full tables are in report.xlsx beside this briefing."])
    return b.save(Path(out_dir) / "brief.pptx")
