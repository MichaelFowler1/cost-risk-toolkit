# Copyright 2026 Michael Fowler
# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0

"""The programme roll-up workbook.

Programme first, then every element behind it: the headline sheets, the
annual funding profile, the S-curve and tornado when the risk roll-up ran,
the buy sensitivity when it was run, and one projections sheet per element.
Ported unchanged from the desktop tool's WBS module; the roll-up itself lives
in :mod:`cost_core.program.rollup` and is imported lazily so this writer can
be imported on its own.
"""

from __future__ import annotations

from typing import TYPE_CHECKING

import pandas as pd

if TYPE_CHECKING:
    from cost_core.program.rollup import ProgramResult


def save_program_workbook(
    filename: str,
    result: ProgramResult,
    sensitivity: pd.DataFrame | None = None,
):
    """Write the roll-up: programme first, then every element behind it.

    The order is deliberate. The programme total is what gets briefed, but a
    reviewer's first question is which element it came from, so the element
    breakdown sits immediately behind the headline rather than at the end.
    """
    import openpyxl
    from openpyxl.chart import Reference, ScatterChart, Series

    from cost_core.reporting.lot_workbook import (
        _format_chart,
        _money_axis_fmt,
        _money_short,
        _nice_bounds,
    )
    from cost_core.program.rollup import (
        by_fiscal_year,
        element_summary,
        influence_table,
        program_summary,
    )

    element_table = element_summary(result)
    with pd.ExcelWriter(filename, engine="openpyxl") as writer:
        program_summary(result).to_excel(
            writer, sheet_name="Program_Summary", index=False
        )
        element_table.to_excel(
            writer, sheet_name="Program_Elements", index=False
        )
        result.by_lot.to_excel(
            writer, sheet_name="Program_By_Lot", index=False
        )
        annual_sheet = by_fiscal_year(result)
        if len(annual_sheet):
            annual_sheet.to_excel(
                writer, sheet_name="Program_By_FY", index=False
            )
        if result.scurve is not None:
            result.scurve.to_excel(
                writer, sheet_name="Program_SCurve", index=False
            )
        if result.tornado is not None and len(result.tornado):
            result.tornado.to_excel(
                writer, sheet_name="Program_Tornado", index=False
            )
        if sensitivity is not None and len(sensitivity):
            sensitivity.to_excel(
                writer, sheet_name="Buy_Sensitivity", index=False
            )
        infl = influence_table(result)
        if infl is not None:
            infl.to_excel(
                writer, sheet_name="Element_Influence", index=False
            )
        used = {
            "Program_Summary", "Program_Elements", "Program_By_Lot",
            "Program_SCurve", "Program_Tornado", "Buy_Sensitivity",
            "Element_Influence", "Program_By_FY",
        }
        for r in result.elements:
            r.projections.to_excel(
                writer, sheet_name=_sheet_name(r.name, used), index=False
            )

    wb = openpyxl.load_workbook(filename)

    # Annual funding, stacked so the element mix within each year shows.
    # Drawn from the by-fiscal-year table rather than the lot table, because
    # two lots in one year are one year of funding, not two bars.
    annual = by_fiscal_year(result)
    if len(annual) and "Program_By_FY" in wb.sheetnames:
        from openpyxl.chart import BarChart

        ws = wb["Program_By_FY"]
        last = len(annual) + 1
        cols = list(annual.columns)
        bar = BarChart()
        bar.type = "col"
        bar.grouping = "stacked"
        bar.overlap = 100
        _format_chart(
            bar, f"{result.program}: funding by fiscal year",
            "Fiscal Year", "Cost ($)", 20, 11,
        )
        # A fiscal year is a label, not a quantity: the shared chart format
        # would otherwise render 2028 as "2,028".
        bar.x_axis.numFmt = "0"
        bar.y_axis.numFmt = _money_axis_fmt(
            0.0, float(annual["Total ($)"].max())
        )
        for r in result.elements:
            bar.series.append(
                Series(
                    Reference(ws, min_col=cols.index(r.name) + 1, min_row=1,
                              max_row=last),
                    title_from_data=True,
                )
            )
        bar.set_categories(Reference(ws, min_col=1, min_row=2, max_row=last))
        ws.add_chart(bar, f"{_col_letter(len(cols) + 2)}2")

    if result.tornado is not None and len(result.tornado):
        from openpyxl.chart import BarChart

        wst = wb["Program_Tornado"]
        cols = list(result.tornado.columns)
        share_col = cols.index("variance_share") + 1
        last_t = len(result.tornado) + 1
        tor = BarChart()
        tor.type = "bar"          # horizontal, which is what makes it a tornado
        # On a horizontal bar chart openpyxl's x_axis is still the category
        # axis, drawn down the side, and y_axis is the value axis along the
        # bottom. Naming them the other way round leaves the shares reading
        # as "0" under the shared numeric format.
        _format_chart(
            tor, f"{result.program}: share of program variance",
            "WBS element", "Share of variance", 18, 10,
        )
        tor.y_axis.numFmt = "0%"
        tor.legend = None
        tor.series.append(
            Series(
                Reference(wst, min_col=share_col, min_row=1, max_row=last_t),
                title_from_data=True,
            )
        )
        tor.set_categories(
            Reference(wst, min_col=1, min_row=2, max_row=last_t)
        )
        wst.add_chart(tor, f"{_col_letter(len(cols) + 2)}2")

    if sensitivity is not None and len(sensitivity):
        wss = wb["Buy_Sensitivity"]
        cols = list(sensitivity.columns)
        last_s = len(sensitivity) + 1
        unit_col = cols.index("Cost per Unit ($)") + 1
        chart = ScatterChart()
        _format_chart(
            chart, f"{result.program}: unit cost against buy size",
            "Buy multiplier", "Cost per unit ($)", 18, 10,
        )
        chart.x_axis.numFmt = "0.0"
        chart.legend = None
        lo, hi = _nice_bounds(
            float(sensitivity["Cost per Unit ($)"].min()),
            float(sensitivity["Cost per Unit ($)"].max()),
        )
        chart.y_axis.numFmt = _money_axis_fmt(lo, hi)
        chart.y_axis.scaling.min = max(0.0, lo)
        chart.y_axis.scaling.max = hi
        s_unit = Series(
            values=Reference(wss, min_col=unit_col, min_row=1, max_row=last_s),
            xvalues=Reference(wss, min_col=1, min_row=2, max_row=last_s),
            title_from_data=True,
        )
        s_unit.marker.symbol = "circle"
        s_unit.marker.size = 8
        s_unit.smooth = False
        chart.series.append(s_unit)
        wss.add_chart(chart, f"{_col_letter(len(cols) + 2)}2")

    if result.scurve is not None and len(result.scurve):
        _add_program_scurve(wb, result, _format_chart, _money_short,
                            _nice_bounds, _money_axis_fmt, ScatterChart,
                            Series, Reference)

    wb.save(filename)


def _col_letter(idx: int) -> str:
    from openpyxl.utils import get_column_letter

    return get_column_letter(idx)


def _sheet_name(name: str, taken: set[str] | None = None) -> str:
    """A legal, unique Excel sheet name for this element.

    Excel refuses several characters and truncates at 31, which is short
    enough that two real WBS names can collide: "1.1 Air Vehicle Structural
    Assembly Group A" and "... Group B" share their first 31 characters. Left
    alone that either errors or drops an element's sheet, so a numeric suffix
    is added when the truncated name is already spoken for.
    """
    from cost_core.program.rollup import ProgramError

    cleaned = "".join(c for c in name if c not in r"[]:*?/\\").strip()
    base = (cleaned[:31] or "Element").strip()
    if taken is None or base not in taken:
        if taken is not None:
            taken.add(base)
        return base
    for n in range(2, 100):
        suffix = f" ({n})"
        candidate = (cleaned[: 31 - len(suffix)]).strip() + suffix
        if candidate not in taken:
            taken.add(candidate)
            return candidate
    raise ProgramError(
        f"Cannot make a unique sheet name for {name!r}; shorten the element "
        "names so their first 31 characters differ."
    )


def _add_program_scurve(
    wb, result, _format_chart, _money_short, _nice_bounds,
    _money_axis_fmt, ScatterChart, Series, Reference,
):
    ws = wb["Program_SCurve"]
    last = len(result.scurve) + 1
    for row in range(2, last + 1):
        ws.cell(row=row, column=1).number_format = "0%"
        ws.cell(row=row, column=2).number_format = "#,##0"

    curve = ScatterChart()
    _format_chart(
        curve,
        f"{result.program}: probability the program comes in at or below",
        "Program Total ($)", "Cumulative Probability", 20, 11,
    )
    curve.y_axis.numFmt = "0%"
    curve.y_axis.scaling.min = 0
    curve.y_axis.scaling.max = 1

    lo = float(result.scurve["Program Total ($)"].min())
    hi = float(result.scurve["Program Total ($)"].max())
    curve.x_axis.numFmt = _money_axis_fmt(lo, hi)
    nice_lo, nice_hi = _nice_bounds(lo, hi)
    curve.x_axis.scaling.min = max(0.0, nice_lo)
    curve.x_axis.scaling.max = nice_hi

    from openpyxl.chart.series import SeriesLabel

    line = Series(
        values=Reference(ws, min_col=1, min_row=1, max_row=last),
        xvalues=Reference(ws, min_col=2, min_row=2, max_row=last),
        title_from_data=True,
    )
    line.marker.symbol = "none"
    line.smooth = True
    line.tx = SeriesLabel(v="Program total distribution")
    curve.series.append(line)

    from openpyxl.chart.shapes import GraphicalProperties

    pct = result.scurve["Percentile"].round(4)
    for i, (label, level, colour, symbol) in enumerate(
        [("P50", 0.50, "0072B2", "circle"), ("P80", 0.80, "D55E00", "diamond")]
    ):
        hit = result.scurve.loc[pct == round(level, 4), "Program Total ($)"]
        if hit.empty:
            continue
        cost = float(hit.iloc[0])
        y_col, x_col = 4 + i * 2, 5 + i * 2
        ws.cell(row=1, column=y_col, value=f"{label}  {_money_short(cost)}")
        ws.cell(row=2, column=y_col, value=level).number_format = "0%"
        ws.cell(row=1, column=x_col, value=f"{label} cost")
        ws.cell(row=2, column=x_col, value=cost).number_format = "#,##0"

        mark = Series(
            values=Reference(ws, min_col=y_col, min_row=1, max_row=2),
            xvalues=Reference(ws, min_col=x_col, min_row=2, max_row=2),
            title_from_data=True,
        )
        mark.marker.symbol = symbol
        mark.marker.size = 11
        style = GraphicalProperties(solidFill=colour)
        style.line.solidFill = "FFFFFF"
        style.line.width = 19050
        mark.marker.graphicalProperties = style
        mark.graphicalProperties.line.noFill = True
        curve.series.append(mark)

    ws.add_chart(curve, "I2")
