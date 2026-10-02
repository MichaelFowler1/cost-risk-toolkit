# Copyright 2026 Michael Fowler
# SPDX-License-Identifier: PolyForm-Noncommercial-1.0.0

"""
gui.py - cost-core in a window, for people who live in Excel, not a terminal.

``ce-core gui`` (or the desktop shortcut that ``ce-core gui --shortcut`` or the
window's Settings make) opens it. The left side lists the jobs in plain words;
the right side walks through three steps: get the template (it opens in
Excel), choose the filled-in workbook, and run. The answer shows in the window:
the key numbers as tiles, the main chart, and what it means in plain English,
with buttons that open the report, the briefing and the results folder.

When something is wrong the window says where: "Check my workbook" lists the
problems it can find without running anything, and a run that stops marks the
cell it stopped on in a copy of the workbook, so fixing it is a click rather
than a row number. A report left open in Excel, which would make the next run
fail, is caught before the run starts.

It is a window onto the same commands, not a second implementation: every run
is ``python -m cost_core <command> ...`` in a child process, exactly what the
command line does, so every check, test and refusal message applies unchanged.
Tkinter ships with Python, so the window adds no dependency, no server and
nothing a lab's allowlisting has to approve beyond Python itself.

The functions above the window class decide everything (the command a task
runs, where results go, what the output said, what the key numbers are, where
a problem is) and are tested without a screen; the class lays out widgets and
calls them.
"""

from __future__ import annotations

import json
import math
import os
import queue
import re
import subprocess
import sys
import threading
from dataclasses import dataclass
from pathlib import Path
from typing import Dict, List, Optional, Sequence, Tuple


# ------------------------------------------------------------------- tasks
@dataclass(frozen=True)
class Task:
    """One job the window offers."""

    key: str            # the template and demo topic
    title: str          # what it's for, in the estimator's words
    command: str        # the ce-core command it runs
    flag: str           # how that command is given the file
    need: str           # what to bring
    get: str            # what comes out
    units: bool = True  # whether the command takes --units
    template: Optional[str] = None  # template topic, if it has one
    demo: Optional[str] = None      # demo topic, if it has one
    filetypes: Tuple[Tuple[str, str], ...] = (("Excel workbook", "*.xlsx"),
                                              ("CSV file", "*.csv"))


XLSX = (("Excel workbook", "*.xlsx"),)
SPEC = (("Excel workbook", "*.xlsx"), ("JSON spec", "*.json"))

TASKS: Tuple[Task, ...] = (
    Task("cost-risk", "How sure is my estimate?", "cost-risk", "--data",
         "Your estimate in Excel: each WBS element with a low, most likely and high, "
         "and any risks that may or may not happen.",
         "The chance the estimate is exceeded, what it takes to be 50%, 80% or 90% sure, "
         "which elements drive the risk, and the P80 reserve shared out.",
         template="cost-risk", demo="cost-risk"),
    Task("cer", "Fit a CER and price new programs", "cer", "--data",
         "Past programs with their cost and technical drivers (weight, power, lines of "
         "code), and the new programs to price.",
         "The CER fitted three ways (OLS, MUPE, ZMPE), what each driver does to cost, and "
         "each new program's estimate with a range.",
         template="cer", demo="cer"),
    Task("phase", "Phase an estimate into a budget", "phase", "--data",
         "Each estimate line's total in base-year dollars, its start year and years, a "
         "spending profile, and the inflation index.",
         "Spend by fiscal year and appropriation in then-year and base-year dollars, what "
         "inflation adds, and the peak year.",
         template="phase", demo="phase", filetypes=XLSX),
    Task("inflate", "Convert base-year and then-year dollars", "inflate", "--data",
         "A table of amounts with the fiscal year each is spent in, and your agency's "
         "published inflation index.",
         "Every amount converted, with the factor and index beside it, and the totals by "
         "fiscal year.",
         units=False, template="inflate", demo="inflate",
         filetypes=(("CSV file", "*.csv"), ("Excel workbook", "*.xlsx"))),
    Task("evm", "Where is the program heading? (EVM)", "evm", "--data",
         "Monthly BCWS, BCWP and ACWP, by control account if you have them, in Excel or "
         "CSV.",
         "CPI, SPI, the warning signs, the accounts owing a variance report, and the final "
         "cost and finish as a range.",
         template="evm", demo="evm"),
    Task("schedule", "Check a schedule (DCMA 14-point)", "schedule-check", "--mspdi",
         "A Microsoft Project schedule saved as XML (File, Save As, XML Format).",
         "The 14 checks, which pass and fail, the tasks behind each failure, and the "
         "critical path.",
         units=False, demo="schedule", filetypes=(("Project XML", "*.xml"),)),
    Task("jcl", "Cost and schedule together (JCL)", "jcl", "--spec",
         "The activities with their durations, links, costs and risks (the template "
         "shows how), or a Microsoft Project file named in it.",
         "The chance of meeting a budget and a date together, the 70% line, and which "
         "activities drive the schedule.",
         units=False, template="jcl", demo="jcl", filetypes=SPEC),
    Task("aoa", "Compare alternatives (AoA)", "aoa", "--spec",
         "Each alternative's cost lines phased by year, its effectiveness, and which one "
         "is the status quo.",
         "Life-cycle cost of each, how often each is cheapest, which are dominated, and "
         "whether each pays for itself against the status quo.",
         units=False, template="aoa", demo="aoa", filetypes=SPEC),
    Task("portfolio", "Choose what to fund", "portfolio", "--spec",
         "Candidate programs, their funding options and value, and the budget by year.",
         "The best set to fund within each year's budget, what an extra dollar buys, and "
         "the chance each year goes over.",
         units=False, template="portfolio", demo="portfolio", filetypes=SPEC),
    Task("open", "I just have a file", "open", "",
         "Any file cost-core reads: an estimate, EVM data, a Project XML, a spec.",
         "cost-core works out what it is and runs the right analysis on it.",
         units=False, filetypes=(("Files cost-core reads", "*.xlsx *.csv *.xml *.json *.zip"),
                                 ("All files", "*.*"))),
)

#: The units the window offers, as the words the commands take.
UNITS = ("as entered", "dollars", "thousands", "millions", "billions")

#: The status quo choice that leaves it to the workbook.
AS_IN_WORKBOOK = "(as in the workbook)"

#: Dollar bases the conversion offers; any BY year can be typed.
BASES = ("ty", "by2024", "by2025", "by2026", "by2027", "by2028")


def task(key: str) -> Task:
    return next(t for t in TASKS if t.key == key)


#: Plain-word explanations, shown from the Help window and on hover.
GLOSSARY: Tuple[Tuple[str, str], ...] = (
    ("Point estimate", "The single number in the estimate, before any uncertainty."),
    ("P50, P80", "The cost you'd be 50% or 80% sure of not exceeding. P80 is the usual "
     "funding target; the gap between the point estimate and it is the reserve."),
    ("Confidence of the point estimate", "How often the simulated cost came in at or "
     "under the point estimate. 30% means it's exceeded 70% of the time."),
    ("Correlation", "How much two elements tend to overrun together, from -1 to 1. "
     "Leaving it out makes the P80 too low."),
    ("CER", "Cost estimating relationship: an equation predicting cost from technical "
     "drivers (weight, power, lines of code), fitted to past programs."),
    ("OLS, MUPE, ZMPE", "Three ways to fit a CER. OLS is ordinary least squares; MUPE and "
     "ZMPE fit percentage errors, which suits cost, where bigger programs miss by bigger "
     "amounts. MUPE is the usual choice."),
    ("Prediction interval", "The range one new program is expected to land in, at the "
     "chosen confidence. Wider than the uncertainty in the line itself, on purpose."),
    ("Typical miss (CV)", "How far, as a percentage, programs typically sit from the "
     "CER's prediction. Under 20% is good for cost data."),
    ("R-squared", "How much of the variation the CER explains, 0 to 1. On a log-log CER "
     "it's high for almost any cost data, so the typical miss matters more."),
    ("p-value", "Below 0.10, a driver's effect is clear in the data; above it, it can't "
     "be told apart from zero."),
    ("Outside the data", "A new program whose drivers are beyond the programs the CER was "
     "fitted on. The CER has no evidence there."),
    ("Base-year and then-year dollars", "Base-year (constant) dollars take inflation out, "
     "for comparing; then-year dollars are what's actually spent each year, for budgets."),
    ("Appropriation", "The kind of money: RDT&E, Procurement, O&M, MILCON. Each has its "
     "own inflation index."),
    ("Rayleigh profile", "The usual shape of development spending: rising fast, peaking "
     "early, tailing off."),
    ("Status quo", "Keeping what you have: the baseline the alternatives are measured "
     "against."),
    ("SIR", "Savings-to-investment ratio: what an alternative saves in running costs over "
     "what it costs up front, both in present value. Above 1 means it pays for itself."),
    ("Payback year", "The year the discounted savings have repaid the extra investment."),
    ("IRR", "The real return on the extra investment, to compare with the discount rate."),
    ("Uniform annual cost", "An alternative's present value spread evenly over its own "
     "years: the fair comparison when service lives differ."),
    ("CPI, SPI", "Cost and schedule performance indices: earned value over actual cost, "
     "and over planned value. Below 1 is over cost or behind schedule."),
    ("TCPI", "The efficiency the rest of the work needs to finish on the budget. Far above "
     "the CPI so far is a warning sign."),
    ("EAC, IEAC", "Estimate at completion: the contractor's (EAC) and independent ones "
     "computed from performance (IEAC)."),
    ("SPI(t)", "Schedule performance on earned schedule, in time rather than dollars; "
     "unlike SPI it doesn't drift back to 1 near the end."),
    ("DCMA 14-point", "A standard set of checks on a schedule's logic: missing links, "
     "leads and lags, hard constraints, high float and so on."),
    ("JCL", "Joint cost and schedule confidence: the chance of meeting a budget AND a "
     "date together, always lower than either alone."),
)


# ----------------------------------------------------------------- places
def documents_dir() -> Path:
    """The user's Documents folder, wherever Windows has put it (OneDrive
    often moves it)."""
    if sys.platform == "win32":
        try:
            import ctypes
            from ctypes import wintypes

            buf = ctypes.create_unicode_buffer(wintypes.MAX_PATH)
            # CSIDL_PERSONAL = 5, SHGFP_TYPE_CURRENT = 0
            if ctypes.windll.shell32.SHGetFolderPathW(None, 5, None, 0, buf) == 0 and buf.value:
                return Path(buf.value)
        except (AttributeError, OSError):
            pass
    return Path.home() / "Documents"


def work_dir() -> Path:
    """Where templates and examples go: Documents\\cost-core."""
    return documents_dir() / "cost-core"


def results_dir(workbook: Path) -> Path:
    """Results go beside the workbook, in '<name> results', as ``ce-core
    open`` puts them, so the two never disagree."""
    workbook = Path(workbook)
    return workbook.parent / f"{workbook.stem} results"


def next_free(out: Path) -> Path:
    """'<name> results 2', '... 3' and so on: a folder not yet used."""
    out = Path(out)
    k = 2
    while (out.parent / f"{out.name} {k}").exists():
        k += 1
    return out.parent / f"{out.name} {k}"


def fresh_name(path: Path) -> Path:
    """'my_estimate 2.xlsx', '... 3' and so on: a file name not yet used."""
    path = Path(path)
    k = 2
    while path.with_name(f"{path.stem} {k}{path.suffix}").exists():
        k += 1
    return path.with_name(f"{path.stem} {k}{path.suffix}")


def locked(out: Path) -> List[str]:
    """Files in ``out`` (or one folder in, where ``open`` writes) that are
    open in Excel or PowerPoint, which the next run couldn't overwrite.
    Office leaves a ~$ owner file beside an open file and holds the file
    itself locked; either is taken as open."""
    out = Path(out)
    if not out.is_dir():
        return []
    found = []
    files = [f for f in out.iterdir() if f.is_file()]
    for sub in (d for d in out.iterdir() if d.is_dir()):
        files += [f for f in sub.iterdir() if f.is_file()]
    for f in sorted(files):
        if f.suffix.lower() not in (".xlsx", ".pptx", ".csv") or f.name.startswith("~$"):
            continue
        name = str(f.relative_to(out))
        if (f.parent / f"~${f.name}").exists() or (f.parent / f"~${f.name[2:]}").exists():
            found.append(name)
            continue
        try:
            with open(f, "r+b"):
                pass
        except PermissionError:
            found.append(name)
        except OSError:
            pass
    return found


def settings_path() -> Path:
    base = os.environ.get("APPDATA")
    return (Path(base) / "cost-core" / "gui.json") if base else \
        Path.home() / ".cost-core-gui.json"


def load_prefs(path: Optional[Path] = None) -> Dict:
    """What the window remembers: units, marking, recent files per task. A
    missing or broken file is an empty memory, never an error."""
    try:
        data = json.loads((path or settings_path()).read_text(encoding="utf-8"))
        return data if isinstance(data, dict) else {}
    except (OSError, ValueError):
        return {}


def save_prefs(prefs: Dict, path: Optional[Path] = None) -> None:
    path = path or settings_path()
    try:
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(json.dumps(prefs, indent=1), encoding="utf-8")
    except OSError:
        pass


def recent(prefs: Dict, key: str) -> List[str]:
    """The task's recent files that are still there."""
    return [p for p in prefs.get("recent", {}).get(key, []) if Path(p).exists()]


def add_recent(prefs: Dict, key: str, path: str, keep: int = 6) -> List[str]:
    """Put ``path`` first in the task's recent files, without repeats."""
    recent = [p for p in prefs.setdefault("recent", {}).get(key, []) if p != path]
    recent.insert(0, path)
    prefs["recent"][key] = recent[:keep]
    return prefs["recent"][key]


# ------------------------------------------------------------- commands
def python_for_children() -> str:
    """The interpreter child runs use: python.exe beside pythonw.exe, so
    their output can be read, with the console window kept hidden."""
    exe = Path(sys.executable)
    if exe.name.lower() == "pythonw.exe":
        console = exe.with_name("python.exe")
        if console.exists():
            return str(console)
    return str(exe)


def run_argv(t: Task, workbook: Path, units: str = "as entered", marking: str = "",
             out: Optional[Path] = None, extras: Optional[Dict[str, str]] = None) -> List[str]:
    """The ``ce-core`` arguments that run ``t`` on ``workbook``."""
    workbook = Path(workbook)
    out = out or results_dir(workbook)
    extras = extras or {}
    if t.command == "open":
        argv = ["open", str(workbook), "--out", str(out)]
    else:
        argv = [t.command, t.flag, str(workbook), "--out", str(out)]
    if t.key == "inflate":
        argv += ["--index", extras.get("index", ""), "--from", extras.get("from", "by2026"),
                 "--to", extras.get("to", "ty")]
    if t.key == "aoa" and extras.get("status_quo"):
        argv += ["--status-quo", extras["status_quo"]]
    if t.units and units and units != "as entered":
        argv += ["--units", units]
    if marking.strip() and t.key != "inflate":
        argv += ["--marking", marking.strip()]
    return argv


def demo_argv(t: Task, marking: str = "") -> Tuple[List[str], Path]:
    out = work_dir() / "examples" / t.demo
    argv = ["demo", t.demo, "--out", str(out)]
    if marking.strip() and t.key != "inflate":
        argv += ["--marking", marking.strip()]
    return argv, out


def run_cost_core(argv: Sequence[str], timeout: float = 3600,
                  started=None) -> Tuple[int, str]:
    """Run ``python -m cost_core <argv>`` and return its exit code and
    everything it printed. A separate process, so a run can't leave state
    behind in the window, and the same code path as the command line.
    ``started`` is handed the process, so the window can stop it on close."""
    flags = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    env = dict(os.environ, PYTHONIOENCODING="utf-8", MPLBACKEND="Agg")
    p = subprocess.Popen([python_for_children(), "-m", "cost_core", *map(str, argv)],
                         stdout=subprocess.PIPE, stderr=subprocess.STDOUT,
                         stdin=subprocess.DEVNULL, text=True, encoding="utf-8",
                         errors="replace", creationflags=flags, env=env)
    if started is not None:
        started(p)
    try:
        output, _ = p.communicate(timeout=timeout)
    except subprocess.TimeoutExpired:
        p.kill()
        p.communicate()
        return 1, f"[ERROR] It ran for over {timeout / 60:.0f} minutes and was stopped."
    return p.returncode, output or ""


# -------------------------------------------------------------- reading
def summary(output: str) -> List[str]:
    """The plain-English sentences a run printed under 'What this means'."""
    lines, inside, current = [], False, ""
    for raw in output.splitlines():
        if raw.strip() == "What this means:":
            inside, current = True, ""
            continue
        if not inside:
            continue
        if raw.startswith("  - "):
            if current:
                lines.append(current)
            current = raw[4:].strip()
        elif raw.startswith("    ") and current:
            current += " " + raw.strip()
        else:
            if current:
                lines.append(current)
            current, inside = "", False
    if current:
        lines.append(current)
    return lines


def notes(output: str) -> List[str]:
    """The run's notes: things it did that the reader should know about."""
    return [ln.split("note:", 1)[1].strip() for ln in output.splitlines()
            if ln.strip().startswith("note:")]


def errors(output: str) -> List[str]:
    """Why a run stopped, as the commands put it, without the log prefix,
    and with the commands' advice put in terms of the window's buttons."""
    if re.search(r"PermissionError|Permission denied", output):
        file = re.search(r"Permission denied: '([^']+)'", output)
        # the last part of the path, split by hand: a Windows path on another
        # system has no separators pathlib knows
        name = re.split(r"[\\/]+", file.group(1).replace("\\\\", "\\"))[-1] if file else ""
        return [f"{name or 'A results file'} couldn't "
                "be saved. It's probably open in Excel or PowerPoint: close it and press "
                "Run again. If it isn't, the folder may be read-only; move your workbook "
                "somewhere you can save to."]
    found = []
    for ln in output.splitlines():
        m = re.match(r"\s*\[(ERROR|CRITICAL)\]\s*(.*)", ln)
        if m:
            found.append(re.sub(r"^[\w -]+ failed: ", "", m.group(2)))
        elif found and ln.startswith("  ") and not ln.strip().startswith(("File ", "^")):
            found[-1] += "\n" + ln.strip()
    if not found:
        tail = [ln for ln in output.splitlines() if ln.strip()]
        found = tail[-3:] or ["It stopped without saying why. Try 'Check my workbook', "
                              "or run it again."]
    words = [(r"To see this command work first: ce-core demo [\w-]+",
              "To see it work first, press 'See an example'."),
             (r"For a file to fill in: ce-core template [\w-]+",
              "For a workbook to fill in, press 'Get the template'."),
             (r"ce-core template [\w-]+ writes (one|a sheet)[^.\n]*",
              "'Get the template' gives you one")]
    for pattern, plain in words:
        found = [re.sub(pattern, plain, f) for f in found]
    return found


def written(out: Path) -> Dict[str, Path]:
    """The report, briefing and folder a run left, when they exist. ``open``
    puts its results one folder further in, under the file's own name."""
    out = Path(out)
    found = {}
    candidates = [out] + ([p for p in out.iterdir() if p.is_dir()] if out.is_dir() else [])
    for folder in candidates:
        for name, key in (("report.xlsx", "report"), ("brief.pptx", "brief")):
            if key not in found and (folder / name).exists():
                found[key] = folder / name
    if "report" not in found and converted(out) is not None:
        found["report"] = converted(out)
    if out.is_dir():
        found["folder"] = out
    return found


# ---------------------------------------------------------- the answer
#: The chart each kind of result draws, and a file only it writes.
SIGNATURES = {"cost-risk": ("cost_risk_s_curve.png", "cost_risk_drivers.png"),
              "cer": ("cer_fit.png", "estimates.csv"),
              "phase": ("phasing.png", "phasing_long.csv"),
              "evm": ("evm.png", "forecast_percentiles.csv"),
              "portfolio": (None, "selected.csv"),   # before jcl: both write frontier.csv
              "jcl": ("jcl.png", "frontier.csv"),
              "aoa": ("aoa_s_curves.png", "s_curves.csv"),
              "schedule": (None, "dcma.csv")}


def converted(out: Path) -> Optional[Path]:
    """The table ``inflate`` wrote, which is named after the input."""
    out = Path(out)
    if not out.is_dir():
        return None
    for f in sorted(out.glob("*.csv")):
        try:
            head = f.open(encoding="utf-8").readline()
        except OSError:
            continue
        if "factor_" in head and "index_used" in head:
            return f
    return None


def result_kind(out: Path) -> Optional[str]:
    """Which analysis left these results, from the files only it writes."""
    out = Path(out)
    for kind, (_, marker) in SIGNATURES.items():
        if (out / marker).exists():
            return kind
    return "inflate" if converted(out) else None


def _money(v: float, units: str) -> str:
    from cost_core import plain

    return plain._money3(float(v), plain.units_label(units or ""))


def recorded_units(out: Path) -> str:
    """The units a run recorded: in assumptions.json, or else on the
    report's Assumptions sheet (cost risk writes no JSON). Empty if neither
    says."""
    out = Path(out)
    try:
        units = json.loads((out / "assumptions.json").read_text(encoding="utf-8")).get("units")
        if units:
            return str(units)
    except (OSError, ValueError, AttributeError):
        pass
    try:
        from openpyxl import load_workbook

        wb = load_workbook(out / "report.xlsx", read_only=True)
        try:
            sheet = next((wb[n] for n in wb.sheetnames if n.strip().lower() == "assumptions"),
                         None)
            for row in (sheet.iter_rows(values_only=True) if sheet is not None else ()):
                if row and str(row[0]).strip().lower() == "units" and len(row) > 1 and row[1]:
                    return str(row[1])
        finally:
            wb.close()
    except Exception:  # noqa: BLE001 - no report, no units: the tiles go unlabelled
        pass
    return ""


def answer(out: Path, units: str = "", kind: Optional[str] = None
           ) -> Tuple[List[Tuple[str, str]], Optional[Path]]:
    """The key numbers of a result as (label, value) tiles, and its main
    chart. Read from the files the run wrote, so they are the report's own
    numbers. ``kind`` is the job that ran, when known: a folder can hold an
    older run of another job too. Anything unreadable gives fewer tiles,
    never an error."""
    import pandas as pd

    out = Path(out)
    if out.is_dir() and not result_kind(out):
        inner = [p for p in out.iterdir() if p.is_dir() and result_kind(p)]
        if inner:
            out = max(inner, key=lambda p: p.stat().st_mtime)
    if kind not in SIGNATURES and kind != "inflate":
        kind = result_kind(out)
    if kind is None:
        return [], None
    units = recorded_units(out) or units
    if units in ("as entered", None):
        units = ""
    m = lambda v: _money(v, units)  # noqa: E731
    chart = SIGNATURES.get(kind, (None,))[0]
    chart = out / chart if chart and (out / chart).exists() else None
    tiles: List[Tuple[str, str]] = []
    try:
        if kind == "cost-risk":
            s = pd.read_csv(out / "summary.csv").set_index("statistic")["value"]
            tiles = [("Point estimate", m(s["point_estimate"])),
                     ("Chance it's enough", f"{s['point_estimate_percentile']:.0f}%"),
                     ("To be 50% sure", m(s["p50"])), ("To be 80% sure", m(s["p80"]))]
        elif kind == "cer":
            meth = pd.read_csv(out / "methods.csv")
            chosen = meth[meth["chosen"].astype(str) == "yes"].iloc[0]
            est = pd.read_csv(out / "estimates.csv")
            if len(est):
                r = est.iloc[0]
                tiles.append((str(r["name"]), m(r["estimate"])))
                tiles.append((f"{r['level']:.0%} range", f"{m(r['lower'])} to {m(r['upper'])}"))
            tiles += [("Typical miss", f"{chosen['cv']:.0%}"),
                      ("Programs fitted", f"{int(chosen['n_obs'])}")]
        elif kind == "phase":
            long = pd.read_csv(out / "phasing_long.csv")
            col = "then_year" if long["then_year"].notna().any() else "base_year"
            per = long.groupby("fiscal_year")[col].sum()
            tiles = [("Base-year total", m(long["base_year"].sum()))]
            if col == "then_year":
                tiles.append(("Then-year total", m(long["then_year"].sum())))
            tiles += [("Years", f"FY{per.index.min()} to FY{per.index.max()}"),
                      ("Peak year", f"FY{per.idxmax()}: {m(per.max())}")]
        elif kind == "evm":
            s = pd.read_csv(out / "summary.csv").drop_duplicates("measure") \
                .set_index("measure")["value"]
            tiles = [("CPI", f"{float(s['CPI']):.2f}"), ("SPI", f"{float(s['SPI']):.2f}"),
                     ("Final cost, P50", m(float(s["EAC P50"]))),
                     ("Final cost, P80", m(float(s["EAC P80"])))]
        elif kind == "jcl":
            s = pd.read_csv(out / "summary.csv").set_index("measure")["value"]
            tiles = [("Plan's joint confidence",
                      f"{float(s['point estimate: joint confidence']):.0%}"),
                     ("Schedule P70", f"{float(s['schedule P70 (months)']):.1f} months"),
                     ("Cost P70", m(float(s["cost P70"])))]
        elif kind == "aoa":
            s = pd.read_csv(out / "summary.csv").sort_values("p50")
            best = s.iloc[0]
            tiles = [("Cheapest at P50", str(best["alternative"])), ("Its P50", m(best["p50"])),
                     ("Cheapest in", f"{best['p_cheapest']:.0%} of simulations")]
            if (out / "economic.csv").exists():
                e = pd.read_csv(out / "economic.csv")
                e = e[e["status_quo"].astype(str) != "yes"].sort_values("net_savings_pv",
                                                                       ascending=False)
                if len(e):
                    r = e.iloc[0]
                    sir = "" if pd.isna(r["sir"]) else f", SIR {r['sir']:.2f}"
                    tiles.append(("Best against the status quo", f"{r['alternative']}{sir}"))
        elif kind == "portfolio":
            s = pd.read_csv(out / "selected.csv")
            funded = s["option"].notna() & (s["option"].astype(str) != "")
            tiles = [("Programs funded", f"{int(funded.sum())} of {len(s)}"),
                     ("Total value", f"{s['value'].sum():,.1f}"),
                     ("Total cost", m(s.loc[funded, "cost"].sum()))]
        elif kind == "inflate":
            table = pd.read_csv(converted(out))
            after = [c for c in table.columns if c.startswith("amount_")][0]
            tiles = [("Before", f"{table['amount'].sum():,.1f}"),
                     (after.replace("amount_", "After, ").replace("_", "-"),
                      f"{table[after].sum():,.1f}"),
                     ("Rows converted", f"{len(table)}")]
        elif kind == "schedule":
            d = pd.read_csv(out / "dcma.csv")
            assessed = d["passed"].dropna().astype(str)
            passed = int((assessed == "True").sum())
            tiles = [("Checks passed", f"{passed} of {len(assessed)}"),
                     ("Checks failed", f"{len(assessed) - passed}")]
    except (OSError, KeyError, ValueError, IndexError, TypeError):
        pass
    return tiles, chart


# --------------------------------------------------- finding problems
_WHERE = re.compile(r"(?P<sheet>[A-Z][\w&]*(?: [A-Z][\w&]*)*) sheet, row (?P<row>\d+)"
                    r"(?: \((?P<label>[^)]*)\))?:? ?(?P<rest>.*)", re.S)


def problem_location(message: str) -> Optional[Tuple[str, int, str]]:
    """(sheet, Excel row, message) for a message naming a sheet and row."""
    m = _WHERE.search(message)
    if not m:
        return None
    return m.group("sheet").strip(), int(m.group("row")), message


def mark_problem(workbook: Path, message: str) -> Optional[Path]:
    """A copy of ``workbook`` with the cell the message is about outlined in
    red and the message attached to it, beside the original. The column is
    the heading the message names, or the whole row when it names none. The
    user's own file is never changed."""
    from openpyxl import load_workbook
    from openpyxl.comments import Comment
    from openpyxl.styles import Border, PatternFill, Side

    where = problem_location(message)
    workbook = Path(workbook)
    if where is None or workbook.suffix.lower() not in (".xlsx", ".xlsm"):
        return None
    sheet, row, text = where
    try:
        wb = load_workbook(workbook, keep_vba=workbook.suffix.lower() == ".xlsm")
    except Exception:  # noqa: BLE001 - an unreadable workbook just has no marked copy
        return None
    lowered_text = text.lower()
    named = [w for w in wb.worksheets
             if re.search(rf"\b{re.escape(w.title.strip().lower())} sheet, row {row}\b",
                          lowered_text)]
    ws = max(named, key=lambda w: len(w.title)) if named else None
    if ws is None:
        return None
    headings = [(c.column, str(c.value)) for c in ws[1] if c.value is not None]
    lowered = text.lower()
    columns = [col for col, name in sorted(headings, key=lambda h: -len(h[1]))
               if re.search(rf"\b{re.escape(name.lower())}\b", lowered)][:1]
    if not columns:
        columns = [col for col, _ in headings]
    red = Side(style="thick", color="C00000")
    for col in columns:
        cell = ws.cell(row=row, column=col)
        cell.border = Border(left=red, right=red, top=red, bottom=red)
        cell.fill = PatternFill("solid", fgColor="F8CBAD")
    first = ws.cell(row=row, column=columns[0])
    note = Comment(f"cost-core stopped here:\n{text}", "cost-core")
    note.width, note.height = 340, 140
    first.comment = note
    ws.sheet_view.selection[0].activeCell = first.coordinate
    ws.sheet_view.selection[0].sqref = first.coordinate
    wb.active = wb.worksheets.index(ws)
    target = workbook.with_name(f"{workbook.stem} - problem marked{workbook.suffix}")
    try:
        wb.save(target)
    except OSError:
        return None
    return target


#: Required columns a row may still leave blank: EVM's future months have a
#: plan but no earned value or actuals yet.
BLANK_ALLOWED = {("evm", "Data", "Earned Value (BCWP)"), ("evm", "Data", "Actual Cost (ACWP)"),
                 ("evm", "Data", "Contractor EAC")}


def _scan_tables(kind: str, path: Path) -> List[str]:
    """Problems visible cell by cell: a required column left blank in a row
    that has data, and text in a column of numbers."""
    from openpyxl import load_workbook

    from cost_core.template_style import TABLES

    tables = TABLES.get(kind)
    if not tables or path.suffix.lower() not in (".xlsx", ".xlsm"):
        return []
    try:
        wb = load_workbook(path, data_only=True, read_only=True)
    except Exception:  # noqa: BLE001 - the reader will say why it can't open it
        return []
    problems = []
    for ws in wb.worksheets:
        spec = tables.get(ws.title)
        if spec is None:
            continue
        rows = list(ws.iter_rows(values_only=True))
        if not rows:
            continue
        head = [None if h is None else str(h).strip() for h in rows[0]]
        body = [(i + 2, r) for i, r in enumerate(rows[1:])
                if any(v is not None and str(v).strip() != "" for v in r)]
        for j, name in enumerate(head):
            if name is None:
                continue
            values = [(n, r[j] if j < len(r) else None) for n, r in body]
            filled = [(n, v) for n, v in values if v is not None and str(v).strip() != ""]
            required = (spec.get(name) or spec.get("*") or (False,))[0]
            if required and spec.get(name) is not None                     and (kind, ws.title, name) not in BLANK_ALLOWED:
                for n, v in values:
                    if v is None or str(v).strip() == "":
                        problems.append(f"{ws.title} sheet, row {n}: {name} is blank, and "
                                        "it's required.")
            numbers = [v for _, v in filled if isinstance(v, (int, float))
                       and not isinstance(v, bool)]
            if filled and len(numbers) >= len(filled) / 2:
                for n, v in filled:
                    if isinstance(v, str) and not _numberish(v):
                        problems.append(f"{ws.title} sheet, row {n}: {name} is {v!r}, "
                                        "which isn't a number.")
    wb.close()
    return problems


def _numberish(text: str) -> bool:
    t = text.strip().replace(",", "").replace("$", "").rstrip("%")
    t = re.sub(r"^(FY|BY|CY)\s*", "", t, flags=re.I)
    try:
        float(t)
        return True
    except ValueError:
        return False


def check_workbook(t: Task, path: Path) -> Dict[str, List[str]]:
    """What's wrong with a workbook, without running the analysis: the
    cell-by-cell problems first, then what the command's own reader says
    (it stops at the first problem it meets), then its notes."""
    import logging
    import warnings

    path = Path(path)
    problems = _scan_tables(t.key, path)
    notes_: List[str] = []
    logging.disable(logging.WARNING)
    try:
        with warnings.catch_warnings():
            warnings.simplefilter("ignore")
            notes_ = _read_with_command(t, path)
    except Exception as e:  # noqa: BLE001 - every reader's refusal, shown as written
        text = str(e) or type(e).__name__
        where = problem_location(text)
        seen = {problem_location(p)[:2] for p in problems if problem_location(p)}
        if text not in problems and (where is None or where[:2] not in seen):
            problems.append(text)
    finally:
        logging.disable(logging.NOTSET)
    return {"problems": problems, "notes": notes_}


def _read_with_command(t: Task, path: Path) -> List[str]:
    """Read ``path`` the way ``t``'s command does; return its notes."""
    if t.key == "cost-risk":
        from cost_core.costrisk import read_workbook
        return list(read_workbook(path).notes)
    if t.key == "cer":
        from cost_core.cer.study import analyse, read_workbook
        return list(analyse(read_workbook(path)).notes)
    if t.key == "phase":
        from cost_core.phasing import phase, read_workbook
        return list(phase(read_workbook(path)).inputs.notes)
    if t.key == "evm":
        from cost_core.evm import EvmData
        return list(EvmData.read(path).notes)
    if t.key == "schedule":
        from cost_core.schedule.mspdi import read_mspdi
        read_mspdi(path)
        return []
    if t.key == "jcl":
        from cost_core.schedule.spec import load_project
        load_project(path)
        return []
    if t.key == "aoa":
        from cost_core.aoa.spec import load_spec
        load_spec(path)
        return []
    if t.key == "portfolio":
        from cost_core.portfolio.spec import load_portfolio
        load_portfolio(path)
        return []
    if t.key == "inflate":
        import pandas as pd
        frame = pd.read_excel(path) if path.suffix.lower() in (".xlsx", ".xlsm") \
            else pd.read_csv(path)
        if frame.empty:
            raise ValueError(f"{path.name} has no rows to convert.")
        return []
    if t.key == "open":
        from cost_core import opener
        return [f"This looks like {opener.plan(path).what}."]
    return []


def alternatives(path: Path) -> List[str]:
    """The alternatives an AoA workbook or spec names, for the status quo
    choice; empty when they can't be read."""
    path = Path(path)
    try:
        if path.suffix.lower() == ".json":
            return [a["name"] for a in json.loads(path.read_text(encoding="utf-8"))
                    .get("alternatives", [])]
        import pandas as pd
        sheets = pd.read_excel(path, sheet_name=None)
        for name, frame in sheets.items():
            if name.strip().lower() == "alternatives" and len(frame.columns):
                return [str(v).strip() for v in frame.iloc[:, 0].dropna()]
        lines = next((f for n, f in sheets.items() if n.strip().lower() == "lines"), None)
        if lines is not None and len(lines.columns):
            return list(dict.fromkeys(str(v).strip() for v in lines.iloc[:, 0].dropna()))
    except Exception:  # noqa: BLE001 - no list is a fine answer
        pass
    return []


def open_path(path: Path) -> None:
    """Open a file or folder the way double-clicking it would."""
    path = Path(path)
    if sys.platform == "win32":
        os.startfile(str(path))  # noqa: S606 - the user's own file, by their click
    elif sys.platform == "darwin":
        subprocess.Popen(["open", str(path)])
    else:
        subprocess.Popen(["xdg-open", str(path)])


def make_template(t: Task, path: Path) -> Path:
    """Write the template for ``t`` at ``path`` (and inflate's amounts)."""
    from cost_core import templates

    path.parent.mkdir(parents=True, exist_ok=True)
    written_path = templates.write(t.template, path)
    templates.write_companion(t.template, written_path)
    return written_path


def template_path(t: Task) -> Path:
    from cost_core.cli import TEMPLATE_FILES

    return work_dir() / TEMPLATE_FILES[t.template]


# ------------------------------------------------------------- shortcut
def make_shortcut(folder: Optional[Path] = None) -> Path:
    """A desktop shortcut that opens the window, with no console behind it.

    A real Windows shortcut (.lnk) to pythonw, made through PowerShell's
    WScript.Shell; if that isn't allowed, a .cmd that starts it instead.
    """
    if sys.platform != "win32":
        raise OSError("The desktop shortcut is for Windows; run ce-core gui instead.")
    exe = Path(sys.executable)
    pythonw = exe.with_name("pythonw.exe") if exe.with_name("pythonw.exe").exists() else exe
    no_window = getattr(subprocess, "CREATE_NO_WINDOW", 0)
    if folder is None:
        q = subprocess.run(["powershell", "-NoProfile", "-Command",
                            "[Environment]::GetFolderPath('Desktop')"],
                           capture_output=True, text=True, creationflags=no_window)
        folder = Path(q.stdout.strip() or Path.home() / "Desktop")
    folder = Path(folder)
    lnk = folder / "cost-core.lnk"
    script = ("$s = (New-Object -ComObject WScript.Shell).CreateShortcut($args[0]); "
              "$s.TargetPath = $args[1]; $s.Arguments = '-m cost_core.gui'; "
              "$s.WorkingDirectory = $args[2]; "
              "$s.Description = 'cost-core: cost risk, CERs, phasing, EVM, schedules'; "
              "$s.Save()")
    work = work_dir()
    work.mkdir(parents=True, exist_ok=True)
    p = subprocess.run(["powershell", "-NoProfile", "-Command", script, str(lnk), str(pythonw),
                        str(work)], capture_output=True, text=True, creationflags=no_window)
    if p.returncode == 0 and lnk.exists():
        return lnk
    cmd = folder / "cost-core.cmd"
    cmd.write_text(f'@echo off\r\nstart "" "{pythonw}" -m cost_core.gui\r\n', encoding="utf-8")
    return cmd


# --------------------------------------------------------------- window
class Tooltip:
    """A note that appears when the mouse rests on a widget."""

    def __init__(self, widget, text: str, delay: int = 500) -> None:
        self.widget, self.text, self.delay = widget, text, delay
        self.tip = self.job = None
        widget.bind("<Enter>", self.schedule, add="+")
        widget.bind("<Leave>", self.hide, add="+")
        widget.bind("<ButtonPress>", self.hide, add="+")

    def schedule(self, _=None):
        self.job = self.widget.after(self.delay, self.show)

    def show(self):
        import tkinter as tk

        if self.tip or not self.text:
            return
        x = self.widget.winfo_rootx() + 20
        y = self.widget.winfo_rooty() + self.widget.winfo_height() + 4
        self.tip = tk.Toplevel(self.widget)
        self.tip.wm_overrideredirect(True)
        self.tip.wm_geometry(f"+{x}+{y}")
        tk.Label(self.tip, text=self.text, justify="left", background="#ffffe8",
                 relief="solid", borderwidth=1, wraplength=380, padx=8, pady=5).pack()

    def hide(self, _=None):
        if self.job:
            self.widget.after_cancel(self.job)
            self.job = None
        if self.tip:
            self.tip.destroy()
            self.tip = None


class App:
    """The window. Layout and wiring only; the functions above decide."""

    def __init__(self, root) -> None:
        import tkinter as tk
        from tkinter import ttk

        from cost_core import __version__

        self.tk, self.ttk, self.root = tk, ttk, root
        self.prefs = load_prefs()
        root.title(f"cost-core {__version__}")
        scale = max(root.winfo_fpixels("1i") / 96.0, 1.0)
        width = min(int(1120 * scale), root.winfo_screenwidth() - 40)
        height = min(int(760 * scale), root.winfo_screenheight() - 80)
        root.geometry(f"{width}x{height}")
        root.minsize(min(int(960 * scale), width), min(int(640 * scale), height))
        style = ttk.Style(root)
        if "vista" in style.theme_names():
            style.theme_use("vista")
        base = "Segoe UI" if sys.platform == "win32" else "TkDefaultFont"
        self.base = base
        style.configure("Title.TLabel", font=(base, 15, "bold"))
        style.configure("Step.TLabel", font=(base, 11, "bold"))
        style.configure("Big.TButton", font=(base, 11, "bold"), padding=(14, 6))
        style.configure("Task.TButton", anchor="w", padding=(10, 7))
        style.configure("TileLabel.TLabel", font=(base, 9), foreground="#555")
        style.configure("TileValue.TLabel", font=(base, 14, "bold"), foreground="#1F4E79")

        outer = ttk.Frame(root, padding=10)
        outer.pack(fill="both", expand=True)
        left = ttk.Frame(outer)
        left.pack(side="left", fill="y", padx=(0, 12))
        ttk.Label(left, text="What do you want to do?", style="Step.TLabel").pack(
            anchor="w", pady=(0, 6))
        self.task_buttons = {}
        for t in TASKS:
            b = ttk.Button(left, text=t.title, style="Task.TButton", width=36,
                           command=lambda k=t.key: self.choose(k))
            b.pack(fill="x", pady=2)
            Tooltip(b, f"{t.get}")
            self.task_buttons[t.key] = b
        ttk.Separator(left).pack(fill="x", pady=10)
        ttk.Button(left, text="What do these words mean?", command=self.glossary).pack(
            fill="x", pady=2)
        ttk.Button(left, text="Settings", command=self.settings).pack(fill="x", pady=2)

        right = ttk.Frame(outer)
        right.pack(side="left", fill="both", expand=True)
        self.title = ttk.Label(right, style="Title.TLabel")
        self.title.pack(anchor="w")
        self.need = ttk.Label(right, wraplength=680, justify="left")
        self.need.pack(anchor="w", pady=(6, 0))
        self.get = ttk.Label(right, wraplength=680, justify="left", foreground="#444")
        self.get.pack(anchor="w", pady=(2, 8))

        step1 = ttk.Frame(right)
        step1.pack(fill="x", pady=3)
        ttk.Label(step1, text="1.", style="Step.TLabel", width=3).pack(side="left")
        self.example_btn = ttk.Button(step1, text="See an example", command=self.example)
        self.example_btn.pack(side="left")
        Tooltip(self.example_btn, "Runs this job on invented numbers so you can see what "
                                  "comes out before you fill anything in.")
        self.template_btn = ttk.Button(step1, text="Get the template (opens in Excel)",
                                       command=self.template)
        self.template_btn.pack(side="left", padx=8)
        Tooltip(self.template_btn, "A workbook to fill in, saved in Documents\\cost-core. "
                                   "Yellow cells are yours; hover over a heading for what "
                                   "goes there.")

        step2 = ttk.Frame(right)
        step2.pack(fill="x", pady=3)
        ttk.Label(step2, text="2.", style="Step.TLabel", width=3).pack(side="left")
        self.file_label = ttk.Label(step2, text="Your file:")
        self.file_label.pack(side="left")
        self.file_var = tk.StringVar()
        self.file_box = ttk.Combobox(step2, textvariable=self.file_var)
        self.file_box.pack(side="left", fill="x", expand=True, padx=6)
        self.file_box.bind("<<ComboboxSelected>>", lambda _: self.file_changed())
        self.file_box.bind("<Return>", lambda _: self.file_changed())
        self.file_box.bind("<FocusOut>", lambda _: self.file_changed())
        ttk.Button(step2, text="Choose...", command=self.browse).pack(side="left")
        self.edit_btn = ttk.Button(step2, text="Open in Excel", command=self.edit_file)
        self.edit_btn.pack(side="left", padx=(6, 0))

        self.extras = ttk.Frame(right)   # packed only when a job has extra choices
        self.index_var = tk.StringVar(value=self.prefs.get("index", ""))
        self.from_var = tk.StringVar(value=self.prefs.get("from", "by2026"))
        self.to_var = tk.StringVar(value=self.prefs.get("to", "ty"))
        self.sq_var = tk.StringVar()

        opts = self.opts = ttk.Frame(right)
        opts.pack(fill="x", pady=3)
        ttk.Label(opts, text="", width=3).pack(side="left")
        ttk.Label(opts, text="Units:").pack(side="left")
        self.units_var = tk.StringVar(value=self.prefs.get("units", "as entered"))
        self.units_box = ttk.Combobox(opts, textvariable=self.units_var, values=UNITS,
                                      width=12, state="readonly")
        self.units_box.pack(side="left", padx=(4, 14))
        Tooltip(self.units_box, "What the money in your workbook is in, for the labels. "
                                "Nothing is converted. Leave it 'as entered' to use the "
                                "workbook's own Units setting; a choice here overrides it.")
        ttk.Label(opts, text="Marking (printed on every page):").pack(side="left")
        self.marking_var = tk.StringVar(value=self.prefs.get("marking", ""))
        mark = self.mark_entry = ttk.Entry(opts, textvariable=self.marking_var, width=18)
        mark.pack(side="left", padx=4)
        Tooltip(mark, "Exactly this text goes at the top and bottom of every slide and "
                      "sheet, e.g. CUI. cost-core doesn't check it.")

        step3 = ttk.Frame(right)
        step3.pack(fill="x", pady=(8, 3))
        ttk.Label(step3, text="3.", style="Step.TLabel", width=3).pack(side="left")
        self.check_btn = ttk.Button(step3, text="Check my workbook", command=self.check)
        self.check_btn.pack(side="left")
        Tooltip(self.check_btn, "Looks for blanks, text in number cells and anything the "
                                "analysis would refuse, in seconds, without running it.")
        self.run_btn = ttk.Button(step3, text="Run", style="Big.TButton", command=self.run)
        self.run_btn.pack(side="left", padx=8)
        self.progress = ttk.Progressbar(step3, mode="indeterminate", length=160)
        self.status = ttk.Label(step3, text="")
        self.status.pack(side="left", padx=10)

        buttons = ttk.Frame(right)
        buttons.pack(fill="x", pady=(6, 2))
        self.report_btn = ttk.Button(buttons, text="Open report (Excel)",
                                     command=lambda: self.open_result("report"))
        self.brief_btn = ttk.Button(buttons, text="Open slides",
                                    command=lambda: self.open_result("brief"))
        self.folder_btn = ttk.Button(buttons, text="Open results folder",
                                     command=lambda: self.open_result("folder"))
        self.where_btn = ttk.Button(buttons, text="Show me where", command=self.show_where)
        self.fix_btn = ttk.Button(buttons, text="Open my workbook to fix it",
                                  command=self.edit_file)
        self.result_buttons = (self.report_btn, self.brief_btn, self.folder_btn)
        self.problem_buttons = (self.where_btn, self.fix_btn)

        self.book = ttk.Notebook(right)
        self.book.pack(fill="both", expand=True, pady=(6, 0))
        page = ttk.Frame(self.book)
        self.book.add(page, text="What it means")
        self.tiles = ttk.Frame(page)
        box = self.text_box = ttk.Frame(page)
        box.pack(fill="both", expand=True)
        self.text = tk.Text(box, wrap="word", height=14, relief="flat", padx=10, pady=8,
                            font=(base, 10), background="#fafafa")
        scroll = ttk.Scrollbar(box, command=self.text.yview)
        self.text.configure(yscrollcommand=scroll.set, state="disabled")
        self.text.tag_configure("head", font=(base, 11, "bold"))
        self.text.tag_configure("error", foreground="#a40000")
        self.text.tag_configure("note", foreground="#555")
        self.text.pack(side="left", fill="both", expand=True)
        scroll.pack(side="left", fill="y")
        self.chart_page = ttk.Frame(self.book)
        self.chart_label = ttk.Label(self.chart_page, anchor="center", cursor="hand2")
        self.chart_label.pack(fill="both", expand=True)
        self.chart_label.bind("<Button-1>", lambda _: self.chart_path and open_path(
            self.chart_path))
        self.chart_path: Optional[Path] = None
        Tooltip(self.chart_label, "Click to open the chart full size.")
        self.book.add(self.chart_page, text="Chart")
        self.book.tab(self.chart_page, state="disabled")

        self.results: Dict[str, Path] = {}
        self.marked: Optional[Path] = None
        self.done: "queue.Queue" = queue.Queue()
        self.proc = None
        self.root.after(100, self.poll)
        self.chart_image = None
        self.current: Task = TASKS[0]
        self.choose(self.prefs.get("task", TASKS[0].key))
        root.protocol("WM_DELETE_WINDOW", self.close)

    # ------------------------------------------------------------- helpers
    def poll(self) -> None:
        """Run what the worker threads handed over. Tk may only be touched
        from the thread running the window, so they never call it."""
        try:
            while True:
                self.done.get_nowait()()
        except queue.Empty:
            pass
        self.root.after(100, self.poll)

    def say(self, parts: Sequence[Tuple[str, str]]) -> None:
        self.text.configure(state="normal")
        self.text.delete("1.0", "end")
        for text, tag in parts:
            self.text.insert("end", text, tag)
        self.text.configure(state="disabled")
        self.book.select(0)

    def show_tiles(self, tiles: Sequence[Tuple[str, str]]) -> None:
        ttk = self.ttk
        for w in self.tiles.winfo_children():
            w.destroy()
        if tiles:
            self.tiles.pack(fill="x", padx=6, pady=(8, 2), before=self.text_box)
        else:
            self.tiles.pack_forget()
        for label, value in tiles:
            card = ttk.Frame(self.tiles, padding=(10, 6), relief="groove")
            card.pack(side="left", padx=(0, 8), fill="y")
            ttk.Label(card, text=label, style="TileLabel.TLabel", wraplength=200).pack(
                anchor="w")
            ttk.Label(card, text=value, style="TileValue.TLabel", wraplength=230).pack(
                anchor="w")

    def show_chart(self, chart: Optional[Path]) -> None:
        self.chart_image = None
        self.chart_path = chart
        if chart is None:
            self.chart_label.configure(image="", text="")
            self.book.tab(self.chart_page, state="disabled")
            return
        try:
            image = self.tk.PhotoImage(file=str(chart))
            self.root.update_idletasks()
            wide = max(self.book.winfo_width() - 20, 640)
            tall = max(self.book.winfo_height() - 40, 360)
            factor = max(1, math.ceil(max(image.width() / wide, image.height() / tall)))
            image = image.subsample(factor, factor)
        except Exception:  # noqa: BLE001 - no picture is better than a crash
            self.book.tab(self.chart_page, state="disabled")
            return
        self.chart_image = image
        self.chart_label.configure(image=image)
        self.book.tab(self.chart_page, state="normal")

    def remember(self) -> None:
        self.prefs.update(units=self.units_var.get(), marking=self.marking_var.get(),
                          task=self.current.key, index=self.index_var.get(),
                          **{"from": self.from_var.get(), "to": self.to_var.get()})
        chosen = self.file_var.get().strip().strip('"')
        if chosen and Path(chosen).exists():
            add_recent(self.prefs, self.current.key, chosen)
            self.file_box.configure(values=recent(self.prefs, self.current.key))
        save_prefs(self.prefs)

    def busy(self, on: bool, words: str = "") -> None:
        state = "disabled" if on else "normal"
        for b in (self.run_btn, self.check_btn, self.example_btn, self.template_btn,
                  *self.task_buttons.values()):
            b.configure(state=state)
        if not on:
            t = self.current
            self.example_btn.configure(state="normal" if t.demo else "disabled")
            self.template_btn.configure(state="normal" if t.template or t.key == "schedule"
                                        else "disabled")
        if on:
            self.progress.pack(side="left", padx=(10, 0))
            self.progress.start(12)
        else:
            self.progress.stop()
            self.progress.pack_forget()
        self.status.configure(text=words)

    def set_buttons(self, results: Dict[str, Path], problem: bool = False) -> None:
        self.results = results
        for b in self.result_buttons + self.problem_buttons:
            b.pack_forget()
        for key, b in (("report", self.report_btn), ("brief", self.brief_btn),
                       ("folder", self.folder_btn)):
            if key in results:
                b.pack(side="left", padx=(0, 8))
        if problem:
            if self.marked is not None:
                self.where_btn.pack(side="left", padx=(0, 8))
            if self.file_var.get().strip():
                self.fix_btn.pack(side="left", padx=(0, 8))

    def build_extras(self) -> None:
        ttk = self.ttk
        for w in self.extras.winfo_children():
            w.destroy()
        # an emptied frame keeps its old height in Tk, so take it out and put
        # it back only when there is something to show
        self.extras.pack_forget()
        t = self.current
        if t.key in ("inflate", "aoa"):
            self.extras.pack(fill="x", before=self.opts)
        if t.key == "inflate":
            row = ttk.Frame(self.extras)
            row.pack(fill="x", pady=3)
            ttk.Label(row, text="", width=3).pack(side="left")
            ttk.Label(row, text="Index file:").pack(side="left")
            ttk.Entry(row, textvariable=self.index_var).pack(side="left", fill="x",
                                                             expand=True, padx=6)
            ttk.Button(row, text="Choose...", command=self.browse_index).pack(side="left")
            row2 = ttk.Frame(self.extras)
            row2.pack(fill="x", pady=3)
            ttk.Label(row2, text="", width=3).pack(side="left")
            ttk.Label(row2, text="From:").pack(side="left")
            ttk.Combobox(row2, textvariable=self.from_var, values=BASES, width=10).pack(
                side="left", padx=(4, 12))
            ttk.Label(row2, text="To:").pack(side="left")
            ttk.Combobox(row2, textvariable=self.to_var, values=BASES, width=10).pack(
                side="left", padx=4)
            ttk.Label(row2, text="ty is then-year; by2026 is base-year 2026 (type any year).",
                      foreground="#555").pack(side="left", padx=8)
        elif t.key == "aoa":
            row = ttk.Frame(self.extras)
            row.pack(fill="x", pady=3)
            ttk.Label(row, text="", width=3).pack(side="left")
            ttk.Label(row, text="Status quo:").pack(side="left")
            self.sq_box = ttk.Combobox(row, textvariable=self.sq_var, width=32,
                                       postcommand=self.refresh_alternatives)
            self.sq_box.pack(side="left", padx=6)
            ttk.Label(row, text="The others are measured against it.",
                      foreground="#555").pack(side="left")
            self.sq_var.set(AS_IN_WORKBOOK)
            self.refresh_alternatives()

    def refresh_alternatives(self) -> None:
        if self.current.key != "aoa" or not hasattr(self, "sq_box"):
            return
        names = alternatives(Path(self.file_var.get().strip().strip('"'))) \
            if self.file_var.get().strip() else []
        self.sq_box.configure(values=[AS_IN_WORKBOOK] + names)
        if names and self.sq_var.get() not in names:
            self.sq_var.set(AS_IN_WORKBOOK)

    def extras_values(self) -> Dict[str, str]:
        status_quo = self.sq_var.get().strip()
        return {"index": self.index_var.get().strip(), "from": self.from_var.get().strip(),
                "to": self.to_var.get().strip(),
                "status_quo": "" if status_quo == AS_IN_WORKBOOK else status_quo}

    # ------------------------------------------------------------- actions
    def choose(self, key: str) -> None:
        if self.file_var.get().strip():
            self.remember()
        try:
            t = task(key)
        except StopIteration:
            t = TASKS[0]
        self.current = t
        self.title.configure(text=t.title)
        self.need.configure(text=f"What you need: {t.need}")
        self.get.configure(text=f"What you get: {t.get}")
        self.example_btn.configure(state="normal" if t.demo else "disabled")
        self.template_btn.configure(state="normal" if t.template or t.key == "schedule"
                                    else "disabled")
        self.units_box.configure(state="readonly" if t.units else "disabled")
        self.mark_entry.configure(state="disabled" if t.key == "inflate" else "normal")
        self.file_label.configure(text="Amounts:" if t.key == "inflate" else "Your file:")
        self.status.configure(text="")
        files = recent(self.prefs, t.key)
        self.file_box.configure(values=files)
        self.file_var.set(files[0] if files else "")
        self.build_extras()
        self.marked = None
        self.set_buttons({})
        self.show_tiles([])
        self.show_chart(None)
        hint = ("Start with 'See an example' to watch it work on invented numbers, or "
                "'Get the template' for a workbook to fill in with your own. Then "
                "'Check my workbook' and 'Run'.")
        if t.key == "open":
            hint = "Choose any file and press Run: cost-core works out what it is."
        self.say([(t.title + "\n\n", "head"), (hint, "")])

    def file_changed(self) -> None:
        self.refresh_alternatives()
        self.remember()

    def browse(self) -> None:
        from tkinter import filedialog

        start = self.file_var.get().strip()
        path = filedialog.askopenfilename(
            title="Choose your file", filetypes=list(self.current.filetypes),
            initialdir=str(Path(start).parent) if start else str(work_dir()))
        if path:
            self.file_var.set(path)
            self.file_changed()

    def browse_index(self) -> None:
        from tkinter import filedialog

        path = filedialog.askopenfilename(
            title="Choose the inflation index table",
            filetypes=[("Index table", "*.csv *.xlsx")], initialdir=str(work_dir()))
        if path:
            self.index_var.set(path)
            self.remember()

    def edit_file(self) -> None:
        from tkinter import messagebox

        path = Path(self.file_var.get().strip())
        if not self.file_var.get().strip() or not path.exists():
            messagebox.showinfo("cost-core", "Choose your file first.")
            return
        open_path(path)

    def template(self) -> None:
        from tkinter import messagebox

        t = self.current
        if t.key == "schedule":
            from cost_core.templates import SCHEDULE_HOWTO

            self.say([("Getting a schedule out of Microsoft Project\n\n", "head"),
                      (SCHEDULE_HOWTO, "")])
            return
        path = template_path(t)
        if path.exists():
            keep = messagebox.askyesnocancel(
                "cost-core", f"You already have {path.name} in {path.parent}.\n\n"
                "Yes: open the one you have.\nNo: start a fresh copy beside it "
                "(yours is kept).")
            if keep is None:
                return
            if keep:
                self.use_template(t, path)
                open_path(path)
                return
            path = fresh_name(path)
        try:
            path = make_template(t, path)
        except (OSError, ValueError) as e:
            messagebox.showerror("cost-core", f"Couldn't write the template: {e}")
            return
        self.use_template(t, path)
        self.say([(f"{path.name} is open in Excel\n\n", "head"),
                  (f"It's saved in {path.parent}. The yellow cells are yours: replace the "
                   "example with your own numbers, and add rows in the yellow space below. "
                   "Dark blue headings are required; hover over any heading for what goes "
                   "there. Save it, then press 'Check my workbook' and 'Run'.", "")])
        open_path(path)

    def use_template(self, t: Task, path: Path) -> None:
        if t.key == "inflate":
            from cost_core.templates import companion_path

            self.index_var.set(str(path))
            self.file_var.set(str(companion_path(path)))
        else:
            self.file_var.set(str(path))
        self.file_changed()

    def example(self) -> None:
        argv, out = demo_argv(self.current, self.marking_var.get())
        self.start(argv, out, "Running the example...", problem_file=None,
                   kind=self.current.key)

    def ready(self) -> Optional[Path]:
        """The chosen file, if there is one; says what's missing if not."""
        from tkinter import messagebox

        text = self.file_var.get().strip().strip('"')
        if not text:
            messagebox.showinfo("cost-core", "Choose your file first (step 2), or try "
                                "'See an example'.")
            return None
        path = Path(text)
        if not path.exists():
            messagebox.showerror("cost-core", f"There's no file at\n{path}")
            return None
        if self.current.key == "inflate" and not Path(self.index_var.get().strip()).exists():
            messagebox.showinfo("cost-core", "Choose the inflation index table too.")
            return None
        return path

    def check(self) -> None:
        path = self.ready()
        if path is None:
            return
        self.remember()
        self.busy(True, "Checking...")
        t = self.current

        def work():
            found = check_workbook(t, path)
            self.done.put(lambda: self.checked(path, found))

        threading.Thread(target=work, daemon=True).start()

    def checked(self, path: Path, found: Dict[str, List[str]]) -> None:
        self.busy(False)
        self.show_tiles([])
        self.show_chart(None)
        problems = found["problems"]
        self.marked = mark_problem(path, problems[0]) if problems else None
        if problems:
            parts = [(f"{len(problems)} thing{'s' if len(problems) > 1 else ''} to fix\n\n",
                      "head")]
            parts += [(f"• {p}\n\n", "error") for p in problems[:20]]
            if len(problems) > 20:
                parts.append((f"... and {len(problems) - 20} more.\n\n", "error"))
            parts.append((f"Fix {'them' if len(problems) > 1 else 'it'} in your workbook, "
                          "save it, and check again.", ""))
            self.status.configure(text="Needs fixing")
        else:
            parts = [("Nothing to fix: it's ready to run\n\n", "head")]
            self.status.configure(text="Ready")
        if found["notes"]:
            parts.append(("\nWorth knowing\n\n", "head"))
            parts += [(f"• {n}\n\n", "note") for n in found["notes"]]
        self.say(parts)
        self.set_buttons({}, problem=bool(problems))

    def run(self) -> None:
        from tkinter import messagebox

        path = self.ready()
        if path is None:
            return
        self.remember()
        out = results_dir(path)
        busy_files = locked(out)
        if busy_files:
            if not messagebox.askyesno(
                    "cost-core", f"{', '.join(busy_files)} from the last run "
                    f"{'is' if len(busy_files) == 1 else 'are'} open, so it can't be "
                    "replaced.\n\nSave this run's results to a new folder instead? (No: "
                    "close it yourself, then press Run again.)"):
                return
            out = next_free(out)
        argv = run_argv(self.current, path, self.units_var.get(), self.marking_var.get(),
                        out=out, extras=self.extras_values())
        self.start(argv, out, "Running... this can take a minute.", problem_file=path,
                   kind=None if self.current.key == "open" else self.current.key)

    def start(self, argv: List[str], out: Path, words: str,
              problem_file: Optional[Path], kind: Optional[str] = None) -> None:
        self.set_buttons({})
        self.show_tiles([])
        self.show_chart(None)
        self.busy(True, words)
        self.say([(words + "\n", "head")])
        units = self.units_var.get() if self.current.units else ""

        def keep(p):
            self.proc = p

        def work():
            code, output = run_cost_core(argv, started=keep)
            self.proc = None
            self.done.put(lambda: self.finish(code, output, out, units, problem_file, kind))

        threading.Thread(target=work, daemon=True).start()

    def finish(self, code: int, output: str, out: Path, units: str,
               problem_file: Optional[Path], kind: Optional[str] = None) -> None:
        self.busy(False)
        if code != 0:
            found = errors(output)
            self.marked = mark_problem(problem_file, found[0]) \
                if problem_file is not None and found else None
            self.say([("It stopped before finishing\n\n", "head")]
                     + [(e + "\n\n", "error") for e in found]
                     + [("Fix that in the workbook, save it, and press Run again."
                         + (" 'Show me where' opens a copy with the cell marked in red."
                            if self.marked else ""), "")])
            self.status.configure(text="Stopped")
            self.set_buttons({}, problem=True)
            return
        self.marked = None
        found = written(out)
        self.set_buttons(found)
        tiles, chart = answer(out, units, kind)
        self.show_tiles(tiles)
        self.show_chart(chart)
        lines = summary(output)
        parts = [("What this means\n\n", "head")]
        parts += [(f"• {ln}\n\n", "") for ln in lines] or [
            ("Done. The results are in the report and the folder.\n\n", "")]
        extra = notes(output)
        if extra:
            parts.append(("Worth knowing\n\n", "head"))
            parts += [(f"• {n}\n\n", "note") for n in extra]
        parts.append((f"Everything is saved in {out}.", "note"))
        self.say(parts)
        self.status.configure(text="Done")

    def show_where(self) -> None:
        if self.marked is not None and self.marked.exists():
            open_path(self.marked)

    def open_result(self, key: str) -> None:
        if key in self.results:
            open_path(self.results[key])

    def child(self, key: str, title: str, width: int, height: int):
        """A window over the main one, or the one already open, raised."""
        old = getattr(self, f"_{key}", None)
        if old is not None and old.winfo_exists():
            old.deiconify()
            old.lift()
            old.focus_force()
            return None
        win = self.tk.Toplevel(self.root)
        setattr(self, f"_{key}", win)
        win.title(title)
        win.transient(self.root)
        scale = max(self.root.winfo_fpixels("1i") / 96.0, 1.0)
        x = self.root.winfo_rootx() + 60
        y = self.root.winfo_rooty() + 40
        win.geometry(f"{int(width * scale)}x{int(height * scale)}+{x}+{y}")
        return win

    def glossary(self) -> None:
        tk, ttk = self.tk, self.ttk
        win = self.child("glossary", "What these words mean", 640, 620)
        if win is None:
            return
        text = tk.Text(win, wrap="word", padx=14, pady=10, font=(self.base, 10),
                       relief="flat")
        scroll = ttk.Scrollbar(win, command=text.yview)
        text.configure(yscrollcommand=scroll.set)
        text.tag_configure("term", font=(self.base, 10, "bold"))
        for term, meaning in GLOSSARY:
            text.insert("end", term + "\n", "term")
            text.insert("end", meaning + "\n\n")
        text.configure(state="disabled")
        text.pack(side="left", fill="both", expand=True)
        scroll.pack(side="left", fill="y")

    def settings(self) -> None:
        from tkinter import messagebox

        ttk = self.ttk
        win = self.child("settings", "Settings", 560, 330)
        if win is None:
            return
        frame = ttk.Frame(win, padding=16)
        frame.pack(fill="both", expand=True)

        def shortcut():
            try:
                path = make_shortcut()
                messagebox.showinfo("cost-core", f"Made {path.name} on your desktop. "
                                    "Double-click it to open cost-core.", parent=win)
            except OSError as e:
                messagebox.showerror("cost-core", str(e), parent=win)

        def send_to():
            from cost_core import opener

            try:
                opener.install_send_to()
                messagebox.showinfo("cost-core", "Done. Right-click any workbook, choose "
                                    "Send to, then cost-core, and it runs the right "
                                    "analysis.", parent=win)
            except (opener.OpenError, OSError) as e:
                messagebox.showerror("cost-core", str(e), parent=win)

        for label, about, action in (
                ("Put cost-core on my desktop", "A shortcut that opens this window.",
                 shortcut),
                ("Add cost-core to my right-click menu", "Right-click a file, Send to, "
                 "cost-core.", send_to),
                ("Open my cost-core folder", f"Templates and examples: {work_dir()}",
                 lambda: (work_dir().mkdir(parents=True, exist_ok=True),
                          open_path(work_dir())))):
            row = ttk.Frame(frame)
            row.pack(fill="x", pady=6)
            ttk.Button(row, text=label, width=36, command=action).pack(side="left")
            ttk.Label(row, text=about, wraplength=230, foreground="#555").pack(
                side="left", padx=10)
        ttk.Label(frame, text="Units and marking are remembered from the main window.",
                  foreground="#555").pack(anchor="w", pady=(14, 0))

    def close(self) -> None:
        if self.proc is not None and self.proc.poll() is None:
            from tkinter import messagebox

            if not messagebox.askyesno("cost-core", "A run is still going. Stop it and "
                                       "close?"):
                return
            self.proc.kill()
        self.remember()
        self.root.destroy()


def main(argv: Optional[Sequence[str]] = None) -> None:
    """Open the window, or with --shortcut, put a cost-core shortcut on the
    desktop."""
    args = list(sys.argv[1:] if argv is None else argv)
    if "--shortcut" in args:
        path = make_shortcut()
        print(f"Made {path}. Double-click it to open cost-core.")
        return
    import tkinter as tk

    sharp()
    root = tk.Tk()
    App(root)
    root.lift()
    root.attributes("-topmost", True)
    root.after(400, lambda: root.attributes("-topmost", False))
    root.focus_force()
    root.mainloop()


def sharp() -> None:
    """Draw at the screen's real resolution on Windows, not blurred up from
    96 dpi; Tk then scales its fonts itself."""
    if sys.platform == "win32":
        try:
            import ctypes

            ctypes.windll.shcore.SetProcessDpiAwareness(1)
        except (AttributeError, OSError):
            pass


if __name__ == "__main__":
    main()
